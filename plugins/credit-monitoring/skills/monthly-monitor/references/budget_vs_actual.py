#!/usr/bin/env python3
"""
Budget-vs-actual analysis module for the credit portfolio monitor.

Compares company ACTUALS (from the canonical monthly financials model) against a
management-prepared BUDGET (from a pinned budget workbook) for the latest month
and YTD, on Revenue, Gross Profit / Gross Margin %, EBITDA, and total Opex plus
key opex groups. Reports variance ($ and %) and flags material variances.

BORROWER-AGNOSTIC
=================
Nothing here is hardcoded to a particular borrower. Everything borrower-specific
is supplied by config (borrower-config.json + budget-map.json):

  * model_lines          canonical metric key -> the borrower's model row label
                         (revenue, gross_profit, ebitda_bridge; optional
                         financing_cost). Read from the model via monitor_lib.
  * opex_groups          display name -> model row label, for the opex breakdown.
                         Keys align 1:1 with the budget-map opex_groups.
  * comparability_items  OPTIONAL list of one-time items to also show ex-item
                         (e.g. a one-time refund sitting in COGS). Each item names
                         a model_line, the metrics it applies_to, and an add/
                         subtract operation. When the list is empty NO ex-item
                         lines are produced -- the general case.
  * ebitda.addback_financing_cost  whether BUDGET EBITDA also adds back the
                         financing_cost line when present in the budget-map.
  * analysis_notes       OPTIONAL free-text notes (e.g. a chart-of-accounts
                         reconciliation) appended verbatim to comparability_notes.

DESIGN
======
* ACTUALS come from the model THROUGH the shared library `monitor_lib` -- this
  module never opens the model itself, never parses numbers itself, never does
  its own period math. It uses ml.open_model / Model.get_line / ml.safe_float /
  ml.conclusion exactly like covenant_compliance.py.
* BUDGET layouts vary per borrower, so budget extraction is DATA-DRIVEN by a
  budget-map.json (file, sheet, period->column, metric->row, opex groups). The
  budget workbook is read with openpyxl directly (it is NOT the model and has no
  standardized layout), but every number is parsed with ml.safe_float.
* COMPARABILITY of EBITDA: ACTUAL EBITDA is read from the model's single EBITDA
  bridge row (model_lines.ebitda_bridge -- the SAME row covenant-compliance
  reads, one system-wide definition, not recomputed here). BUDGET EBITDA is
  computed with the IDENTICAL bridge formula -- Net Income + Interest Expense
  (+ Financing Cost if enabled and present) + Income Taxes + D&A - Interest
  Income -- so actual and budget EBITDA are like-for-like.

Usage (config-driven -- preferred):
    python budget_vs_actual.py --borrower-config <borrower-config.json> --out <output .json>
  model <- paths.model_output_path; budget-map <- paths.budget_map_path; the
  model_lines / opex_groups / comparability_items / ebitda blocks come from the
  borrower-config. The latest month defaults to the model's last populated month
  and YTD start to January of that year; both can be overridden with the flags.
  A default January start is trimmed to the months the model actually carries,
  and the output says so; an explicit --ytd-start the model cannot cover is a
  SKIP instead, since narrowing an analyst's own window is not the engine's call.

Usage (explicit paths -- still requires --borrower-config for model_lines):
    python budget_vs_actual.py \
        --borrower-config <borrower-config.json> \
        --model <model workbook .xlsx> \
        --budget-map <budget-map.json> \
        --out <output .json> \
        [--latest-month "May 2026"] [--ytd-start "Jan 2026"] [--borrower "..."]
"""

import argparse
import os
import sys

# --- Resolve the shared lib robustly: walk up from this file to the plugin
# root, then lib/ (same pattern as covenant_compliance.py). ---
_HERE = os.path.dirname(os.path.abspath(__file__))
_PLUGIN_ROOT = os.path.abspath(os.path.join(_HERE, "..", "..", ".."))
_LIB_DIR = os.path.join(_PLUGIN_ROOT, "lib")
if _LIB_DIR not in sys.path:
    sys.path.insert(0, _LIB_DIR)

import monitor_lib as ml
import rows_emit

import openpyxl
from openpyxl.utils import get_column_letter

# Default materiality: flag when |var%| >= 10% AND |var$| >= $50,000. The dollar
# side is a DOLLAR figure and the variances it is compared against are in the
# model's unit, so compute() puts it on that unit. Unconverted on a
# model in thousands the bar became $50,000,000: a 20% revenue overrun stopped
# being material, the flag list came back empty, and the output still printed
# the rule as "|var$| >= 50,000".
DEFAULT_MATERIAL_PCT = 0.10
DEFAULT_MATERIAL_ABS = 50000.0


ConfigError = ml.ConfigError


# ===========================================================================
# Config surface (borrower-agnostic) -- everything borrower-specific is here
# ===========================================================================
def load_analysis_config(cfg):
    """Pull the budget-vs-actual config surface out of a borrower-config dict.

    Returns a dict with model_lines / opex_groups / comparability_items /
    addback_financing_cost / analysis_notes. Validates the model lines this
    module requires and normalizes the comparability items.
    """
    model_lines = cfg.get("model_lines") or {}
    lines = {
        "revenue": ml.require_line(model_lines, "revenue", "budget-vs-actual"),
        "gross_profit": ml.require_line(model_lines, "gross_profit", "budget-vs-actual"),
        "ebitda_bridge": ml.require_line(model_lines, "ebitda_bridge", "budget-vs-actual"),
        # optional -- only used for a comparability note
        "financing_cost": model_lines.get("financing_cost"),
    }
    opex_groups = cfg.get("opex_groups") or {}
    if not isinstance(opex_groups, dict):
        raise ConfigError("opex_groups must be an object {display_name: model_label}")

    items = []
    for raw in (cfg.get("comparability_items") or []):
        if not raw.get("model_line"):
            raise ConfigError("each comparability_items entry needs a 'model_line'")
        items.append({
            "name": raw.get("name") or raw["model_line"],
            "model_line": raw["model_line"],
            "applies_to": [a.lower() for a in (raw.get("applies_to") or [])],
            "operation": (raw.get("operation") or "add").lower(),
            "note": raw.get("note", ""),
        })

    ebitda_cfg = cfg.get("ebitda") or {}
    return {
        "lines": lines,
        "opex_groups": opex_groups,
        "comparability_items": items,
        "addback_financing_cost": bool(ebitda_cfg.get("addback_financing_cost", False)),
        "analysis_notes": list(cfg.get("analysis_notes") or []),
    }


# ===========================================================================
# ACTUALS -- read from the model through monitor_lib only
# ===========================================================================
def _sum_line(model, label, months):
    """Sum a model IS row over a list of month labels (None cells -> 0)."""
    total = 0.0
    any_found = False
    for mth in months:
        v = model.get_line(label, month=mth)
        if v is not None:
            total += v
            any_found = True
    return total if any_found else None


def _comparability_adjustments(model, items, latest_month, ytd_months):
    """Resolve each comparability item's model-line value for the latest month
    and YTD, and accumulate the adjustment applied to each affected metric.

    An item's value is ADDED to (operation 'add', default) or SUBTRACTED from
    (operation 'subtract') the reported metric to strip the item's effect. The
    'add' default reproduces the canonical case of a contra amount (e.g. a refund
    sitting as a negative in COGS): adding it back removes the boost it gave to
    gross profit / EBITDA.
    """
    adj = {"gross_profit": {"latest": 0.0, "ytd": 0.0},
           "ebitda": {"latest": 0.0, "ytd": 0.0}}
    per_item = []
    for it in items:
        vm = model.get_line(it["model_line"], month=latest_month) or 0.0
        vy = _sum_line(model, it["model_line"], ytd_months) or 0.0
        sign = -1.0 if it["operation"] == "subtract" else 1.0
        sm, sy = sign * vm, sign * vy
        for metric in it["applies_to"]:
            if metric in adj:
                adj[metric]["latest"] += sm
                adj[metric]["ytd"] += sy
        per_item.append({"name": it["name"], "model_line": it["model_line"],
                         "applies_to": it["applies_to"], "operation": it["operation"],
                         "latest": round(vm, 2), "ytd": round(vy, 2),
                         "note": it["note"]})
    return adj, per_item


def actual_metrics(model, lines, opex_groups, items, latest_month, ytd_months):
    """Build the actuals block (latest month + YTD) entirely via monitor_lib."""
    rev_m = model.get_line(lines["revenue"], month=latest_month)
    rev_y = _sum_line(model, lines["revenue"], ytd_months)
    gp_m = model.get_line(lines["gross_profit"], month=latest_month)
    gp_y = _sum_line(model, lines["gross_profit"], ytd_months)

    # EBITDA: read straight from the model's single EBITDA bridge row -- the SAME
    # row covenant_compliance.py uses -- so the two modules report identical
    # actual EBITDA. NOT recomputed here.
    eb_m = model.get_line(lines["ebitda_bridge"], month=latest_month)
    eb_y = _sum_line(model, lines["ebitda_bridge"], ytd_months)

    # Comparability adjustments (empty list -> zero adjustment, no ex-item lines).
    adj, per_item = _comparability_adjustments(model, items, latest_month, ytd_months)

    # opex groups
    opex_m, opex_y = {}, {}
    for name, label in ml.iter_label_map(opex_groups):
        opex_m[name] = model.get_line(label, month=latest_month) or 0.0
        opex_y[name] = _sum_line(model, label, ytd_months) or 0.0
    total_opex_m = sum(opex_m.values())
    total_opex_y = sum(opex_y.values())

    fin_m = fin_y = None
    if lines.get("financing_cost"):
        fin_m = model.get_line(lines["financing_cost"], month=latest_month) or 0.0
        fin_y = _sum_line(model, lines["financing_cost"], ytd_months) or 0.0

    def _adj(base_m, base_y, metric):
        return ({"latest": (base_m + adj[metric]["latest"]) if base_m is not None else None,
                 "ytd": (base_y + adj[metric]["ytd"]) if base_y is not None else None})

    return {
        "revenue": {"latest": rev_m, "ytd": rev_y},
        "gross_profit": {"latest": gp_m, "ytd": gp_y},
        "gross_profit_adj": _adj(gp_m, gp_y, "gross_profit"),
        "ebitda": {"latest": eb_m, "ytd": eb_y},
        "ebitda_adj": _adj(eb_m, eb_y, "ebitda"),
        "total_opex": {"latest": total_opex_m, "ytd": total_opex_y},
        "opex_groups_latest": opex_m,
        "opex_groups_ytd": opex_y,
        "financing_cost": {"latest": fin_m, "ytd": fin_y},
        "comparability": per_item,
        "adjustments": adj,
    }


# ===========================================================================
# BUDGET -- data-driven from budget-map.json (openpyxl read; ml.safe_float parse)
# ===========================================================================
class BudgetReader:
    def __init__(self, budget_map, addback_financing_cost=False):
        self.map = budget_map
        self.addback_financing_cost = addback_financing_cost
        # Malformed budget-map -> ConfigError with the missing surface named,
        # never a raw KeyError (a raw KeyError is an engine bug, and main()'s
        # skip path deliberately does not catch it).
        for req in ("sheet", "period_columns", "metric_rows"):
            if req not in budget_map:
                raise ConfigError(
                    f"budget-map is missing required key {req!r}; see "
                    "config/budget-map.template.json")
        if not (budget_map.get("file") or budget_map.get("file_mount")):
            raise ConfigError(
                "budget-map needs 'file' (or 'file_mount') -- the budget workbook path")
        # AS-GIVEN workbook path (the configured string, before sandbox
        # resolution) -- what source refs record, so the downstream
        # artifact can join store hyperlinks onto it.
        self.given_path = str(budget_map.get("file") or budget_map.get("file_mount"))
        path = ml.resolve_path(
            budget_map.get("file_mount") or budget_map["file"])
        wb = openpyxl.load_workbook(path, data_only=True)
        self.sheet_name = budget_map["sheet"]
        try:
            self.ws = wb[budget_map["sheet"]]
        except KeyError:
            raise ConfigError(
                f"budget-map sheet {budget_map['sheet']!r} not found in the budget "
                f"workbook (sheets: {', '.join(wb.sheetnames)})") from None
        self.period_columns = budget_map["period_columns"]
        self.metric_rows = budget_map["metric_rows"]

    def _cell(self, row, col):
        return ml.safe_float(self.ws.cell(row=row, column=col).value) or 0.0

    def _rows_value(self, rows_spec, col):
        """A metric_rows / opex_groups entry -> value in a single column.
        rows_spec may be an int, a list of rows to sum, or {add:[],subtract:[]}.
        """
        if isinstance(rows_spec, int):
            return self._cell(rows_spec, col)
        if isinstance(rows_spec, list):
            return sum(self._cell(r, col) for r in rows_spec)
        if isinstance(rows_spec, dict):
            v = sum(self._cell(r, col) for r in rows_spec.get("add", []))
            v -= sum(self._cell(r, col) for r in rows_spec.get("subtract", []))
            return v
        raise ValueError(f"unrecognized rows spec: {rows_spec!r}")

    def _col(self, period_iso):
        """Column index for a period. A month the budget does not reach is a
        SKIP (main turns this into the skip record), and the reason is read by
        an analyst, so it says what the budget covers rather than asking for a
        column to be added: a budget whose later months are all #N/A has no
        month there to map, and setup excluded those columns on purpose."""
        try:
            return self.period_columns[period_iso]
        except KeyError:
            # period_columns also carries a '_comment' key on many deals, so the
            # months are the keys that read as one; period_stamp picks the last.
            months = [k for k in self.period_columns if ml.month_label(k)]
            through = (f"this budget runs through "
                       f"{ml.month_label(ml.period_stamp(months))}" if months
                       else "this budget names no months")
            raise ConfigError(
                f"there is no budgeted {ml.month_label(period_iso) or period_iso} "
                f"to compare the actuals against — {through}") from None

    def metric(self, name, period_iso):
        """Budget value for a named metric at one period (ISO month-end)."""
        col = self._col(period_iso)
        try:
            rows = self.metric_rows[name]
        except KeyError:
            raise ConfigError(
                f"budget-map metric_rows has no entry for {name!r}; add its row(s) "
                "to budget-map.json") from None
        return self._rows_value(rows, col)

    def metric_sum(self, name, period_iso_list):
        return sum(self.metric(name, p) for p in period_iso_list)

    def opex_group(self, rows_spec, period_iso):
        col = self._col(period_iso)
        return self._rows_value(rows_spec, col)

    def opex_group_sum(self, rows_spec, period_iso_list):
        return sum(self.opex_group(rows_spec, p) for p in period_iso_list)

    def _has_metric(self, name):
        return name in self.metric_rows

    def ebitda_components(self):
        """The (metric, sign) terms of the budget EBITDA bridge, exactly as
        ebitda() sums them -- ONE list drives both the value and its per-figure
        source refs, so the two can never drift apart."""
        comps = [("net_income", 1), ("interest_expense", 1), ("taxes", 1),
                 ("depreciation_amortization", 1)]
        if self.addback_financing_cost and self._has_metric("financing_cost"):
            comps.append(("financing_cost", 1))
        if self._has_metric("interest_income"):
            comps.append(("interest_income", -1))
        return comps

    def ebitda(self, period_iso_list):
        """Budget EBITDA computed with the IDENTICAL formula the model's EBITDA
        bridge uses, so budget is like-for-like with the bridge-sourced actual:
            Net Income + Interest Expense + Financing Cost (if enabled AND present
            in the budget) + Income Taxes + D&A - Interest Income.
        Driven by the budget-map metric_rows (see ebitda_components). The
        financing-cost add-back is gated by ebitda.addback_financing_cost in
        borrower-config."""
        return sum(sign * self.metric_sum(name, period_iso_list)
                   for name, sign in self.ebitda_components())

    # -- provenance: which cells a read touched, no value change ---
    @staticmethod
    def _rows_list(rows_spec):
        """Every workbook row a rows spec touches (add and subtract alike)."""
        if isinstance(rows_spec, int):
            return [rows_spec]
        if isinstance(rows_spec, list):
            return list(rows_spec)
        if isinstance(rows_spec, dict):
            return (list(rows_spec.get("add", []))
                    + list(rows_spec.get("subtract", [])))
        return []

    def rows_provenance(self, rows_spec, period_iso_list):
        """Per-row source-ref detail for a rows spec over the given periods:
        [{sheet, row, cells, period}, ...] naming exactly the cells
        _rows_value read. Pure provenance -- reads no values."""
        cols = sorted({self._col(p) for p in period_iso_list})
        contiguous = cols == list(range(cols[0], cols[-1] + 1))
        stamps = sorted(p[:7] for p in period_iso_list)
        period = stamps[0] if stamps[0] == stamps[-1] else f"{stamps[0]}..{stamps[-1]}"
        out = []
        for r in self._rows_list(rows_spec):
            if len(cols) == 1:
                cells = f"{get_column_letter(cols[0])}{r}"
            elif contiguous:
                cells = (f"{get_column_letter(cols[0])}{r}:"
                         f"{get_column_letter(cols[-1])}{r}")
            else:  # non-contiguous period->column mapping: list the cells
                cells = ",".join(f"{get_column_letter(c)}{r}" for c in cols)
            out.append({"sheet": self.sheet_name, "row": r,
                        "cells": cells, "period": period})
        return out


def budget_metrics(reader, latest_iso, ytd_iso_list):
    rev_m = reader.metric("revenue", latest_iso)
    rev_y = reader.metric_sum("revenue", ytd_iso_list)
    gp_m = reader.metric("gross_profit", latest_iso)
    gp_y = reader.metric_sum("gross_profit", ytd_iso_list)
    eb_m = reader.ebitda([latest_iso])
    eb_y = reader.ebitda(ytd_iso_list)

    opex_m, opex_y = {}, {}
    for name, spec in ml.iter_label_map(reader.map.get("opex_groups")):
        opex_m[name] = reader.opex_group(spec, latest_iso)
        opex_y[name] = reader.opex_group_sum(spec, ytd_iso_list)
    total_opex_m = sum(opex_m.values())
    total_opex_y = sum(opex_y.values())

    return {
        "revenue": {"latest": rev_m, "ytd": rev_y},
        "gross_profit": {"latest": gp_m, "ytd": gp_y},
        "ebitda": {"latest": eb_m, "ytd": eb_y},
        "total_opex": {"latest": total_opex_m, "ytd": total_opex_y},
        "opex_groups_latest": opex_m,
        "opex_groups_ytd": opex_y,
    }


# ===========================================================================
# Variance + materiality
# ===========================================================================
def variance(actual, budget):
    """(var_abs, var_pct) with var_abs = actual - budget. var_pct vs |budget|."""
    if actual is None or budget is None:
        return None, None
    var_abs = actual - budget
    denom = abs(budget)
    var_pct = (var_abs / denom) if denom else None
    return var_abs, var_pct


def is_material(var_abs, var_pct, mat_pct, mat_abs):
    if var_abs is None or var_pct is None:
        return False
    return abs(var_pct) >= mat_pct and abs(var_abs) >= mat_abs


def _line(metric, scope, actual, budget, mat_pct, mat_abs, lower_is_better=False,
          note=None, source_ref=None, calc=None):
    var_abs, var_pct = variance(actual, budget)
    material = is_material(var_abs, var_pct, mat_pct, mat_abs)
    # favorability: for cost lines lower actual is better.
    favorable = None
    if var_abs is not None:
        favorable = (var_abs < 0) if lower_is_better else (var_abs > 0)
    rec = {
        "metric": metric,
        "scope": scope,
        "actual": round(actual, 2) if actual is not None else None,
        "budget": round(budget, 2) if budget is not None else None,
        "var_abs": round(var_abs, 2) if var_abs is not None else None,
        "var_pct": round(var_pct, 4) if var_pct is not None else None,
        "favorable": favorable,
        "material": material,
    }
    if note:
        rec["note"] = note
    if source_ref:
        rec["source_ref"] = source_ref
    if calc:
        rec["calc"] = calc
    return rec


# ===========================================================================
# Per-figure provenance + calculation records
#
# Each result record carries source_ref (into the top-level source_files list)
# covering BOTH sides -- the model cells the actual came from and the budget
# cells + budget-map section the budget came from -- plus a plain-English calc
# note. validate_output.py gates the output on both.
# ===========================================================================
def _fmt(v):
    """Numbers in calc notes: thousands-separated, 2dp; None -> n/a."""
    return "n/a" if v is None else f"{v:,.2f}"


def _fmt_pct(v):
    return "n/a" if v is None else f"{v:.2%}"


def _rows_desc(rows_spec):
    """Plain-English description of a budget-map rows spec."""
    if isinstance(rows_spec, int):
        return f"row {rows_spec}"
    if isinstance(rows_spec, list):
        return "rows " + "+".join(str(r) for r in rows_spec)
    if isinstance(rows_spec, dict):
        add = "+".join(str(r) for r in rows_spec.get("add", []))
        sub = "+".join(str(r) for r in rows_spec.get("subtract", []))
        return f"rows {add}" + (f" minus {sub}" if sub else "")
    return f"rows {rows_spec!r}"


def _scope_how(scope, latest_month, ytd_months):
    """How a scope reads: the single latest month, or the summed YTD span."""
    if scope == "latest":
        return str(latest_month)
    return f"{ytd_months[0]}..{ytd_months[-1]} summed ({len(ytd_months)} mo)"


def _model_side(idx, model, label, scope, latest_month, ytd_months):
    """(refs, desc) for a model-row read at scope 'latest' or 'YTD' -- the
    exact cells actual_metrics read, via Model.provenance (window-aware)."""
    if scope == "latest":
        p = model.provenance(label, month=latest_month, window=1)
        how = _scope_how(scope, latest_month, ytd_months)
    else:
        # _sum_line reads each YTD month independently and skips the ones the
        # model can't serve, so the ref describes the SERVED span -- on a young
        # model that is shorter than the fiscal YTD span (and provenance would
        # rightly refuse the full-span window).
        served = [m for m in ytd_months if model.get_line(label, month=m) is not None]
        if not served:
            return [], f"model row '{label}' not found in the model (read as 0/None)"
        n = model._month_index(served[-1]) - model._month_index(served[0]) + 1
        p = model.provenance(label, month=served[-1], window=n)
        how = f"{served[0]}..{served[-1]} summed ({len(served)} mo)"
    if p is None:
        return [], f"model row '{label}' not found in the model (read as 0/None)"
    return ([idx.ref(model.given_path, **p)],
            f"model '{p['label']}' ({p['sheet']} {p['cells']}), {how}")


def _budget_side(idx, reader, budget_map_file, rows_spec, periods, section, how):
    """(refs, desc) for a budget read: the workbook cells actually read plus
    the budget-map section that decided which budget rows match the metric."""
    refs = []
    if rows_spec is None:
        desc = f"no budget-map entry ({section}); budget not available"
    else:
        for p in reader.rows_provenance(rows_spec, periods):
            refs.append(idx.ref(reader.given_path, **p))
        desc = f"budget '{reader.sheet_name}' {_rows_desc(rows_spec)}, {how}"
    refs.append(idx.ref(budget_map_file, section=section))
    return refs, desc


def _budget_ebitda_side(idx, reader, budget_map_file, periods, how):
    """(refs, desc) for budget EBITDA: one cell ref per bridge component read
    (labeled with the component name), plus the budget-map metric_rows block
    that mapped them. Mirrors BudgetReader.ebitda_components exactly."""
    refs, terms = [], []
    for name, sign in reader.ebitda_components():
        for p in reader.rows_provenance(reader.metric_rows[name], periods):
            refs.append(idx.ref(reader.given_path, label=name, **p))
        terms.append(f"{'-' if sign < 0 else '+'} {name}")
    refs.append(idx.ref(budget_map_file, section="metric_rows"))
    formula = " ".join(terms)[2:]  # drop the leading "+ "
    return refs, (f"budget EBITDA = {formula} per budget-map metric_rows, "
                  f"sheet '{reader.sheet_name}', {how}")


def _var_calc(actual, budget, a_desc, b_desc):
    """The one-line calc note for an actual-vs-budget variance record."""
    va, vp = variance(actual, budget)
    s = (f"variance {_fmt(va)} = actual {_fmt(actual)} ({a_desc}) - "
         f"budget {_fmt(budget)} ({b_desc})")
    if vp is not None:
        s += f"; var_pct {vp:.1%} = variance / |budget|"
    return s


# ===========================================================================
# Compute
# ===========================================================================
def compute(model, reader, acfg, latest_month, ytd_months, unit, borrower=None,
            mat_pct=DEFAULT_MATERIAL_PCT, mat_abs_dollars=DEFAULT_MATERIAL_ABS,
            budget_map_file=None, budget_file=None, ytd_note=None):
    """unit is the borrower's ReportingUnit (monitor_lib.reporting_unit(cfg)).
    mat_abs_dollars is the materiality bar in DOLLARS, as a person states it;
    every comparison below uses it on the model's unit, and every disclosure
    prints the dollar figure.

    budget_map_file is the AS-CONFIGURED path string of budget-map.json --
    the map is itself a source (it decided which budget rows match each model
    line), so refs point into it by section. Defaults to a generic placeholder
    so legacy callers still produce a valid output."""
    mat_abs = unit.from_dollars(mat_abs_dollars)
    lines = acfg["lines"]
    opex_groups = acfg["opex_groups"]
    items = acfg["comparability_items"]
    bridge = lines["ebitda_bridge"]

    latest_iso = ml.month_end_iso(latest_month)
    ytd_iso = [ml.month_end_iso(m) for m in ytd_months]

    A = actual_metrics(model, lines, opex_groups, items, latest_month, ytd_months)
    B = budget_metrics(reader, latest_iso, ytd_iso)

    # ---- Per-figure provenance: ONE SourceIndex for the run;
    # a (refs, desc) pair per metric/scope mirroring exactly the reads
    # actual_metrics / budget_metrics performed above. ----
    _mp = getattr(model, "given_path", None)
    # budget-map.json is our record of which budget rows match which model
    # lines; the budget itself is the borrower's file, so that is the upstream.
    idx = ml.SourceIndex(source_map=ml.model_source_map(_mp), model_path=_mp,
                         upstream_map=({budget_map_file: budget_file}
                                       if budget_map_file and budget_file else None))
    budget_map_file = budget_map_file or "budget-map.json"
    how_m = _scope_how("latest", latest_month, ytd_months)
    how_y = _scope_how("YTD", latest_month, ytd_months)

    def AP(label, scope):
        """Actual-side provenance: the model row/cells the value came from."""
        return _model_side(idx, model, label, scope, latest_month, ytd_months)

    def BP(name, scope):
        """Budget-side provenance for a metric_rows entry."""
        return _budget_side(idx, reader, budget_map_file,
                            reader.metric_rows.get(name),
                            [latest_iso] if scope == "latest" else ytd_iso,
                            f"metric_rows.{name}",
                            how_m if scope == "latest" else how_y)

    def BPG(name, scope):
        """Budget-side provenance for an opex_groups entry."""
        return _budget_side(idx, reader, budget_map_file,
                            (reader.map.get("opex_groups") or {}).get(name),
                            [latest_iso] if scope == "latest" else ytd_iso,
                            f"opex_groups.{name}",
                            how_m if scope == "latest" else how_y)

    rev_a_m, rev_b_m = AP(lines["revenue"], "latest"), BP("revenue", "latest")
    rev_a_y, rev_b_y = AP(lines["revenue"], "YTD"), BP("revenue", "YTD")
    gp_a_m, gp_b_m = AP(lines["gross_profit"], "latest"), BP("gross_profit", "latest")
    gp_a_y, gp_b_y = AP(lines["gross_profit"], "YTD"), BP("gross_profit", "YTD")
    eb_a_m, eb_a_y = AP(bridge, "latest"), AP(bridge, "YTD")
    eb_b_m = _budget_ebitda_side(idx, reader, budget_map_file, [latest_iso], how_m)
    eb_b_y = _budget_ebitda_side(idx, reader, budget_map_file, ytd_iso, how_y)

    # Which metrics carry a comparability adjustment (drives whether ex-item
    # lines are emitted at all -- the general, no-items case emits none).
    gp_items = [it for it in A["comparability"] if "gross_profit" in it["applies_to"]]
    eb_items = [it for it in A["comparability"] if "ebitda" in it["applies_to"]]
    item_names = ", ".join(dict.fromkeys(  # de-dup, preserve order
        it["name"] for it in (gp_items + eb_items))) or "comparability items"

    results = []
    exceptions = []

    def add(rec):
        results.append(rec)
        if rec.get("material"):
            exceptions.append({
                "metric": rec["metric"], "scope": rec["scope"],
                "actual": rec["actual"], "budget": rec["budget"],
                "var_abs": rec["var_abs"], "var_pct": rec["var_pct"],
                "favorable": rec["favorable"],
                "rule": f"|var%|>={mat_pct:.0%} AND |var$|>=${mat_abs_dollars:,.0f}",
            })

    def cmp_line(metric, scope, a_val, b_val, a_prov, b_prov, **kw):
        """A variance record carrying its own source refs + calc note: refs
        cover BOTH sides (model cells; budget cells + budget-map section)."""
        (a_refs, a_desc), (b_refs, b_desc) = a_prov, b_prov
        return _line(metric, scope, a_val, b_val, mat_pct, mat_abs,
                     source_ref=a_refs + b_refs,
                     calc=_var_calc(a_val, b_val, a_desc, b_desc), **kw)

    # Revenue
    add(cmp_line("Revenue", "latest", A["revenue"]["latest"],
                 B["revenue"]["latest"], rev_a_m, rev_b_m))
    add(cmp_line("Revenue", "YTD", A["revenue"]["ytd"],
                 B["revenue"]["ytd"], rev_a_y, rev_b_y))

    # Gross Profit (as-reported)
    gp_note = None
    if gp_items:
        gp_note = (f"As-reported; includes comparability item(s) [{item_names}] "
                   "that distort the underlying figure. See ex-item line.")
    add(cmp_line("Gross Profit (as-reported)", "latest", A["gross_profit"]["latest"],
                 B["gross_profit"]["latest"], gp_a_m, gp_b_m))
    add(cmp_line("Gross Profit (as-reported)", "YTD", A["gross_profit"]["ytd"],
                 B["gross_profit"]["ytd"], gp_a_y, gp_b_y, note=gp_note))

    # Gross Profit (YTD ex-item) -- only when comparability items apply to GP.
    gp_item_refs = []
    if gp_items:
        gp_item_refs = [r for it in gp_items for r in AP(it["model_line"], "YTD")[0]]
        gp_adj = A["adjustments"]["gross_profit"]["ytd"]
        gp_ex = A["gross_profit_adj"]["ytd"]
        add(_line(f"Gross Profit (YTD ex-{item_names})", "YTD",
                  gp_ex, B["gross_profit"]["ytd"],
                  mat_pct, mat_abs,
                  note="Removes the one-time comparability item(s); the "
                       "underlying comparable YTD gross profit.",
                  source_ref=gp_a_y[0] + gp_item_refs + gp_b_y[0],
                  calc=(f"ex-item YTD gross profit {_fmt(gp_ex)} = as-reported "
                        f"{_fmt(A['gross_profit']['ytd'])} "
                        f"{'+' if gp_adj >= 0 else '-'} {_fmt(abs(gp_adj))} "
                        f"adjustment from comparability item(s) [{item_names}] "
                        f"(each item's model line summed over "
                        f"{ytd_months[0]}..{ytd_months[-1]}); variance = actual "
                        f"- budget {_fmt(B['gross_profit']['ytd'])} ({gp_b_y[1]})")))

    # Gross Margin % (as-reported; plus ex-item YTD when applicable)
    def gm(gp, rev):
        return (gp / rev) if rev else None
    a_gm_m = gm(A["gross_profit"]["latest"], A["revenue"]["latest"])
    b_gm_m = gm(B["gross_profit"]["latest"], B["revenue"]["latest"])
    a_gm_y = gm(A["gross_profit"]["ytd"], A["revenue"]["ytd"])
    b_gm_y = gm(B["gross_profit"]["ytd"], B["revenue"]["ytd"])

    def _margin_row(label, scope, a, b, note, source_ref=None, calc=None):
        # `unit` says what KIND of quantity this is, so a row of it carries no
        # money scale: a margin and its variance are percentage points, while
        # every other result in this module is money on the model's unit.
        rec = {"metric": label, "scope": scope, "unit": "pct",
               "actual": round(a, 4) if a is not None else None,
               "budget": round(b, 4) if b is not None else None,
               "var_abs": round(a - b, 4) if (a is not None and b is not None) else None,
               "var_pct": None,
               "favorable": (a >= b) if (a is not None and b is not None) else None,
               "material": False, "note": note}
        if source_ref:
            rec["source_ref"] = source_ref
        if calc:
            rec["calc"] = calc
        return rec

    def _margin_calc(a, b, a_gp, a_rev, b_gp, b_rev, a_how):
        return (f"actual margin {_fmt_pct(a)} = gross profit {_fmt(a_gp)} / "
                f"revenue {_fmt(a_rev)} (model, {a_how}); budget margin "
                f"{_fmt_pct(b)} = {_fmt(b_gp)} / {_fmt(b_rev)} (budget "
                f"'{reader.sheet_name}'); var_abs = margin points (actual - budget)")

    results.append(_margin_row(
        "Gross Margin %", "latest", a_gm_m, b_gm_m,
        "margin points (actual - budget)",
        source_ref=gp_a_m[0] + rev_a_m[0] + gp_b_m[0] + rev_b_m[0],
        calc=_margin_calc(a_gm_m, b_gm_m, A["gross_profit"]["latest"],
                          A["revenue"]["latest"], B["gross_profit"]["latest"],
                          B["revenue"]["latest"], how_m)))
    results.append(_margin_row(
        "Gross Margin % (as-reported)", "YTD", a_gm_y, b_gm_y,
        "margin points; as-reported",
        source_ref=gp_a_y[0] + rev_a_y[0] + gp_b_y[0] + rev_b_y[0],
        calc=_margin_calc(a_gm_y, b_gm_y, A["gross_profit"]["ytd"],
                          A["revenue"]["ytd"], B["gross_profit"]["ytd"],
                          B["revenue"]["ytd"], how_y)))
    if gp_items:
        a_gm_y_ex = gm(A["gross_profit_adj"]["ytd"], A["revenue"]["ytd"])
        results.append(_margin_row(
            f"Gross Margin % (ex-{item_names})", "YTD",
            a_gm_y_ex, b_gm_y, "margin points; ex-comparability item(s)",
            source_ref=gp_a_y[0] + gp_item_refs + rev_a_y[0] + gp_b_y[0] + rev_b_y[0],
            calc=(f"actual ex-item margin {_fmt_pct(a_gm_y_ex)} = ex-item YTD "
                  f"gross profit {_fmt(A['gross_profit_adj']['ytd'])} / YTD "
                  f"revenue {_fmt(A['revenue']['ytd'])} (model); budget margin "
                  f"{_fmt_pct(b_gm_y)} = {_fmt(B['gross_profit']['ytd'])} / "
                  f"{_fmt(B['revenue']['ytd'])} (budget '{reader.sheet_name}'); "
                  f"var_abs = margin points (actual - budget)")))

    # EBITDA -- actual sourced from the model's single EBITDA bridge row (the same
    # row covenant-compliance reads); budget mirrors the bridge formula.
    add(cmp_line("EBITDA (as-reported)", "latest", A["ebitda"]["latest"],
                 B["ebitda"]["latest"], eb_a_m, eb_b_m,
                 note=(f"Actual EBITDA is read from the model's EBITDA bridge row "
                       f"'{bridge}' (single, system-wide definition), not recomputed "
                       "-- so it equals what covenant-compliance reads.")))
    eb_note_ytd = (f"Actual EBITDA read from the model's EBITDA bridge row "
                   "(single definition shared with covenant-compliance), not "
                   "recomputed.")
    if eb_items:
        eb_note_ytd += " As-reported YTD EBITDA is distorted by comparability item(s). See ex-item line."
    add(cmp_line("EBITDA (as-reported)", "YTD", A["ebitda"]["ytd"],
                 B["ebitda"]["ytd"], eb_a_y, eb_b_y, note=eb_note_ytd))
    if eb_items:
        eb_item_refs = [r for it in eb_items for r in AP(it["model_line"], "YTD")[0]]
        eb_adj = A["adjustments"]["ebitda"]["ytd"]
        eb_ex = A["ebitda_adj"]["ytd"]
        add(_line(f"EBITDA (YTD ex-{item_names})", "YTD", eb_ex,
                  B["ebitda"]["ytd"], mat_pct, mat_abs,
                  note="Removes the one-time comparability item(s); the "
                       "underlying comparable YTD EBITDA.",
                  source_ref=eb_a_y[0] + eb_item_refs + eb_b_y[0],
                  calc=(f"ex-item YTD EBITDA {_fmt(eb_ex)} = as-reported "
                        f"{_fmt(A['ebitda']['ytd'])} "
                        f"{'+' if eb_adj >= 0 else '-'} {_fmt(abs(eb_adj))} "
                        f"adjustment from comparability item(s) [{item_names}] "
                        f"(each item's model line summed over "
                        f"{ytd_months[0]}..{ytd_months[-1]}); variance = actual "
                        f"- budget {_fmt(B['ebitda']['ytd'])} ({eb_b_y[1]})")))

    # Total Opex (normalized -- like-for-like per the budget-map's opex groups)
    a_opex_m = {n: AP(lab, "latest") for n, lab in ml.iter_label_map(opex_groups)}
    a_opex_y = {n: AP(lab, "YTD") for n, lab in ml.iter_label_map(opex_groups)}
    budget_group_names = [n for n, _v in ml.iter_label_map(reader.map.get("opex_groups"))]
    b_opex_m = {n: BPG(n, "latest") for n in budget_group_names}
    b_opex_y = {n: BPG(n, "YTD") for n in budget_group_names}
    group_names = ", ".join(n for n, _v in ml.iter_label_map(opex_groups)) or "none configured"
    n_groups = sum(1 for _ in ml.iter_label_map(opex_groups))

    def _tot_a_prov(prov_map, how):
        return ([r for p in prov_map.values() for r in p[0]],
                f"sum of the {n_groups} configured opex-group model "
                f"rows [{group_names}], {how}")

    def _tot_b_prov(prov_map, how):
        refs = [r for p in prov_map.values() for r in p[0]]
        if not refs:  # no groups mapped: the (empty) map block is the source
            refs.append(idx.ref(budget_map_file, section="opex_groups"))
        return (refs, f"sum of the budget-map opex_groups rows, budget "
                      f"'{reader.sheet_name}', {how}")

    add(cmp_line("Total Opex (normalized)", "latest", A["total_opex"]["latest"],
                 B["total_opex"]["latest"], _tot_a_prov(a_opex_m, how_m),
                 _tot_b_prov(b_opex_m, how_m), lower_is_better=True))
    add(cmp_line("Total Opex (normalized)", "YTD", A["total_opex"]["ytd"],
                 B["total_opex"]["ytd"], _tot_a_prov(a_opex_y, how_y),
                 _tot_b_prov(b_opex_y, how_y), lower_is_better=True))

    # Key opex groups
    for name, _lab in ml.iter_label_map(opex_groups):
        add(cmp_line(f"Opex: {name}", "latest", A["opex_groups_latest"][name],
                     B["opex_groups_latest"].get(name),
                     a_opex_m[name], b_opex_m.get(name) or BPG(name, "latest"),
                     lower_is_better=True))
        add(cmp_line(f"Opex: {name}", "YTD", A["opex_groups_ytd"][name],
                     B["opex_groups_ytd"].get(name),
                     a_opex_y[name], b_opex_y.get(name) or BPG(name, "YTD"),
                     lower_is_better=True))

    # --- Notes (built generically from config, not hardcoded) ---
    # A shortened YTD window leads, because it changes what every YTD figure
    # below it means.
    notes = [ytd_note] if ytd_note else []
    for it in A["comparability"]:
        notes.append(
            f"Comparability item '{it['name']}' (model line '{it['model_line']}'): "
            f"{it['ytd']:,.0f} YTD, applied to {it['applies_to'] or 'none'} "
            f"({it['operation']}). {it['note']}".strip())
    if A["financing_cost"]["ytd"]:
        notes.append(
            f"Financing cost (model line '{lines['financing_cost']}'): "
            f"{A['financing_cost']['ytd']:,.0f} YTD. Confirm whether it sits above "
            "or below the EBITDA bridge per the credit agreement's definition.")
    notes.append(
        f"Actual EBITDA is sourced from the model's single EBITDA bridge row "
        f"'{bridge}' -- the SAME row covenant-compliance reads -- so the two "
        "modules report identical actual EBITDA (one system-wide definition). "
        "Budget EBITDA is computed with the identical bridge formula (Net Income "
        "+ Interest Expense + Financing Cost if enabled and present + Income Taxes "
        "+ D&A - Interest Income) so it is like-for-like. The financing-cost "
        f"add-back is {'ENABLED' if acfg['addback_financing_cost'] else 'DISABLED'} "
        "via ebitda.addback_financing_cost in borrower-config.")
    notes.extend(acfg.get("analysis_notes") or [])

    return ml.conclusion(
        module="budget-vs-actual",
        period=latest_month,   # normalized to the ISO YYYY-MM stamp
        results=results,
        exceptions=exceptions,
        source_index=idx,
        unit=unit,
        engine=("data-driven budget-map v4 (config-driven, borrower-agnostic) "
                "+ per-figure source refs/calc"),
        borrower=borrower or reader.map.get("borrower"),
        latest_month=latest_month,
        ytd_months=ytd_months,
        materiality_rule=(f"|var%| >= {mat_pct:.0%} AND |var$| >= "
                          f"${mat_abs_dollars:,.0f}"),
        budget_source={"file": reader.map.get("file"),
                       "sheet": reader.map.get("sheet"),
                       "basis": "management-prepared / unaudited"},
        comparability_notes=notes,
        comparability_items_ytd=[{"name": it["name"], "ytd": it["ytd"]}
                                 for it in A["comparability"]],
        financing_cost_ytd=(round(A["financing_cost"]["ytd"], 2)
                            if A["financing_cost"]["ytd"] is not None else None),
        methodology=("Actuals from the company-prepared/unaudited monthly financials "
                     "model (read via monitor_lib). Budget is management-prepared/"
                     "unaudited, extracted via budget-map.json. ACTUAL EBITDA is read "
                     "from the model's single EBITDA bridge row (the same definition "
                     "covenant-compliance uses); budget EBITDA mirrors that bridge "
                     "formula for like-for-like comparison."),
    )


# ===========================================================================
# CLI
# ===========================================================================
def _derive_ytd_months(latest_month, ytd_start):
    """Inclusive month-label list from ytd_start through latest_month."""
    start_y, start_m = ml._parse_label(ytd_start)
    end_y, end_m = ml._parse_label(latest_month)
    n = (end_y - start_y) * 12 + (end_m - start_m) + 1
    if n < 1:
        raise ValueError("ytd_start is after latest_month")
    return ml.trailing_months(latest_month, n)


def _ytd_within_model(model, revenue_label, ytd_months):
    """The trailing run of `ytd_months` the model actually carries, ending at the
    latest month.

    A YTD window the model does not reach is the one comparison error this module
    cannot survive: the actual side reads each month independently and skips the
    ones the model has no column for, while the budget side sums every month asked
    of it. A default January start against a model that begins in March therefore
    compared four months of actual against six months of budget, and every YTD
    variance in the output was that gap rather than performance.

    Contiguous from the latest month backwards, not merely "the months present":
    a YTD figure with a hole in the middle is not a period a reader can name.
    """
    covered = set()
    rev = model.get_line(revenue_label)
    for i, mo in enumerate(model.months or []):
        if rev and i < len(rev) and rev[i] is not None:
            covered.add(mo)
    kept = []
    for mo in reversed(ytd_months):
        if mo not in covered:
            break
        kept.append(mo)
    return list(reversed(kept))


def resolve_ytd_window(model, revenue_label, latest_month, ytd_start=None):
    """The YTD window to compare on, and the note the output owes a reader.

    Returns (ytd_months, note). Raises ConfigError -- which main turns into a skip
    record -- when no honest window exists.

    An EXPLICIT ytd_start is the analyst's, so a window the model cannot cover is
    refused rather than narrowed. The default January start is the engine's own
    convenience, so the engine trims it to what the model carries and says so in
    the note: both sides then sum the same months, which is the whole point.
    """
    if ytd_start:
        asked = _derive_ytd_months(latest_month, ytd_start)
        covered = _ytd_within_model(model, revenue_label, asked)
        if len(covered) != len(asked):
            have = (f"only reads from {covered[0]}" if covered
                    else "does not read that far back")
            raise ConfigError(
                f"the requested year-to-date window starts at {asked[0]}, but the "
                f"financials model {have} — either start the window there, or add "
                f"the earlier months to the model first")
        return asked, None

    end_y, _ = ml._parse_label(latest_month)
    asked = _derive_ytd_months(latest_month, f"Jan {end_y}")
    kept = _ytd_within_model(model, revenue_label, asked)
    if not kept:
        raise ConfigError(
            f"the financials model carries no figures for {latest_month}, so there "
            "is no year-to-date period to compare the budget against")
    if len(kept) == len(asked):
        return kept, None
    note = (f"Year-to-date covers {kept[0]} to {kept[-1]} "
            f"({len(kept)} month{'s' if len(kept) > 1 else ''}), not from "
            f"{asked[0]}: the financials model does not carry the earlier months of "
            f"the year. Both the actual and the budget side are summed over this "
            f"same shorter window, so the YTD comparison is like-for-like — but it "
            f"is not a full year to date.")
    return kept, note


def _last_populated_month(model, revenue_label):
    """The model's last populated month label: the latest month whose Revenue
    row has a value, falling back to the last month column. Used only to pick the
    reporting period -- it does not affect any computed figure."""
    rev = model.get_line(revenue_label)  # per-month list, oldest -> newest
    if rev:
        for i in range(len(rev) - 1, -1, -1):
            if rev[i] is not None:
                return model.months[i]
    return model.months[-1]


def main():
    ap = argparse.ArgumentParser(description="Budget-vs-actual analysis")
    ap.add_argument("--borrower-config",
                    help="borrower-config.json (supplies model_lines / opex_groups "
                         "/ comparability_items / ebitda; REQUIRED for model line mapping)")
    ap.add_argument("--model")
    ap.add_argument("--budget-map")
    ap.add_argument("--out", required=True)
    ap.add_argument("--latest-month",
                    help='override; default = model last populated month')
    ap.add_argument("--ytd-start",
                    help='override; default = January of the latest month\'s year')
    ap.add_argument("--borrower")
    ap.add_argument("--run-id",
                    help="the run this output belongs to, minted once by the "
                         "orchestrator for the whole run (or set CREDIT_MONITOR_RUN_ID)")
    args = ap.parse_args()
    if args.run_id:
        os.environ["CREDIT_MONITOR_RUN_ID"] = args.run_id   # run_stamp() reads it

    model_path = args.model
    budget_map_path = args.budget_map
    borrower = args.borrower
    cfg = {}

    if args.borrower_config:
        cfg = ml.load_config(args.borrower_config)
        borrower = borrower or cfg.get("borrower_name")
        paths = cfg.get("paths", {})
        model_path = model_path or paths.get("model_output_path")
        budget_map_path = budget_map_path or paths.get("budget_map_path")

    if not args.borrower_config:
        ap.error("--borrower-config is required (it supplies model_lines / "
                 "opex_groups / comparability_items / ebitda).")
    if not (model_path and budget_map_path):
        ap.error("could not resolve --model / --budget-map (supply them or a "
                 "--borrower-config with paths.model_output_path + budget_map_path).")

    latest_month = args.latest_month   # best-known period if we skip below

    try:
        acfg = load_analysis_config(cfg)
        # Which unit this model's figures are in; the materiality bar is stated
        # in dollars and has to meet them on the same footing.
        unit = ml.reporting_unit(cfg)

        model = ml.open_model(model_path)
        # Period first: once the model is open the run's period is known, so a
        # budget-map failure below still writes a skip record WITH a period.
        latest_month = args.latest_month or _last_populated_month(model, acfg["lines"]["revenue"])

        budget_map = ml.load_config(budget_map_path)
        reader = BudgetReader(budget_map,
                              addback_financing_cost=acfg["addback_financing_cost"])

        # The YTD window has to be a span BOTH sides cover, or every YTD variance
        # reports the gap between the windows instead of performance.
        ytd_months, ytd_note = resolve_ytd_window(
            model, acfg["lines"]["revenue"], latest_month, args.ytd_start)
    except (ConfigError, ValueError, FileNotFoundError) as e:
        # A known bad-input condition (missing config surface, unmapped budget
        # month, unreadable model, ...) is a SKIP RECORD, not a traceback: the
        # output stays valid (validate_output.py accepts it) and says why the
        # module could not compute. Never a plausible guess.
        # Deliberately NOT catching KeyError: every known bad-input KeyError
        # site raises ConfigError with a clear message (BudgetReader.__init__,
        # _col, metric), so a raw KeyError here is an engine BUG and must
        # crash loudly, not wear a skip record's costume.
        out = ml.skip_record("budget-vs-actual", e, period=latest_month, borrower=borrower)
    else:
        # compute() is deliberately OUTSIDE the broad guard above, matching
        # the sibling engines: an engine bug must crash loudly, not wear a
        # skip record's costume. The ONE exception is ConfigError — the
        # engine's own typed bad-input signal — because budget reads part of
        # its input surface lazily during compute (_col / metric hit the
        # workbook per month): an unmapped budget month is a documented skip
        # condition, not a bug. budget_map_file: the AS-CONFIGURED budget-map
        # path string (pre-sandbox-resolution) that source refs record.
        try:
            out = compute(model, reader, acfg, latest_month, ytd_months, unit,
                          borrower=borrower, budget_map_file=budget_map_path,
                          budget_file=(cfg.get("paths") or {}).get("budget_path"),
                          ytd_note=ytd_note)
        except ConfigError as e:
            out = ml.skip_record("budget-vs-actual", e, period=latest_month, borrower=borrower)

    out_path = ml.write_output(out, args.out)
    rows_emit.write_rows(out, out_path, deal_name=borrower)
    if out.get("skipped"):
        print(f"Wrote {out_path}: SKIPPED (not computed) -- {out['reason']}")
    else:
        print(f"Wrote {out_path}: {len(out['results'])} results, "
              f"{len(out['exceptions'])} exceptions")


if __name__ == "__main__":
    main()
