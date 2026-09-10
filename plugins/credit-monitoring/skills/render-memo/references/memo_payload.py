#!/usr/bin/env python3
"""Gather ONE payload carrying every figure the monitoring memo prints or draws.

WHY THIS EXISTS
A memo assembled separately from archived results can drift from its source figures
or retain data from a worked example. This payload is the memo's only data source.

So no figure reaches the memo any other way than through this file. The renderer
composes HTML from what this emits; it reads no archive and computes no figure.

WHAT IT EMITS
  identity   borrower, period, model path, and the exact archive files the payload
             was built from -- validate_memo.py checks the memo against this
  windows    period_windows.windows(): the display and analysis windows, the
             prior-year T3M pair, quarters, fiscal years, annualisation factors
  sections   one entry per block of the five-tab report, in document order, each
             with `figures` (value + source_ref + calc note + the store URL,
             joined once here) and whatever else that block shows -- the deal
             record's company facts and key terms, the covenant grid with its
             preamble, history, far-out list and stress inputs, the statement
             tables, the source-document registry, the mechanical data-quality
             notes and the numbered follow-ups
  reads      the archived module outputs for the period, the borrower's model,
             and the control files its settings name: the deal record, the
             distilled covenant terms, the pledged-asset map, the record of
             which package filled each model column, and the restatement log
  groups     a component set and its stated total emitted as ONE traced group, so
             a reader cannot foot the parts to a different number than the total
  absent     one record per thing the memo cannot fill: a module that was skipped,
             a covenant that could not be tested, a month only partly readable, a
             column with no data. This is the memo's "could not cover" section
  degrade    per-figure link degradation, in the vocabulary report-structure's
             "Degradation" section already fixes: no store, no url_base, no
             upstream

DETERMINISM
Two runs over the same archive produce a byte-identical payload: no timestamps, no
run ids minted here, archive files read in sorted order, JSON written with sorted
keys. `render_memo.py` inherits that, which is how re-rendering the same run twice
gives the same file.

UNITS
Money is emitted in DOLLARS. Each archived envelope says what one model figure
stands for (`reporting_unit.scale`), and a figure says what kind of quantity it is
in its own `unit` field -- `usd` scales, `ratio` / `%` / `x` / `days` do not. A
memo that understates a borrower a thousand-fold reads as a collapsing business,
and a margin multiplied by a thousand claims a 350x margin, so the scaling is done
here, once, against the figure's own declared kind.

Usage:
  python memo_payload.py --borrower-config <borrower-config.json> \
                         --period YYYY-MM --out <memo-payload.json>
                         [--archive <dir>] [--model <path>] [--trend-months 18]
"""
import argparse
import json
import os
import re
import sys
import tempfile


def _plugin_lib():
    here = os.path.dirname(os.path.abspath(__file__))
    for up in (os.path.join(here, "..", "..", "..", "lib"),
               os.path.join(here, "..", "..", "lib")):
        if os.path.isdir(up):
            sys.path.insert(0, os.path.abspath(up))
            return
    raise RuntimeError("cannot locate the plugin lib directory")


_plugin_lib()
import monitor_lib as ml  # noqa: E402
import openpyxl           # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _HERE)
import period_windows  # noqa: E402
import performance_views  # noqa: E402

SCHEMA = "credit-monitor-memo-payload/v1"

# Money scales by the envelope's reporting unit; a ratio, a percentage, a
# multiple and a day count do not. An unknown kind is left alone -- guessing
# is what produces the thousand-fold error this list exists to prevent.
SCALES = {"usd", "dollars", "$"}

# Which memo section each module's figures belong to. `three-statement-analysis`
# is split further by each result's own `section` field.
MODULE_SECTIONS = {
    "covenant-compliance": "covenant_detail",
    "budget-vs-actual": "snapshot_performance",
    "ap-ar-aging": "balance_review",
    "revenue-quality": "income_review",
    "liquidity-runway": "balance_review",
    "three-statement-analysis": None,          # per-result, see TSA_SECTIONS
    "revenue-quality-business-lines": "income_review",
}

TSA_SECTIONS = {
    "income": "income_review",
    "balance_sheet": "balance_review",
    "cash_flow": "cashflow_review",
    "working_capital": "balance_review",
    "leverage": "balance_review",
}

# The memo's fixed section order, one entry per block of the five-tab report
# (report-structure.md "Section order (fixed)"). The payload carries them in
# this order so the renderer never decides it. The old one-page `snapshot_exposure`
# is gone: the facility it stated is the Deal summary tab's key-terms card, and
# the funded-debt group it carried is that tab's capital structure.
SECTIONS = [
    "header",                # above the tabs
    "covenant_strip",        # overview + covenants
    "snapshot_performance",  # overview
    "snapshot_collateral",   # overview
    "exec_summary",          # overview, prose only
    "trend_chart",           # overview
    "deal_company",          # deal
    "deal_terms",            # deal
    "deal_transaction",      # deal
    "deal_security",         # deal
    "deal_capstructure",     # deal
    "covenant_detail",       # covenants
    "income_review",         # financials
    "cashflow_review",       # financials
    "balance_review",        # financials
    "files_registry",        # files
    "data_quality",          # files
    "flags",                 # files
]

# What each section carries beyond the `figures / chips / groups` seed every
# section has. Seeded up front so the payload's shape is the same on every
# borrower and the renderer never has to ask whether a key exists: an input the
# run did not have leaves the seed value and adds an `absent` record saying why.
# The two module-gated blocks -- `income_review.revenue_quality` and
# `balance_review.aging` -- are deliberately NOT seeded: a module that did not
# run leaves its key off entirely, and the `module_skipped` record covers it.
SECTION_SEEDS = {
    "header": {"our_view": None, "deal_team": [], "borrower_legal_name": None,
               "one_liner": None},
    "covenant_strip": {"prior_breaches": [], "either_or": []},
    "snapshot_performance": {"performance_views": None},
    "snapshot_collateral": {"collateral": None},
    "trend_chart": {"months": [], "series": {}, "notes": {},
                    "basis_break_index": None, "quarters": []},
    "deal_company": {"one_liner": None, "facts": [], "source_ref": None},
    "deal_terms": {"rows": [], "covenant_summary_row": None},
    "deal_transaction": {"elements": {}, "gaps": [], "sources": []},
    "deal_security": {"as_of": None, "columns": [], "rows": []},
    "deal_capstructure": {"rows": [], "funded_debt": None},
    "covenant_detail": {"preamble": None, "grid": [], "history": [],
                        "far_out": [], "stress_inputs": None},
    "income_review": {"table": None},
    "cashflow_review": {"table": None},
    "balance_review": {"table": None, "liquidity": None},
    "files_registry": {"documents": []},
    "data_quality": {"notes": []},
    "flags": {"items": []},
}

# What to call each module in front of a reader. The memo is read by fund analysts,
# so a note about a missing input names the work in their words rather than
# printing the module's own slug.
MODULE_NAMES = {
    "covenant-compliance": "Covenant compliance",
    "budget-vs-actual": "Performance against plan",
    "ap-ar-aging": "Receivables and payables ageing",
    "revenue-quality": "Revenue quality",
    "liquidity-runway": "Liquidity and runway",
    "three-statement-analysis": "Income, balance sheet and cash flow review",
    "revenue-quality-business-lines": "Revenue by business line",
}

# Words that mean something to whoever wrote the engine and nothing to the analyst
# reading the memo. A sentence containing one is kept in the payload for an
# engineer and left off the page.
ENGINE_WORDS = ("--", "engine", "input-gated", "argparse", "cli", "stderr",
                "traceback", "exit code", "config key", "json", "schema")

# A covenant result whose status is one of these carries no verdict, so it is an
# `absent` record rather than a figure.
NOT_TESTED = {"not_evaluable", "not_tested", "n/e", "not_computable"}

# Files that live in the archive folder and follow its naming, but are not module
# outputs: see read_archive.
NOT_A_MODULE = {"coverage", "certified-values"}


class PayloadError(Exception):
    """A bad input the operator can act on -- reported as a plain message."""


def recorded_path(*paths):
    """First supported logical path among configured or resolved candidates."""
    for path in paths:
        if not path:
            continue
        for candidate in (path, ml.unresolve_path(path)):
            try:
                if ml.parse_store_path(candidate):
                    return candidate
            except ml.ConfigError:
                pass
    return None


# ------------------------------------------------------------------ small helpers
def module_name(mod):
    """The module's name as a reader would say it."""
    return MODULE_NAMES.get(str(mod)) or str(mod).replace("-", " ").capitalize()


def plain_reason(text):
    """A module's own explanation, with the engine-facing sentences dropped.

    Modules write a reason for whoever reads the run, and the good part of it is
    plain: which report was not filed, when it last was. The tail is often written
    for an engineer ("the engine needs --ar-latest, so the module is input-gated
    off"), and a fund analyst reading that learns nothing and worries. Keep the
    sentences a reader can act on; the full text stays in the payload as `detail`.
    """
    text = str(text or "").strip()
    if not text:
        return ""
    parts = re.split(r"(?<=[.!?])\s+", text)
    kept = [p for p in parts
            if not any(w in p.lower() for w in ENGINE_WORDS)]
    return " ".join(kept).strip() or parts[0].strip()


def seat(sections, section, what="figure"):
    """The section a figure, group or record says it belongs to.

    A name the payload does not carry is a bug in whoever emitted it, and the one
    thing it must not do is disappear: a figure quietly dropped is a number the
    memo never prints and no gate can miss. So an unknown name stops the run with
    the name it used and the names that exist.
    """
    try:
        return sections[section]
    except KeyError:
        raise PayloadError(
            f"a {what} says it belongs to a memo section called {section!r}, and "
            f"the memo has no such section. Its sections are: "
            f"{', '.join(SECTIONS)}.")


# Keys of a deal-record block that describe where the block CAME FROM rather than
# what it says. A fact row already carries its own source, so folding these into
# the row's text puts the filing note on the page beside the fact.
RECORD_BOOKKEEPING = ("source", "source_tier", "provenance", "as_of", "status")


def flatten_value(v):
    """One deal-record value as one line of reader-facing text, or None.

    The deal record states a fact as a string, as a dict of named parts, or as a
    list of them. A dict joins its own values with "; " and a list joins with
    " * ", so the page shows the whole fact rather than the first part of it.
    Notes to the next reader of the file (an underscore key) and the record's own
    filing details are left out -- they are not the fact.
    """
    if v is None or v == "" or v == [] or v == {}:
        return None
    if isinstance(v, dict):
        parts = [flatten_value(x) for k, x in v.items()
                 if not str(k).startswith("_")
                 and str(k).lower() not in RECORD_BOOKKEEPING]
        parts = [p for p in parts if p]
        return "; ".join(parts) or None
    if isinstance(v, (list, tuple)):
        parts = [flatten_value(x) for x in v]
        parts = [p for p in parts if p]
        return " · ".join(parts) or None
    if isinstance(v, bool):
        return "yes" if v else "no"
    return str(v).strip() or None


def first_value(d, *keys):
    """The first of several field names a record might use, or None.

    The same fact is spelled differently across the deals -- amortisation and
    amortization, `ecf_sweep` on the facility and again under its fees -- and a
    memo that reads only one spelling prints "not on file" against a fact that is
    on file.
    """
    for k in keys:
        if not isinstance(d, dict):
            return None
        if d.get(k) not in (None, "", [], {}):
            return d.get(k)
    return None


_PCT = re.compile(r"(\d+(?:\.\d+)?)\s*%")
_NO_FIGURE = ("none", "not ", "no ", "n/a", "nil")


def leading_pct(text):
    """A percentage a deal record states at the front of its own sentence.

    The discount at which the loan was bought is written as prose ("2.00% -- each
    advance purchased at 98.00% of principal"), and the security summary's cost
    column is that first figure. Only the OPENING of the sentence is read, and
    only when the sentence does not begin by saying there is none -- a percentage
    picked out of the middle of a paragraph is as likely to be the purchase price
    or a fee as the discount.
    """
    s = str(text or "").strip()
    if not s or any(s.lower().startswith(w) for w in _NO_FIGURE):
        return None
    m = _PCT.search(s[:12])
    return float(m.group(1)) / 100.0 if m else None


def _us_date(iso):
    """An ISO date as the memo writes a date: 2026-06-30 -> 6/30/2026."""
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", str(iso or ""))
    if not m:
        return None
    return f"{int(m.group(2))}/{int(m.group(3))}/{m.group(1)}"


def _leading_iso_date(text):
    """The ISO date a spec field opens with, or None.

    A covenant's first test date is often written with the reason beside it
    ("2024-07-19 (Closing Date; covenant applies at all times ...)"), so the date
    is read off the front and the rest is left as the note it is.
    """
    m = re.match(r"\s*(\d{4}-\d{2}-\d{2})", str(text or ""))
    return m.group(1) if m else None


def _iso_plus_months(iso, months):
    """An ISO date shifted by whole months, clamped to the month's own length."""
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", str(iso or ""))
    if not m:
        return None
    import calendar
    y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    n = y * 12 + mo - 1 + months
    y2, mo2 = n // 12, n % 12 + 1
    return f"{y2:04d}-{mo2:02d}-{min(d, calendar.monthrange(y2, mo2)[1]):02d}"


def _month_offset(period, months):
    """An ISO YYYY-MM shifted by whole months, or None."""
    if not re.fullmatch(r"\d{4}-\d{2}", str(period or "")):
        return None
    n = int(period[:4]) * 12 + int(period[5:7]) - 1 + months
    if n < 0:
        return None
    return f"{n // 12:04d}-{n % 12 + 1:02d}"


def read_pledged(model, spec, period, to_usd=None):
    """One pledged component's reported value, whether it is a row or a sum of rows.

    A map names a component either by a single model row ("Total - 13000 -
    Inventory") or by rows to add and subtract, which is how a net figure like
    fixed assets after accumulated depreciation is expressed. Returns
    (value, refs, missing_rows) -- and a value only when every named row resolved,
    because a net figure short one of its parts is not that figure. Values come
    back in dollars.
    """
    to_usd = to_usd or (lambda v: v)
    if model is None or not spec:
        return None, None, []
    if isinstance(spec, str):
        v = to_usd(_num(model.get_line(spec, month=period, sheet="Balance Sheet")))
        ref = model.provenance(spec, sheet="Balance Sheet", month=period)
        return v, ([ref] if ref else None), ([] if v is not None else [spec])
    if not isinstance(spec, dict):
        return None, None, []
    total, refs, missing = 0.0, [], []
    for sign, key in ((1, "add"), (-1, "subtract")):
        for row in (spec.get(key) or []):
            v = to_usd(_num(model.get_line(row, month=period, sheet="Balance Sheet")))
            if v is None:
                missing.append(row)
                continue
            total += sign * v
            ref = model.provenance(row, sheet="Balance Sheet", month=period)
            if ref:
                refs.append(ref)
    if missing or not refs:
        return None, (refs or None), missing
    return total, refs, []


def _rows_of(spec):
    """Every model row a component names, for saying which one is missing."""
    if isinstance(spec, str):
        return [spec]
    if isinstance(spec, dict):
        return [r for k in ("add", "subtract") for r in (spec.get(k) or [])]
    return []


def _settled(v):
    """A control-file value, or None when it is still a note-to-self.

    Deal context is drafted before it is signed off, and an unsettled field says so
    in its own text ("AGENT PLACEHOLDER - watchlist / trajectory"). Printing that
    into a credit memo states a view the fund has not given, so an unsettled
    value is treated as no value.
    """
    if not isinstance(v, str):
        return v
    flag = ("placeholder", "tbd", "tbu", "to be updated", "fill in", "xxx",
            "no analyst", "not a credit opinion", "must be replaced",
            "auto-set", "provisional")
    return None if any(f in v.lower() for f in flag) else v


# The three fields of the deal record that state the fund's own view of the credit.
JUDGMENT_FIELDS = ("watchlist", "trajectory", "thesis_status")

# Fields retired from the judgment record. An older deal-context can still carry
# these -- the field and its own disclaimer note -- from before they were retired.
# Both are dropped before sign-off classification runs: left in, an old retired
# note that reads as unsettled (e.g. "provisional; no analyst sign-off") is not
# recognised as belonging to any current field, so it is read as a WHOLE-BLOCK
# warning and silently suppresses every field the analyst actually did sign off.
RETIRED_JUDGMENT_FIELDS = ("risk_rating",)


def _drop_retired(ov):
    """A deal record's `our_view` with retired fields and their notes removed."""
    retired_notes = {f"_{f}_note" for f in RETIRED_JUDGMENT_FIELDS}
    return {k: v for k, v in ov.items()
            if k not in RETIRED_JUDGMENT_FIELDS and k not in retired_notes}


def settled_view(ov):
    """The fund's standing view, keeping only the fields that have been signed off.

    Onboarding writes this block before an analyst has seen it, and it says so in a
    note BESIDE the fields -- `_WARNING` for the whole block, `_trajectory_note` for
    one of them -- rather than inside every value. So the long fields carried the
    marker in their own text and were dropped, while the one-word fields ("on",
    "stable") printed on page one as the fund's view of the credit. A note that
    disclaims a field disclaims it whether the disclaimer sits in the value or
    beside it: a block-level warning governs all three, a `_<field>_note` governs
    that field.

    Returns the view and the names of the fields a disclaimer removed, so the memo
    can say why it shows no view instead of simply showing none.
    """
    ov = _drop_retired(ov or {})
    notes = {k: v for k, v in ov.items() if str(k).startswith("_")}
    whole_block = any(isinstance(v, str) and _settled(v) is None
                      for k, v in notes.items() if not _names_a_field(k))
    out, disclaimed = {}, []
    for k in JUDGMENT_FIELDS:
        note = notes.get(f"_{k}_note")
        if whole_block or (isinstance(note, str) and _settled(note) is None):
            out[k] = None
            if ov.get(k) not in (None, ""):
                disclaimed.append(k)
            continue
        out[k] = _settled(ov.get(k))
        if out[k] is None and ov.get(k) not in (None, ""):
            disclaimed.append(k)
    return out, disclaimed


def _names_a_field(key):
    """Whether a note key belongs to one judgment field rather than the block."""
    return any(str(key) == f"_{f}_note" for f in JUDGMENT_FIELDS)


# What the deal record's `static.deal_team_provenance` says when the analyst
# confirmed the names. A record that carries the stamp and does not say this holds
# a list nobody looked at, so the header states no deal team. The officer-title
# rule is `ml.names_an_officer`, shared with the setup engine that proposes names.
ANALYST_STAMP = "analyst"


def settled_deal_team(static):
    """The deal-team names the header may print, and why any were removed.

    Onboarding proposes these names and the analyst confirms them, which stamps the
    list. Two rules keep a borrower's officer off the banner: a list the deal record
    itself says nobody confirmed is dropped whole, and an entry carrying an officer
    title is dropped whatever the stamp says. A record written before the stamp
    existed keeps its names -- the stamp governs when it is there.

    Returns the names and one sentence for the run's `absent` record, or None when
    every name on file survives.
    """
    static = static or {}
    raw = static.get("deal_team") or []
    if not isinstance(raw, list):
        raw = str(raw).split(",")
    names = [str(n).strip() for n in raw if str(n).strip()]
    if not names:
        return [], None

    stamp = str(static.get("deal_team_provenance") or "").strip()
    if stamp and ANALYST_STAMP not in stamp.lower():
        return [], ("the deal record says the analyst did not confirm these names, "
                    "so the header states no deal team.")

    kept = [n for n in names
            if _settled(n) is not None and not ml.names_an_officer(n)]
    if len(kept) < len(names):
        return kept, ("a name on the deal-team list carries an officer title, so it "
                      "reads as an officer of the borrower rather than one of ours "
                      "and is left off the header.")
    return kept, None


def _slug(s):
    s = re.sub(r"[^a-z0-9]+", "-", str(s or "").lower()).strip("-")
    return s or "x"


def _period_label(period):
    """"2026-05" as the canonical month label monitor_lib's period maths uses."""
    return ml._label_from(int(str(period)[:4]), int(str(period)[5:7]))


def _num(x):
    """A figure as a float, or None. Accepts the string forms the archives carry."""
    return ml.safe_float(x)


def _read_json(path):
    with open(path, "r", encoding="utf-8") as fh:
        return json.load(fh)


def _maybe_json(path):
    """Read a JSON file, or None when it is absent or unreadable. Used for the
    optional control files -- a missing collateral map is a memo without a
    collateral box, not a failed run."""
    if not path:
        return None
    try:
        return _read_json(ml.resolve_path(path))
    except (OSError, ValueError):
        return None


def _url(store, path_in_store, stores):
    """The store URL for one ref, joined from the ENVELOPE's own store table.

    Returns (url, degrade_reason). A missing store or a store with no `url_base`
    yields (None, reason) -- the caller emits a figure that still answers, with
    no link and a note, and never a guessed URL.

    The address comes from `monitor_lib.store_link`, so a memo link opens the
    file the way a dashboard link does. Joining the path here instead put the
    workbook in the reader's Downloads folder.
    """
    if not store:
        return None, "no_store"
    entry = (stores or {}).get(store)
    if not isinstance(entry, dict) or not entry.get("url_base"):
        return None, "no_url_base"
    if not path_in_store:
        return entry["url_base"].rstrip("/"), None
    return ml.store_link(entry["url_base"],
                         str(path_in_store).replace("\\", "/")), None


# ------------------------------------------------------------------ the archive
_PERIOD_FILE = re.compile(r"^(?P<module>.+)-(?P<period>\d{4}-\d{2})"
                          r"(?P<skip>\.skipped)?\.json$")


def read_archive(archive_dir, period):
    """Every archived MODULE output for `period`, keyed by module, in sorted order.

    Three things sitting in the same folder are not module outputs and carry no
    memo figures: the sidecars an engine writes beside its output
    (`*.json.rows.json`), the coverage record (a record OF the run), and the
    certified values transcribed from the borrower's compliance certificate (an
    INPUT the covenant engine reads, whose figures reach the memo through that
    engine's own results). Listing any of them would make the memo's identity
    block claim a result file that produced nothing.
    """
    if not os.path.isdir(archive_dir):
        raise PayloadError(
            f"No results folder to build the memo from. Looked in: {archive_dir}")
    found = {}
    for name in sorted(os.listdir(archive_dir)):
        if name.endswith(".rows.json"):
            continue
        m = _PERIOD_FILE.match(name)
        if not m or m.group("period") != period:
            continue
        module = m.group("module")
        path = os.path.join(archive_dir, name)
        try:
            data = _read_json(path)
        except ValueError as e:
            raise PayloadError(f"{name} is not readable as a result file: {e}")
        if module in NOT_A_MODULE:
            continue
        found[module] = {"path": path, "name": name, "data": data,
                         "skipped": bool(m.group("skip")) or bool(data.get("skipped"))}
    if not found:
        raise PayloadError(
            f"No results for {period} in {archive_dir} — run the monthly monitor "
            f"for this month before rendering its memo.")
    return found


# ------------------------------------------------------------------ figures
def _figure(fid, section, label, value, unit, ref, calc, env, extra=None):
    """One memo figure: its value in dollars where it is money, the trace it came
    with, and the store link joined once here.

    Returns (figure, degrade_reasons). A figure the payload cannot state safely
    comes back with `unsafe` set and no value -- see the unit rule below.
    """
    stores = env.get("stores") or {}
    scale = float(((env.get("reporting_unit") or {}).get("scale")) or 1.0)
    v = _num(value)
    kind = (unit or "").strip().lower()
    unsafe = None
    if v is not None and kind in SCALES and scale != 1.0:
        v = v * scale
    elif v is not None and not kind and scale != 1.0:
        # The result does not say what kind of quantity it is, and this borrower's
        # model is not in dollars. Multiplying would claim a 1000x ratio; not
        # multiplying would understate a money figure a thousand-fold, which reads as
        # a collapsing business. Neither is publishable, so the figure is withheld and
        # said out loud. The fix belongs in the module that omitted the unit.
        unsafe = (f"this figure does not say whether it is money, and this borrower's "
                  f"model is kept in units of {scale:,.0f} dollars, so it cannot be "
                  f"stated safely")
        v = None
    hops, degrade = [], []
    for r in (ref or []):
        if not isinstance(r, dict):
            continue
        url, why = _url(r.get("store"), r.get("path_in_store"), stores)
        if why:
            degrade.append(why)
        ups = []
        for u in (r.get("upstream") or []):
            if not isinstance(u, dict):
                continue
            uurl, uwhy = _url(u.get("store"), u.get("path_in_store"), stores)
            if uwhy:
                degrade.append(uwhy)
            ups.append({k: u.get(k) for k in
                        ("store", "path_in_store", "sheet", "section", "quote")
                        if u.get(k) is not None} | {"url": uurl})
        if not r.get("upstream"):
            degrade.append("no_upstream")
        hops.append({k: r.get(k) for k in
                     ("store", "path_in_store", "sheet", "row", "label",
                      "cells", "section", "period")
                     if r.get(k) is not None} | {"url": url, "upstream": ups})
    out = {
        "id": fid,
        "section": section,
        "label": label,
        "value": v,
        "unit": unit,
        "scaled_by": scale if (kind in SCALES and scale != 1.0) else 1.0,
        "calc": calc,
        "hops": hops,
        "source_ref": ref or [],
    }
    if unsafe:
        out["unsafe"] = unsafe
        degrade.append("unit_unstated")
    if extra:
        out.update(extra)
    return out, sorted(set(degrade))


def is_analytical(entry):
    """Whether this row is an analytical reading rather than a covenant of the loan.

    A spec may declare it outright with `analytical: true`, which is the preferred
    form. Older specs may instead open the row's name with "ANALYTICAL" and its note
    with "ANALYTICAL ROW, NOT A COVENANT IN THE AGREEMENT", so the name is read too
    -- otherwise a reading the agreement does not contain sits on the front of the
    memo wearing a breach, and a reader counts one covenant too many.
    """
    if entry.get("analytical") is not None:
        return bool(entry.get("analytical"))
    name = str(entry.get("name") or "").strip().lower()
    note = str(entry.get("note") or "").strip().lower()
    return name.startswith("analytical") or note.startswith("analytical row")


# The two things a covenant's spec name carries that the memo states elsewhere:
# the "either/or leg" tag, which the group chip above the leg already says, and a
# trailing parenthetical restating the formula or the covenant's other title.
_LEG_TAG = re.compile(r"(\s*[-–—]{1,2}\s*either\s*/\s*or\s+leg\s*$"
                      r"|\s*\(\s*either\s*/\s*or\s+leg\s*\)\s*$)", re.I)
_TRAILING_PAREN = re.compile(r"\s*\([^()]*\)\s*$")


def covenant_short_name(name):
    """A covenant's name the way a chip and a one-line summary say it.

    The full spec name stays on the detail grid, where a reader is reading one
    covenant closely. On a chip it has to be read at a glance, so the leg tag and
    the parenthetical come off and the ASCII double hyphen becomes an em dash.
    """
    s = str(name or "").strip()
    for _ in range(2):
        s = _LEG_TAG.sub("", s).strip()
    s = _TRAILING_PAREN.sub("", s).strip() or s
    return re.sub(r"\s*--\s*", " — ", s).strip()


def covenant_short_names(names):
    """Short names for a set of covenants, keyed by full name.

    Shortening drops a parenthetical, and two covenants can be told apart by that
    parenthetical alone ("Minimum EBITDA (senior)" beside "(total)"). Where the
    short forms collide, every covenant in the collision keeps its full name: a
    long chip is read slowly, a wrong one is read wrongly.
    """
    shorts = {str(n): covenant_short_name(n) for n in names}
    counts = {}
    for s in shorts.values():
        counts[s.lower()] = counts.get(s.lower(), 0) + 1
    return {n: (s if counts[s.lower()] == 1 else n) for n, s in shorts.items()}


def _section_token(section):
    """The section NUMBER at the front of a recorded section reference.

    A spec may record the provenance with the number ("6.8(b), as amended and
    restated by Third Amendment Section 8 (effective 2026-05-07)"). That sentence
    is a trace, not a chip label, so the label keeps "6.8(b)" and the sentence
    stays where a reader is reading closely. A reference that does not start with a
    number is used whole.
    """
    s = re.sub(r"^\s*(?:§|section)\s*", "", str(section or ""), flags=re.I).strip()
    m = re.match(r"\d+(?:\.\d+)*(?:\([^)]{1,8}\))*", s)
    return (m.group(0) if m else s) or None


def _shared_section(sections):
    """The section the legs of a group have in common: 6.7(b)(i) + 6.7(b)(ii) ->
    6.7(b). Used where a group records no section of its own."""
    toks = [t for t in (_section_token(s) for s in sections) if t]
    if not toks:
        return None
    common = toks[0]
    for t in toks[1:]:
        while common and not t.startswith(common):
            common = common[:-1]
    # Stop at a clean boundary, so a trimmed "6.7(" or "6." never reaches a chip.
    common = re.sub(r"[^\d)]+$", "", common or "")
    return common if re.fullmatch(r"\d+(?:\.\d+)*(?:\([^)]{1,8}\))*", common or "") \
        else None


def _is_identifier(text):
    """True where a string is a machine key rather than something to read:
    "s6_7_b", "grp-2". A name a person wrote has spaces in it."""
    s = str(text or "").strip()
    return bool(s) and not re.search(r"\s", s) and bool(re.search(r"[_-]|\d", s))


def either_or_labels(env=None, spec=None):
    """How the memo names each either/or group, keyed by group id.

    A group id is a key in the spec ("s6_7_b"), so printing it puts a slug in the
    largest text on the front of the memo. The reader holds the agreement, so the
    memo prints that group's section of it -- "§6.7(b)" -- and keeps the id for
    the trace.

    Three sources, in order, because a chip must never fall back to the id: the
    group's own recorded section; the section its legs share, where the group
    records none of its own; then a name a person wrote. A group whose spec offers
    none of the three is named for what it is, since "Either/or test" beside its
    own legs is readable and `s6_7_b` is not.
    """
    sources = [g for g in ((env or {}).get("results") or [])
               if isinstance(g, dict) and str(g.get("kind") or "") == "either_or_group"]
    for gid, g in (((spec or {}).get("either_or_groups") or {}).items()):
        if isinstance(g, dict):
            sources.append(dict(g, id=gid))
    legs = {}
    for c in ((spec or {}).get("covenants") or []):
        if isinstance(c, dict) and c.get("either_or_group"):
            ref = c.get("agreement_ref")
            if isinstance(ref, dict):
                legs.setdefault(str(c["either_or_group"]), []).append(
                    flatten_value(ref.get("section")))
    labels = {}
    for g in sources:
        gid = str(g.get("id") or "")
        if not gid or gid in labels:
            continue
        ref = g.get("agreement_ref")
        section = _section_token(flatten_value(ref.get("section"))) \
            if isinstance(ref, dict) else None
        section = section or _shared_section(legs.get(gid) or [])
        name = flatten_value(g.get("name"))
        labels[gid] = (f"§{section}" if section
                       else (name if name and not _is_identifier(name)
                             else "Either/or test"))
    return labels


def _covenant_figures(mod, env, period):
    """Covenant results become figures on the detail grid, chips on the strip and
    a per-covenant history; a test with no verdict becomes an `absent` record.

    The engine emits ONE RESULT PER TEST DATE, not one per covenant -- that series
    is the memo's covenant-history matrix. So each figure's id carries its test
    date: keying on the covenant alone silently served the last test date's number
    as the current month's, which is the class of defect this payload exists to
    make impossible. The current test is the one falling on the reported month's
    end; that is the only one the strip and the grid seat.
    """
    figures, absent, degrade = [], [], []
    chips, history, groups = [], {}, {}
    current_end = ml.month_end_iso(_period_label(period)) if period else None
    scale = float(((env.get("reporting_unit") or {}).get("scale")) or 1.0)
    labels = either_or_labels(env)

    def money(v, primitive):
        """A covenant's own figure in dollars, where the covenant is a balance.

        The strip and the history print these values directly, so a level covenant
        left in the model's own units puts "$750" beside a $750,000 floor on the
        front of the memo. A ratio covenant is a multiple and never scales.
        """
        n = _num(v)
        if n is None or primitive in ("max_ratio", "min_coverage"):
            return n
        return n * scale

    for i, r in enumerate(env.get("results") or []):
        if not isinstance(r, dict):
            continue
        cid = r.get("id") or f"covenant-{i}"
        td = str(r.get("test_date") or "")
        is_current = bool(current_end) and td == current_end
        status = str(r.get("status") or "").strip().lower()
        analytical = is_analytical(r)
        if str(r.get("kind") or "") == "either_or_group":
            # An either/or group is not a covenant: it holds no figure of its own,
            # and its legs carry the history. It is the strip's governing verdict,
            # settled once by the engine so the strip and the detail cannot differ.
            # Read as a covenant it had no `actual`, so the memo published the
            # group as a test that could not be measured in any month.
            if is_current:
                groups[str(cid)] = {"id": str(cid),
                                    "label": labels.get(str(cid)) or str(cid),
                                    "name": r.get("name"),
                                    "status": r.get("status"),
                                    "headroom_pct": _num(r.get("headroom_pct")),
                                    "members": [str(m) for m in (r.get("members") or [])]}
            continue
        hist = history.setdefault(cid, {"id": cid, "name": r.get("name") or cid,
                                        "primitive": r.get("primitive"),
                                        "analytical": analytical,
                                        "tests": []})
        if status in NOT_TESTED or r.get("actual") is None:
            reason = (r.get("not_evaluable_reason") or r.get("reason")
                      or (r.get("notes") if isinstance(r.get("notes"), str) else None)
                      or "the figures this test needs are not in the model for this month")
            hist["tests"].append({"test_date": td, "status": "not_evaluable",
                                  "reason": reason})
            absent.append({
                "kind": "covenant_not_tested",
                "section": "covenant_detail",
                "what": f"{r.get('name') or cid} at {td}" if td else (r.get("name") or cid),
                "id": cid, "test_date": td, "current": is_current,
                "reason": reason, "module": mod,
            })
            continue
        stem = f"covenant.{_slug(cid)}.{_slug(td) or f'test-{i}'}"
        for key in ("actual", "required", "headroom_abs", "headroom_pct",
                    "cert_value", "cert_headroom_abs", "cert_headroom_pct"):
            if r.get(key) is None:
                continue
            if key.endswith("_pct"):
                # Headroom is a proportion the memo prints out of a hundred
                # ("+25.8% of room"), so the payload says percentage rather than
                # ratio -- otherwise the gate compares a printed 25.8 against a
                # stored 0.258 and cannot tell a real figure from an edited one.
                u = "%"
            elif r.get("primitive") in ("max_ratio", "min_coverage"):
                u = "ratio"
            else:
                u = "usd"
            f, dg = _figure(
                f"{stem}.{key}", "covenant_detail",
                f"{r.get('name') or cid} — {key.replace('_', ' ')}"
                + (f" ({td})" if td else ""),
                r.get(key), u, r.get("source_ref"), r.get("calc"), env,
                extra={"covenant_id": cid, "field": key, "test_date": td,
                       "current": is_current, "status": r.get("status"),
                       "cert_status": r.get("cert_status"),
                       "cert_agrees": r.get("cert_agrees"),
                       "cert_delta": _num(r.get("delta_model_vs_cert")),
                       "analytical": analytical,
                       "window": r.get("window"),
                       "primitive": r.get("primitive")})
            figures.append(f)
            degrade.extend(dg)
        primitive = r.get("primitive")
        hist["tests"].append({
            "test_date": td, "status": r.get("status"),
            "cert_status": r.get("cert_status"),
            "actual": money(r.get("actual"), primitive),
            "required": money(r.get("required"), primitive),
        })
        if is_current:
            chips.append({
                "id": cid, "name": r.get("name") or cid,
                "status": r.get("status"), "cert_status": r.get("cert_status"),
                "test_date": td,
                "headroom_pct": _num(r.get("headroom_pct")),
                "either_or_group": r.get("either_or_group"),
                "primitive": primitive,
                "analytical": analytical,
                # Both sides on the front of the memo: the reader who sees only the
                # strip on a credit whose certificate and our own figures disagree
                # has to see that they disagree.
                "required": money(r.get("required"), primitive),
                "actual": money(r.get("actual"), primitive),
                "cert_value": money(r.get("cert_value"), primitive),
                "cert_agrees": r.get("cert_agrees"),
                "cert_delta": money(r.get("delta_model_vs_cert"), primitive),
            })
    shorts = covenant_short_names([c["name"] for c in chips])
    for c in chips:
        c["short_name"] = shorts.get(str(c["name"])) or c["name"]
    for h in history.values():
        h["tests"].sort(key=lambda t: t.get("test_date") or "")
        h["tally"] = {
            k: sum(1 for t in h["tests"] if str(t.get("status") or "").lower() == k)
            for k in ("pass", "at_risk", "breach", "not_evaluable")
        }
    if not chips and current_end:
        absent.append({
            "kind": "covenant_no_test_this_month",
            "section": "covenant_strip",
            "what": "covenant status",
            "reason": f"no covenant falls due on {current_end}, so the strip shows "
                      f"the history instead of this month's verdicts",
            "module": mod,
        })
    return (figures, absent, degrade, chips,
            [history[k] for k in sorted(history)],
            [groups[k] for k in sorted(groups)])


def unresolved_earlier_breaches(history, period):
    """Breaches at test dates before the reported month, for the front of the memo.

    The strip carries the month's own verdicts, so a month that passes shows green
    and nothing else -- and a reader who takes in only the front of a memo on a
    credit that breached in April reads it as compliant. A covenant breach is an
    event until it is waived or cured, and neither a waiver nor a cure is anything
    this run reads, so the front of the memo states the earlier breach and sends the
    reader to the history rather than deciding it is behind us.

    One entry per covenant, the most recent breach, within the trailing year.
    """
    current_end = ml.month_end_iso(_period_label(period)) if period else None
    if not current_end:
        return []
    start = _month_offset(period, -12)
    floor = f"{start}-01" if start else ""
    out = []
    for h in history:
        if h.get("analytical"):
            continue
        earlier = [t for t in (h.get("tests") or [])
                   if str(t.get("status") or "").lower() == "breach"
                   and floor <= str(t.get("test_date") or "") < current_end]
        if not earlier:
            continue
        last = max(earlier, key=lambda t: t.get("test_date") or "")
        out.append({"id": h.get("id"), "name": h.get("name") or h.get("id"),
                    "test_date": last.get("test_date"), "status": "breach",
                    "primitive": h.get("primitive"),
                    "actual": last.get("actual"), "required": last.get("required")})
    return sorted(out, key=lambda x: (x.get("test_date") or "", x.get("name") or ""))


# The liquidity block's fields, the name each carries on the page, and its unit.
# `three-statement-analysis` publishes this OUTSIDE `results`, so nothing that
# walks results ever saw it -- and the memo's runway figure was written by hand
# off a five-month average, matching neither archived figure.
LIQUIDITY_FIELDS = [
    ("cash", "Cash", "usd"),
    ("cash_floor", "Covenant cash floor", "usd"),
    ("burn_1m", "Cash burn, last month", "usd"),
    ("burn_3m_avg", "Cash burn, three-month average", "usd"),
    ("months_to_floor_1m", "Months to the floor, on last month's burn", "months"),
    ("months_to_floor_3m", "Months to the floor, on the three-month burn", "months"),
]


def liquidity_figures(env):
    """The archived runway, published so the memo quotes it rather than redoing it.

    The module already works out how long the cash lasts, on two burn bases, and
    archives both with the rows they were read from. Publishing them as figures is
    what lets the page print them, the gate check them, and the prose quote them --
    a runway a memo computes for itself is a third number nobody can check.
    """
    liq = env.get("liquidity")
    if not isinstance(liq, dict):
        return [], [], [], None
    figs, degrade = [], []
    ref, calc = liq.get("source_ref"), liq.get("calc")
    for key, label, unit in LIQUIDITY_FIELDS:
        if liq.get(key) is None:
            continue
        f, dg = _figure(f"liquidity.{_slug(key)}", "balance_review", label,
                        liq.get(key), unit, ref, calc, env,
                        extra={"metric": label, "field": key})
        figs.append(f)
        degrade.extend(dg)
    return figs, sorted(set(degrade)), _runway_gaps(liq), liquidity_block(liq, figs)


def liquidity_block(liq, figs):
    """The runway block's own rows, in the order the memo prints them.

    The figures carry the values and their traces; this says which of them the
    block shows and under what heading, so the renderer lays the block out
    without deciding what belongs in it.

    Each row's value is READ OFF ITS FIGURE rather than off the module again, so a
    money row is in dollars like every other money figure in the payload. Reading
    the module a second time published this block in the model's own units -- cash
    of 6,662.43 on a borrower whose statements are kept in thousands, which prints
    as $6,662 against a $750,000 covenant floor.
    """
    have = {f.get("field"): f for f in figs}
    rows = [{"id": _slug(key), "field": key, "label": label, "unit": unit,
             "value": have.get(key, {}).get("value"),
             "figure_id": have.get(key, {}).get("id")}
            for key, label, unit in LIQUIDITY_FIELDS
            if liq.get(key) is not None]
    if not rows:
        return None
    return {"rows": rows, "calc": liq.get("calc"),
            "source_ref": liq.get("source_ref") or []}


def _runway_gaps(liq):
    """Why a burn basis shows no runway, when it shows none.

    A basis on which cash GREW has no months to a floor, and dropping the row
    without a word leaves a reader wondering whether it failed to compute. It also
    settles which basis the liquidity lens reads: the rule turns on the bases that
    have a figure.
    """
    out = []
    for burn, months, said in (("burn_1m", "months_to_floor_1m", "last month's burn"),
                               ("burn_3m_avg", "months_to_floor_3m",
                                "the three-month burn")):
        b = _num(liq.get(burn))
        if liq.get(months) is not None or b is None:
            continue
        out.append({
            "kind": "runway_not_measured", "section": "balance_review",
            "what": f"Months to the floor, on {said}",
            "reason": ("cash grew on this basis, so it never reaches the floor"
                       if b <= 0 else
                       "cash is already at or below the floor on this basis")})
    return out


def absent_reads(mod, env, section_of):
    """Every figure a module reported as absent, as a line the memo can print.

    A module now says when the row behind a figure holds nothing rather than
    summing the blanks: `absent` with the reason and one plain sentence
    (coverage_note.absent_fields). Those figures reach the memo as "could not
    cover" lines instead of vanishing — a revenue row that is empty in all 23
    months used to leave a memo full of zeros, and a memo missing a line reads
    the same as a memo whose line was fine.

    The sweep is depth-first because a module may report absence beside a figure
    rather than instead of a whole record: the ageing module's snapshot is real
    and only its days metric has nothing behind it.
    """
    out = []

    def walk(node, depth=0):
        if depth > 12:
            return
        if isinstance(node, dict):
            if node.get("absent") is True:
                what = (node.get("absent_what") or node.get("metric")
                        or node.get("name") or node.get("label") or node.get("line")
                        or node.get("id") or module_name(mod))
                out.append({
                    "kind": "figure_absent",
                    "section": section_of(node),
                    "what": str(what).replace("_", " "),
                    "reason": node.get("absent_note")
                              or "the model holds no figure behind it",
                    "detail": node.get("absent_reason"),
                    "module": mod,
                })
                return
            for k, v in node.items():
                if k != "source_ref":
                    walk(v, depth + 1)
        elif isinstance(node, list):
            for v in node:
                walk(v, depth + 1)

    walk(env.get("results") or [])
    for key in (env.get("figure_collections") or []):
        walk(env.get(key))
    return out


def _result_figures(mod, env, section_of):
    """Every other module's results, one figure per numeric field a memo prints."""
    figures, degrade = [], []
    for i, r in enumerate(env.get("results") or []):
        if not isinstance(r, dict):
            continue
        if r.get("absent") is True:
            # Nothing behind it, so nothing to print: absent_reads has already
            # turned it into a line saying why.
            continue
        section = section_of(r)
        name = (r.get("metric") or r.get("name") or r.get("id")
                or r.get("label") or f"{mod}-{i}")
        unit = r.get("unit")
        base = f"{_slug(mod)}.{_slug(name)}"
        if r.get("scope"):
            base += f".{_slug(r['scope'])}"
        emitted = False
        for key in ("latest", "actual", "value", "ytd", "budget", "prior",
                    "prior_ytd", "var_abs", "var_pct", "mom_pct"):
            if key not in r or r.get(key) is None:
                continue
            # Some modules publish a whole snapshot under one of these names -- the
            # ageing module's `latest` is a dict of buckets, a total and a days
            # metric. That is a block the memo lays out, not a figure, and emitting
            # it as one produced a figure with no value, no source and no
            # calculation: a row the page cannot print and the gate cannot check.
            if not isinstance(r.get(key), (int, float)) and _num(r.get(key)) is None:
                continue
            u = unit
            if key in ("var_pct", "mom_pct"):
                u = "ratio"
            f, dg = _figure(f"{base}.{key}", section,
                            f"{name} — {key.replace('_', ' ')}",
                            r.get(key), u, r.get("source_ref"), r.get("calc"), env,
                            extra={"metric": name, "field": key,
                                   "scope": r.get("scope"),
                                   "material": r.get("material"),
                                   "favorable": r.get("favorable")})
            figures.append(f)
            degrade.extend(dg)
            emitted = True
        if not emitted:
            continue
    return figures, degrade


# ------------------------------------------------------------------ the chart
def basis_break_seat(cfg, months):
    """Where on the chart's own axis the reporting basis changed, and why.

    The chart draws a dashed rule at a basis break so a reader does not compare
    across it, and the break's month is stated in the borrower's settings
    ("Jan 2025", or an ISO month) rather than as a position. Returns (index,
    label, reason): the index of the EARLIEST break that falls inside the months
    the chart draws, or (None, None, None) when the deal has no break or every
    break predates the window. A break before the window needs no rule -- every
    month drawn is on the same basis.
    """
    breaks = ((cfg.get("analysis") or {}).get("basis_breaks") or [])
    seats = []
    for b in breaks:
        raw = b.get("effective") if isinstance(b, dict) else b
        label = ml.month_label(raw)
        if label and label in months:
            seats.append((months.index(label), label,
                          (b.get("reason") if isinstance(b, dict) else None)))
    if not seats:
        return None, None, None
    return min(seats, key=lambda s: s[0])


def trend_notes(unit, break_label, quarters=None):
    """The caption under each of the chart's two views.

    Money on this chart is dollars whatever the borrower's statements are kept
    in, so the dollar view says which -- a reader who knows the package is in
    thousands needs to be told the chart is not. The margin view carries no unit
    at all. Both name the basis break when the window spans one, because that is
    the one place on the chart where two adjacent months are not comparable.
    Both also say what the quarter-end callouts hold, and which arithmetic made
    them: a reader who adds three months of margin and gets a fourth number needs
    to be told the callout did not.
    """
    scaled = unit is not None and unit.scale != 1.0
    abs_note = ("Net revenue, gross profit, EBITDA and operating expense in US "
                "dollars, read from the monthly financial statements"
                + (f", which are kept in {unit.plain}." if scaled else "."))
    pct_note = ("Each line is that figure as a share of net revenue in the same "
                "month, so no dollar amount is shown.")
    if quarters:
        abs_note += (" The box above each quarter-end month holds that quarter's "
                     "three months added together.")
        pct_note += (" The box above each quarter-end month holds the quarter's "
                     "own margin — the three months' figures added up, then "
                     "divided — not the average of the three monthly margins.")
    if break_label:
        spans = (f" The dashed line at {break_label} marks a change in how the "
                 f"borrower reports these figures: months either side of it are "
                 f"not directly comparable.")
        abs_note += spans
        pct_note += spans
        if any(q.get("spans_basis_break") for q in (quarters or [])):
            straddles = (" A quarter box marked † adds months from either side of "
                         "that line together.")
            abs_note += straddles
            pct_note += straddles
    return {"abs": abs_note, "pct": pct_note}


def trend_series(model, cfg, period, months_back, unit=None):
    """The trend chart's bars and lines, read from THIS borrower's model.

    The template keeps the chart's code while the payload supplies its data. The
    series are computed here from the model's own rows so nothing the example
    file carries can reach the page.

    Money is stated in dollars: a model kept in thousands reads 9,750 for
    $9,750,000, and a chart drawn off the raw rows would show a borrower a
    thousandth of its size. Margins are ratios and are left alone.
    """
    to_usd = (lambda v: v) if unit is None else unit.to_dollars
    lines = cfg.get("model_lines") or {}
    groups = cfg.get("opex_groups") or {}
    opex_labels = [v for k, v in sorted(groups.items())
                   if not k.startswith("_") and v]
    want = {
        "net_revenue": lines.get("revenue"),
        "gross_profit": lines.get("gross_profit"),
        "ebitda": lines.get("ebitda_bridge") or lines.get("reported_ebitda"),
    }
    axis = list(model.months or [])
    if period:
        try:
            end = ml._label_from(int(period[:4]), int(period[5:7]))
            if end in axis:
                axis = axis[:axis.index(end) + 1]
        except (ValueError, IndexError):
            pass
    axis = axis[-months_back:]
    out = {"months": axis, "series": {}, "absent": []}
    for key, label in sorted(want.items()):
        if not label:
            out["absent"].append({"series": key,
                                  "reason": "this borrower's account list names no row for it"})
            continue
        vals = []
        for mo in axis:
            vals.append(to_usd(_num(model.get_line(label, month=mo))))
        if all(v is None for v in vals):
            out["absent"].append({"series": key, "row": label,
                                  "reason": f"'{label}' carries no figure in any of these months"})
            continue
        out["series"][key] = {"row": label, "values": vals}
    if opex_labels:
        vals = []
        for mo in axis:
            tot, seen = 0.0, False
            for label in opex_labels:
                v = to_usd(_num(model.get_line(label, month=mo)))
                if v is not None:
                    tot += v
                    seen = True
            vals.append(tot if seen else None)
        if any(v is not None for v in vals):
            out["series"]["opex"] = {"rows": opex_labels, "values": vals}
    else:
        out["absent"].append({"series": "opex",
                              "reason": "this borrower's account list groups no operating expenses"})
    # Margins view: derived here so the browser performs no arithmetic.
    rev = (out["series"].get("net_revenue") or {}).get("values")
    if rev:
        for num_key, out_key in (("gross_profit", "gross_margin"),
                                 ("ebitda", "ebitda_margin"),
                                 ("opex", "opex_pct_of_revenue")):
            src = (out["series"].get(num_key) or {}).get("values")
            if not src:
                continue
            out["series"][out_key] = {
                "derived_from": [num_key, "net_revenue"],
                "values": [(None if (a is None or not b) else a / b)
                           for a, b in zip(src, rev)],
            }
    return out


# The four dollar series the chart draws as levels. A quarter's level is the three
# months added up; a margin is NOT, because the mean of three monthly margins is not
# the quarter's margin -- a quarter is the quarter's numerator over the quarter's
# revenue. The margin keys are rebuilt from the summed levels for that reason.
QUARTER_LEVELS = ("net_revenue", "gross_profit", "ebitda", "opex")
QUARTER_MARGINS = (("gross_margin", "gross_profit"),
                   ("ebitda_margin", "ebitda"),
                   ("opex_pct_of_revenue", "opex"))


def whole_quarters_drawn(months):
    """Every calendar quarter the chart's window holds ALL THREE months of.

    Oldest first, as (label, short, [month labels], [indices]). The window's own
    ends usually clip a quarter, and a clipped quarter is not a quarter -- it is
    left out here rather than totalled and shown low.
    """
    pos, ends = {}, []
    for i, label in enumerate(months or []):
        try:
            year, mnum = ml._parse_label(label)
        except ValueError:
            continue
        pos[(year, mnum)] = i
        if mnum % 3 == 0:
            ends.append((year, mnum))
    out = []
    for year, endmonth in ends:
        wanted = [(year, endmonth - 2), (year, endmonth - 1), (year, endmonth)]
        if any(k not in pos for k in wanted):
            continue
        idx = [pos[k] for k in wanted]
        qnum = (endmonth - 1) // 3 + 1
        out.append((f"Q{qnum} {year}", f"Q{qnum} {str(year)[-2:]}",
                    [months[i] for i in idx], idx))
    return out


def quarter_callouts(months, series, basis_break_index=None):
    """The quarter totals the chart calls out at each quarter-end month.

    A callout is a claim about a quarter, so it is made only where the whole
    quarter is on the chart AND all three of its months carry a figure for that
    series. A quarter missing a month would otherwise read as a low quarter
    rather than a partial one, which is the error a credit reader cannot see and
    cannot correct.

    Aggregation is done here, not in the browser, for the same reason the margins
    are: a figure computed on the page is gated by nothing and traces to no model
    cell.

    Returns a list, oldest first, of {label, short, end_index, months, values,
    spans_basis_break}. `values` carries only the series the quarter is complete
    for; a quarter complete for nothing is not returned at all.
    """
    out = []
    for label, short, qmonths, idx in whole_quarters_drawn(months):
        levels = {}
        for key in QUARTER_LEVELS:
            vals = ((series.get(key) or {}).get("values") or [])
            got = [vals[i] for i in idx if i < len(vals)]
            if len(got) == 3 and all(v is not None for v in got):
                levels[key] = float(sum(got))
        values = dict(levels)
        rev = levels.get("net_revenue")
        for out_key, num_key in QUARTER_MARGINS:
            if rev and num_key in levels:
                values[out_key] = levels[num_key] / rev
        if not values:
            continue
        out.append({
            "label": label,
            "short": short,
            "end_index": idx[-1],
            "months": qmonths,
            "values": values,
            # A quarter with the break inside it adds two reporting bases together.
            # The total is still stated, and it is marked, because hiding it would
            # leave a gap the reader reads as a missing quarter.
            "spans_basis_break": (basis_break_index is not None
                                  and idx[0] < basis_break_index <= idx[-1]),
        })
    return out


# ------------------------------------------------------------------ statement tables
STATEMENT_SHEETS = [
    ("Income Statement", "income_review", True),
    ("Balance Sheet", "balance_review", False),
    ("Cash Flow Statement", "cashflow_review", True),
]


# The statement layout the build writes ends with the plugin's own working area:
# the EBITDA bridge, the covenant inputs, the balance check and our own exposure
# line. Those rows exist for the engines, and a reader who meets them reads a
# tool's scratch pad in the middle of a financial statement. They are named by the
# build's layout contract, in the plugin's own words, so matching them names no
# borrower's account.
WORKING_HEADERS = ("ebitda bridge", "covenant inputs", "memo checks",
                   "balance check", "our exposure")

# The one working row a reader needs, because it is the figure every covenant is
# measured on and the memo cites it by name.
WORKING_KEEP = ("ebitda (covenant bridge)", "ebitda")

# The memo's own names for rows whose model names are ledger long-forms. Standard
# accounting phrases only — nothing here is a borrower's private vocabulary — and
# every rewrite keeps the model's own name in the row's method note.
FINANCE_LABEL_REWRITES = [
    (r"(?i)^net cash provided by or \(used in\) operating activities$",
     "Operating cash flow"),
    (r"(?i)^net cash provided by or \(used in\) financing activities$", "Financing"),
    (r"(?i)^net increase or \(decrease\) in cash$", "Net change in cash"),
    (r"(?i)^accrued expenses\s*[-–—]?\s*lt$", "Accrued expenses — long term"),
    (r"(?i)^prepaid expenses and other current assets$",
     "Prepaid and other current assets"),
    (r"(?i)costs and estimated earnings in excess of billings",
     "Costs in excess of billings"),
    (r"(?i)billings in excess of costs and estimated earnings",
     "Billings in excess of costs"),
    (r"(?i), net of .+$", ", net"),
    (r",\s+&", " &"),
]


def finance_label(label):
    """A row's printed name, with the ledger long-forms said the short way."""
    s = str(label or "")
    for pat, repl in FINANCE_LABEL_REWRITES:
        s = re.sub(pat, repl, s)
    return s.strip() or str(label or "")

# Rows of the cash-flow statement that state a POSITION rather than a movement. The
# year-to-date column adds a row's months up, which is right for a flow and nonsense
# for a balance -- five months of closing cash added together is not the cash.
# Matched as prefixes: borrowers write "Cash, end of period" and "Cash at the end
# of the period" for the same row.
POSITION_ROWS = ("cash, beginning", "cash, end", "cash at the beginning",
                 "cash at the end")


def statement_rows(model, sheet):
    """Every row of one statement, in the order the statement itself has them,
    each marked body or working.

    A statement read out of order is not a statement: revenue below net income, or
    a cash flow opening on inventory, costs a reader the shape they use to read it.
    The workbook already holds the borrower's own order, so that is the order shown
    — down to where the plugin's own working area begins. Working rows are returned
    rather than dropped, because a few of them ARE the statement: the EBITDA bridge
    holds the interest, tax and depreciation detail the body states only in
    aggregates, and the curation step adopts a working row wherever the arithmetic
    proves it belongs.
    """
    meta = model._meta.get(sheet) or {}
    ordered = sorted(((row_no, orig) for row_no, orig in meta.values()),
                     key=lambda pair: pair[0])
    out, in_working = [], False
    for row_no, orig in ordered:
        low = str(orig or "").strip().lower()
        # A build writes the bridge under a header of its own -- "EBITDA bridge",
        # or just "EBITDA:" with the layout's colon.
        if any(h in low for h in WORKING_HEADERS) or re.fullmatch(r"ebitda\s*:", low):
            in_working = True
        out.append((orig, row_no, in_working))
    return out


def plain_row_label(label):
    """A model row's name as a reader says it, with the bookkeeping code left out.

    A borrower numbers the groups in its chart of accounts, so the model carries rows
    called "Total - 12000 - Accounts Receivable". The number is the borrower's own
    filing system: it tells a fund analyst nothing and it makes a table of eleven
    rows hard to read. The row the figure was read from is still named in the
    figure's own note, so nothing is lost by keeping it off the page.
    """
    s = re.sub(r"\s*\b\d{4,6}\b\s*-\s*", " ", str(label or ""))
    s = re.sub(r"\s+-\s+", " ", s)
    return re.sub(r"\s{2,}", " ", s).strip(" -") or str(label or "")


def _has_ledger_code(label):
    return bool(re.search(r"\b\d{4,6}\b", str(label or "")))


def _fold_aliases(entries):
    """One row per figure, where the model carries the same figure twice.

    A chart of accounts usually holds both the ledger group total and the statement
    subtotal over it -- "Total - 12000 - Accounts Receivable" on row 27 and "Total
    Accounts Receivable" on row 28 -- and both carry the identical amount every
    month. Printed as two rows they read as two different receivables balances and
    add up to twice the real one in the reader's head. Identical every month means
    one figure, so one row is shown, the plainer name is the one kept, and the other
    name is stated in the row's note so the trail back to the model is unbroken.
    """
    groups, order = {}, []
    for i, e in enumerate(entries):
        sig = (tuple(e["cells"]), e.get("ytd"))
        if sig not in groups:
            groups[sig] = []
            order.append(sig)
        groups[sig].append((i, e))
    out = []
    for sig in order:
        grp = groups[sig]
        # Keep the row the borrower already names in plain words, so the note names
        # the coded twin rather than restating the row's own printed name; then the
        # shortest name; then whichever the account list names first.
        _i, keep = min(grp, key=lambda pair: (_has_ledger_code(pair[1]["label"]),
                                              len(plain_row_label(pair[1]["label"])),
                                              pair[0]))
        others = [e["label"] for _j, e in grp if e is not keep]
        if others:
            keep["method"] += (" The borrower's own account list carries the same "
                               "amount under "
                               + ", ".join(f'"{o}"' for o in others) + ".")
        out.append(keep)
    return out


def _display_labels(entries):
    """The name each row is printed under: plain where that is unambiguous.

    Two rows whose plain names would collide keep their full model names -- a table
    with two rows called the same thing is worse than one with a ledger code in it.
    """
    plain = [finance_label(plain_row_label(e["label"])) for e in entries]
    for e, p in zip(entries, plain):
        e["model_row_label"] = e["label"]
        e["label"] = e["label"] if plain.count(p) > 1 else p
    return entries


# ------------------------------------------------------------ statement curation
# The memo's statement tables read like a lender's spread, not a ledger dump: the
# EBITDA bridge's interest and depreciation lines sit where a statement states
# them, detail that only restates a subtotal folds into it, and a row too small to
# print as anything but zero is left out. Every fold below is proven by the rows'
# own arithmetic before it is applied — a collapse that does not foot never
# happens — and each one says in the row's method note what it did.

def _close(a, b, tol):
    if a is None or b is None:
        return a is None and b is None
    return abs(a - b) <= tol


def _cells_match(xs, ys, tol, sign=1):
    return len(xs) == len(ys) and all(
        _close(x, None if y is None else sign * y, tol) for x, y in zip(xs, ys))


def _sums_to(components, total, tol):
    """Whether the component rows add to the total row, month by month."""
    if not components:
        return False
    for i, tv in enumerate(total["cells"]):
        if tv is None:
            continue
        if not _close(sum(c["cells"][i] or 0.0 for c in components), tv, tol):
            return False
    return True


def _is_money(entry):
    return (entry.get("unit") or "usd") == "usd"


def _curate_income(body, work, tol):
    """The income statement, closed the way a statement closes.

    The model's body often stops at the tax line and keeps interest, depreciation,
    net income and EBITDA in its working bridge. Each adoption below is arithmetic,
    not label faith: a bridge row is taken only where the body's own figures prove
    it belongs.
    """
    used = set()
    # An aggregate that is exactly minus one bridge row IS that row: a "Total
    # other income and expense" carrying nothing but interest prints as the
    # interest expense it is.
    for b in body:
        if not _is_money(b) or "total" not in b["label"].lower():
            continue
        for w in work:
            if id(w) in used or not _is_money(w):
                continue
            if _cells_match(b["cells"], w["cells"], tol, sign=-1) and \
                    _close(b.get("ytd"),
                           None if w.get("ytd") is None else -w["ytd"], tol):
                b["method"] += (f' The statement carries this as "{b["label"]}"; '
                                f'every month it equals the model\'s '
                                f'"{w["label"]}" line, so it prints under that '
                                f'name with its natural sign.')
                b["label"], b["cells"] = w["label"], list(w["cells"])
                b["ytd"] = w.get("ytd")
                b["row"] = w.get("row")
                used.add(id(w))
                break
    # A bridge row that duplicates a body row adds nothing.
    for w in work:
        if id(w) in used:
            continue
        for b in body:
            if _is_money(w) == _is_money(b) and \
                    _cells_match(b["cells"], w["cells"], tol):
                used.add(id(w))
                break
    # A subtotal whose shown parts do not foot, where exactly one bridge row
    # fills the gap, adopts that row in its place — depreciation inside the
    # operating expenses block.
    i = 0
    while i < len(body):
        s = body[i]
        if _is_money(s) and s["label"].lower().startswith("total"):
            comps, k = [], i - 1
            while k >= 0 and _is_money(body[k]) and \
                    not body[k]["label"].lower().startswith(
                        ("total", "gross", "income", "net")):
                comps.append(body[k])
                k -= 1
            if comps and not _sums_to(comps, s, tol):
                for w in work:
                    if id(w) in used or not _is_money(w):
                        continue
                    if _sums_to(comps + [w], s, tol):
                        w2 = dict(w)
                        w2["method"] += (" Read from the model's EBITDA bridge; "
                                         "it completes the subtotal below it.")
                        body.insert(i, w2)
                        used.add(id(w))
                        i += 1
                        break
        i += 1
    # The bridge's net income closes the statement when the body states none.
    if not any("net income" in b["label"].lower() for b in body):
        for w in work:
            if id(w) not in used and "net income" in w["label"].lower():
                w2 = dict(w)
                w2["method"] += " Read from the model's EBITDA bridge."
                body.append(w2)
                used.add(id(w))
                break
    # EBITDA itself — the figure the covenants are measured on — and its margin.
    eb = None
    for w in work:
        low = w["label"].strip().lower()
        if id(w) not in used and \
                (low == "ebitda" or low.startswith("ebitda (covenant")):
            eb = dict(w)
            used.add(id(w))
            break
    if eb:
        eb["method"] += (" The model's own EBITDA bridge row — the figure the "
                         "covenants are measured on.")
        eb["label"] = "EBITDA (bridge)"
        body.append(eb)
        rev = next((b for b in body if b["label"].strip().lower() == "revenue"),
                   None)
        if rev:
            margin = {
                "label": "EBITDA margin %", "row": eb.get("row"),
                "sheet": eb.get("sheet"), "unit": "pct", "sub": True,
                "method": ("EBITDA (bridge) over revenue, month by month; the "
                           "year to date is the ratio of the two year-to-date "
                           "figures."),
                "cells": [None if (e is None or not r) else e / r
                          for e, r in zip(eb["cells"], rev["cells"])],
            }
            if eb.get("ytd") is not None and rev.get("ytd"):
                margin["ytd"] = eb["ytd"] / rev["ytd"]
            body.append(margin)
    return body


def _curate_cashflow(body, tol):
    """The cash flow, read to its close: operating detail, one line each for
    investing and financing, and the period-end cash it all lands on."""
    # The statement ends at the period-end cash; what follows — supplemental
    # disclosures, the build's own memo checks — is not the statement.
    for i, r in enumerate(body):
        if r["label"].strip().lower().startswith(("cash at the end", "cash, end")):
            body = body[:i + 1]
            break
    # The opening balance is last month's closing balance: one of the two rows
    # carries information.
    body = [r for r in body if not r["label"].strip().lower()
            .startswith(("cash at the beginning", "cash, beginning"))]
    # Investing and financing collapse to their subtotals when their own detail
    # foots — the memo reads operating detail and treats the other two sections
    # as one line each.
    for word in ("investing activities", "financing activities"):
        idx = next((i for i, r in enumerate(body)
                    if word in r["label"].lower()
                    and r["label"].lower().startswith("net cash")), None)
        if idx is None:
            continue
        acc, k = [], idx - 1
        while k >= 0:
            acc.append(body[k])
            if _sums_to(acc, body[idx], tol):
                body[idx]["method"] += (
                    " Its own lines — " + "; ".join(f'"{c["label"]}"'
                                                    for c in reversed(acc))
                    + " — fold into this one, and they foot.")
                del body[k:idx]
                break
            k -= 1
    # The lone line left between operating and financing is the investing
    # section; it is named for what it is.
    op = next((i for i, r in enumerate(body)
               if "operating activities" in r["label"].lower()
               and r["label"].lower().startswith("net cash")), None)
    fin = next((i for i, r in enumerate(body)
                if "financing activities" in r["label"].lower()
                and r["label"].lower().startswith("net cash")), None)
    if op is not None and fin is not None and fin - op == 2:
        mid = body[op + 1]
        mid["method"] += (f' The statement\'s one investing line, '
                          f'"{mid["label"]}".')
        mid["label"] = f"Investing ({mid['label'][0].lower()}{mid['label'][1:]})"
    return body


def _curate_balance(body, tol):
    """The balance sheet to its equity total: the closing identity restates the
    asset total, and detail that foots into a stated net folds into it."""
    for i, r in enumerate(body):
        low = r["label"].strip().lower()
        if low.startswith("total liabilities and") and "equity" in low:
            body = body[:i]
            break
    # A gross row and its "Less:" contra fold into the net row that foots them.
    out, i = [], 0
    while i < len(body):
        if i + 2 < len(body) and body[i + 1]["label"].lower().startswith("less"):
            net = body[i + 2]
            if net["label"].lower().startswith("net") and \
                    _sums_to([body[i], body[i + 1]], net, tol):
                net["method"] += (f' Net of "{body[i]["label"]}" and '
                                  f'"{body[i + 1]["label"]}", which fold into '
                                  f'this row and foot.')
                out.append(net)
                i += 3
                continue
        out.append(body[i])
        i += 1
    body = out
    # Equity detail folds into its stated total.
    eq = next((i for i, r in enumerate(body)
               if "equity" in r["label"].lower()
               and r["label"].lower().startswith("total")), None)
    if eq is not None:
        acc, k = [], eq - 1
        while k >= 0 and not body[k]["label"].lower().startswith("total"):
            acc.append(body[k])
            if _sums_to(acc, body[eq], tol):
                body[eq]["method"] += (
                    " Its own lines — " + "; ".join(f'"{c["label"]}"'
                                                    for c in reversed(acc))
                    + " — fold into this one, and they foot.")
                del body[k:eq]
                break
            k -= 1
    return body


def _row_is_share(model, sheet, row_no):
    """Whether this statement row states a share rather than an amount, read off
    the workbook's own number formats for that row.

    A row counts as a share when every data cell that carries a format states one
    as a percentage. Anything less is treated as money: withholding the scale from
    a money row prints a borrower at a thousandth of its size, so the ambiguous
    case takes the reading that cannot be silently wrong by a factor of a thousand.
    """
    try:
        ws = model.wb[sheet]
        first = (model.layout.get(sheet) or {}).get("first_data_col")
    except (KeyError, AttributeError):
        return False
    if not row_no or not first:
        return False
    formats = []
    for col in range(int(first), ws.max_column + 1):
        cell = ws.cell(row=int(row_no), column=col)
        if cell.value is None:
            continue
        fmt = str(cell.number_format or "")
        if fmt and fmt.lower() != "general":
            formats.append(fmt)
    return bool(formats) and all("%" in f for f in formats)


def statement_tables(model, period, months_back=5, unit=None):
    """One table per statement sheet: the trailing months, plus a year-to-date
    column on the flow statements only (balances do not sum).

    Each row carries the method the memo's trace box states and, per month, the
    workbook cell the figure came from -- the row store and column store
    report-structure's coverage contract asks for, so the renderer stores a trace
    once per row and once per column rather than per cell.

    Every money cell is stated in dollars. A statement row is money, and a model
    kept in thousands would otherwise print a borrower at a thousandth of its size.

    A statement sheet also carries SHARE rows -- a gross margin, an expense as a
    percentage of revenue. The workbook says which they are: the cell's own number
    format is a percentage one. That is read rather than guessed from the label,
    because a row is a share when the borrower formatted it as a share, and a
    label heuristic mistags a dollar row on the next borrower. A share is not
    money, so it is neither scaled nor added down a year (0.30 + 0.31 is not a
    margin), and it prints as a percentage.
    """
    to_usd = (lambda v: v) if unit is None else unit.to_dollars
    scale = float(getattr(unit, "scale", 1) or 1)
    # Small enough that it can only print as a zero: 2.5 of the model's own units,
    # in dollars. A row of nothing but that is left off the table.
    floor = 2.5 * scale
    tol = 2.0 * scale
    axis = list(model.months or [])
    if period:
        try:
            end = _period_label(period)
            if end in axis:
                axis = axis[:axis.index(end) + 1]
        except (ValueError, IndexError):
            pass
    shown = axis[-months_back:]
    ytd_year = shown[-1].split()[-1] if shown else None
    ytd_months = [m for m in axis if m.split()[-1] == ytd_year] if ytd_year else []
    out = {}
    for sheet, section, is_flow in STATEMENT_SHEETS:
        rows = model.sheets.get(sheet)
        if not rows:
            continue
        table = {"sheet": sheet, "months": shown, "flow": is_flow,
                 "ytd_months": ytd_months if is_flow else [], "rows": []}
        if scale != 1:
            table["display"] = {"scale": scale}
        body, work = [], []
        for orig, row_no, in_working in statement_rows(model, sheet):
            label = orig
            share = _row_is_share(model, sheet, row_no)
            keep = (lambda v: v) if share else to_usd
            cells = [keep(_num(v)) for v in
                     [model.get_line(label, month=mo, sheet=sheet) for mo in shown]]
            if all(v is None for v in cells):
                continue
            if all((v or 0) == 0 for v in cells):
                continue
            entry = {
                "label": orig, "row": row_no, "sheet": sheet,
                "unit": "pct" if share else "usd",
                "method": (f"The share the borrower states on its {sheet.lower()}."
                           if share else
                           f"As the borrower reported it on its {sheet.lower()}."),
                "cells": cells,
            }
            if share:
                entry["sub"] = True
            if is_flow and ytd_months and not share:
                low0 = str(orig).strip().lower()
                if not low0.startswith(POSITION_ROWS):
                    vals = [to_usd(_num(model.get_line(label, month=mo,
                                                       sheet=sheet)))
                            for mo in ytd_months]
                    got = [v for v in vals if v is not None]
                    entry["ytd"] = sum(got) if got else None
                    entry["ytd_months_covered"] = len(got)
                    entry["ytd_months_requested"] = len(ytd_months)
                elif low0.startswith(("cash at the end", "cash, end")):
                    # The year to date ENDS at this balance: the column repeats
                    # the period's own closing cash rather than a nonsense sum.
                    entry["ytd"] = cells[-1]
            (work if in_working else body).append(entry)
        body = _fold_aliases(body)
        body = [r for r in body
                if not _is_money(r)
                or any(abs(v) >= floor for v in r["cells"] if v is not None)
                or (r.get("ytd") is not None and abs(r["ytd"]) >= floor)]
        if section == "income_review":
            body = _curate_income(body, work, tol)
        elif section == "cashflow_review":
            body = _curate_cashflow(body, tol)
        elif section == "balance_review":
            body = _curate_balance(body, tol)
        # A share row prints its own year-to-date ratio when the rows it is built
        # from are on the table -- a margin's year is a ratio, never a sum.
        money = [r for r in body if _is_money(r)]
        for srow in body:
            if not _is_money(srow) and srow.get("ytd") is None:
                pair = next(
                    ((num, den) for num in money for den in money
                     if den is not num
                     and num.get("ytd") is not None and den.get("ytd")
                     and all(sv is None
                             or (nv is not None and dv
                                 and abs(nv / dv - sv) <= 0.002)
                             for sv, nv, dv in zip(srow["cells"], num["cells"],
                                                   den["cells"]))), None)
                if pair:
                    num, den = pair
                    srow["ytd"] = num["ytd"] / den["ytd"]
                    srow["method"] += (
                        f' The year to date is year-to-date {num["label"]} over '
                        f'year-to-date {den["label"]}.')
            if not _is_money(srow) and not srow["label"].rstrip().endswith("%"):
                srow["label"] = srow["label"].rstrip() + " %"
        table["rows"] = _display_labels(body)
        if table["rows"]:
            out[section] = table
    return out


# ------------------------------------------------------------------ debt group
def debt_group(env):
    """The funded-debt breakdown as ONE traced group: its components, its stated
    total, and whether they foot. A reader who adds up the parts and
    gets a different number than the total is reading a defect, so the group
    carries the check rather than leaving it to the eye.

    The totals are put in DOLLARS here, like every other money figure the payload
    carries. The module states them on the model's own unit, and this group is the
    figure the capital structure's total row and the leverage numerator both seat:
    a borrower kept in thousands would otherwise show $22,513 of funded debt
    against $21,500,000 of our own loan.
    """
    fd = env.get("funded_debt")
    if not isinstance(fd, dict):
        return None
    scale = float(((env.get("reporting_unit") or {}).get("scale")) or 1.0)
    total = _num(fd.get("ex_leases_latest"))
    incl = _num(fd.get("incl_leases_latest"))
    total = None if total is None else total * scale
    incl = None if incl is None else incl * scale
    comps = [str(c) for c in (fd.get("components") or [])]
    leases = [str(c) for c in (fd.get("lease_components") or [])]
    return {
        "id": "funded_debt",
        "section": "deal_capstructure",
        "label": "Funded debt",
        "basis": fd.get("basis"),
        "components": comps,
        "lease_components": leases,
        # The group headers the module left out because they carry no balance. A
        # reader adding the components up needs to know they are not missing.
        "empty_rows_excluded": [str(x) for x in (fd.get("empty_rows_excluded") or [])],
        "total_ex_leases": total,
        "total_incl_leases": incl,
        "lease_total": (None if (incl is None or total is None) else incl - total),
        "foots": None if (incl is None or total is None) else incl >= total,
    }


# ------------------------------------------------------------------ the security
class BudgetBalanceSheet:
    """The pinned budget's balance-sheet plan, for the Collateral box's vs-Budget column.

    The plan is read through budget-map's `balance_sheet_rows`, whose keys are the
    collateral-map component labels and whose values take the same three shapes
    budget-vs-actual already reads: a row number, a list of rows to sum, or
    {add, subtract}. One convention across both engines, so a map that works for
    the variance module works here.

    Values come back in the budget workbook's OWN units, which is the model's unit
    by convention -- the pinned budget keeps the borrower's own convention, so the
    caller converts the plan the same way it converts the reported level. A missing
    key or a month the budget does not reach is a blank cell, not an error: the
    column is input-gated all the way down to the individual component.
    """

    def __init__(self, budget_map):
        for req in ("sheet", "period_columns"):
            if req not in budget_map:
                raise ml.ConfigError(f"budget-map is missing required key {req!r}; "
                                     "see config/budget-map.template.json")
        given = budget_map.get("file") or budget_map.get("file_mount")
        if not given:
            raise ml.ConfigError("budget-map needs 'file' (or 'file_mount') -- "
                                 "the budget workbook path")
        self.given_path = str(given)
        wb = openpyxl.load_workbook(
            ml.resolve_path(budget_map.get("file_mount") or budget_map["file"]),
            data_only=True)
        self.sheet_name = budget_map["sheet"]
        try:
            self.ws = wb[self.sheet_name]
        except KeyError:
            raise ml.ConfigError(
                f"budget-map sheet {self.sheet_name!r} not found in the budget "
                f"workbook (sheets: {', '.join(wb.sheetnames)})") from None
        self.period_columns = budget_map["period_columns"]
        self.rows = self.plan_rows(budget_map)
        # Every row spec is checked HERE, while the reader is being built, so a map
        # written by hand with a quoted row number gates the whole column with one
        # sentence naming the key. Left to `value()`, it would raise in the middle
        # of the component loop and take the memo down with it.
        for key, spec in self.rows.items():
            self._row_signs(spec, key)

    @staticmethod
    def plan_rows(budget_map):
        """The components a budget-map really plans a balance for.

        A `_`-prefixed key is a note the analyst left for the next reader of the
        file, not a component -- and the template ships one, so a block carrying
        nothing but notes plans nothing. Read in one place, so the check for "is
        there a plan at all" and the rows the plan is read from cannot disagree:
        counting the notes made a borrower with no plan look like a borrower whose
        plan happened to cover none of its pledged lines.
        """
        return {k: v for k, v in
                (budget_map.get("balance_sheet_rows") or {}).items()
                if not str(k).startswith("_")}

    def plans(self, label):
        """Whether the budget plans a balance for this component at all."""
        return label in self.rows

    def _col(self, period):
        """The column holding an ISO YYYY-MM month, matched on the month rather
        than the exact key: setup writes month-ends ("2026-06-30") and the month
        is what identifies the column, so a map keyed a day either way still
        joins. `_comment` keys read as no month and are skipped."""
        for key, col in self.period_columns.items():
            if str(key)[:7] == str(period)[:7] and isinstance(col, int):
                return col
        return None

    @staticmethod
    def _row_signs(spec, label=None):
        """A row spec as (row, sign) pairs.

        Refuses anything the convention does not name, rather than reading it as
        something else. Dropping a bad entry quietly is the worse failure of the
        two: `[55, "61"]` would publish the plan for row 55 alone, and a wrong plan
        beside a right actual is a number a reader acts on. So one bad entry gates
        the whole column, and the message names it.
        """
        where = ("budget-map balance_sheet_rows"
                 + (f"[{label!r}]" if label else ""))

        def rows(seq, sign, key=None):
            out = []
            for r in seq:
                if isinstance(r, bool) or not isinstance(r, int):
                    raise ml.ConfigError(
                        f"{where}{'' if key is None else '.' + key} names {r!r} as "
                        "a budget row; a row is a whole number, so quote nothing "
                        "and add no text")
                out.append((r, sign))
            return out

        if isinstance(spec, bool):                 # a bool is an int in Python
            pairs = None
        elif isinstance(spec, int):
            pairs = [(spec, 1)]
        elif isinstance(spec, list):
            pairs = rows(spec, 1)
        elif isinstance(spec, dict):
            pairs = (rows(spec.get("add") or [], 1, "add")
                     + rows(spec.get("subtract") or [], -1, "subtract"))
        else:
            pairs = None
        if not pairs:
            raise ml.ConfigError(
                f"{where} holds {spec!r}, which names no budget row; it takes a row "
                "number, a list of rows to sum, or {add, subtract}")
        return pairs

    def value(self, label, period):
        """The planned balance for one component at one month, in the budget's own
        units, with the cells it was read from. (None, []) when the budget plans
        no such component, does not reach the month, or leaves every cell empty."""
        spec = self.rows.get(label)
        col = self._col(period)
        if spec is None or col is None:
            return None, []
        total, refs, seen = 0.0, [], False
        for row, sign in self._row_signs(spec, label):
            v = ml.safe_float(self.ws.cell(row=row, column=col).value)
            if v is None:
                continue
            seen = True
            total += sign * v
            refs.append({"sheet": self.sheet_name, "row": row,
                         "cells": f"{get_column_letter(col)}{row}",
                         "period": str(period)[:7]})
        return (total if seen else None), refs


def read_budget_plan(cfg, absent):
    """The pinned budget's balance-sheet plan, or None with one sentence saying why.

    The vs-Budget column is a column, so its absence gates the column and never the
    box: a borrower with no plan on file still gets every pledged component, its
    level and its moves. The reason is written for the analyst who could put a plan
    on file, which is why each case says a different thing.
    """
    path = (cfg.get("paths") or {}).get("budget_map_path")

    def gate(reason):
        absent.append({"kind": "column_absent", "section": "snapshot_collateral",
                       "what": "The against-plan column in Collateral metrics",
                       "reason": reason})
        return None

    if not path:
        return gate("this borrower's settings name no pinned budget, so the "
                    "pledged assets are not compared against a plan.")
    bmap = _maybe_json(path)
    if bmap is None:
        return gate("the pinned budget's map was not on file "
                    f"({path}), so the pledged assets are not compared against "
                    "a plan.")
    if not BudgetBalanceSheet.plan_rows(bmap):
        return gate("the pinned budget carries no balance-sheet plan, so the "
                    "pledged assets are compared against earlier months only.")
    try:
        return BudgetBalanceSheet(bmap)
    except Exception as e:                 # noqa: BLE001 - gate the column, never fail
        # Every way a workbook can refuse to open lands here, and none of them may
        # take the memo down: a file that is not really an .xlsx raises openpyxl's
        # own InvalidFileException, and a truncated one raises BadZipFile, neither
        # of which is an OSError. The column is the thing that gates.
        return gate(f"the pinned budget could not be read this month ({e}), so "
                    "the pledged assets are not compared against a plan.")


def build_collateral(cfg, model, period, funded_debt, absent, unit=None):
    """What secures the loan, from the borrower's collateral map.

    The map is an optional control file: a borrower without one gets a memo that
    says so. A borrower WITH one gets the box, so read the configured path before
    concluding anything is missing -- saying "no pledged-asset map on file" to a
    reader who has one on disk is worse than leaving the box out.

    Each pledged component names a model row, so the box carries the reported
    values and the coverage cuts computed over them, every figure with the row it
    was read from. `funded_debt` is what coverage is measured against; without it
    the components still print and the cuts do not.

    Each component is also compared against the pinned budget's balance-sheet plan
    for the same month, joined on the component's own label (see
    `read_budget_plan`). The plan gates that one column: a borrower without one
    gets the box with the against-plan cells empty and one sentence saying why.
    """
    path = (cfg.get("paths") or {}).get("collateral_map_path") \
        or cfg.get("collateral_map_path")
    if not path:
        absent.append({"kind": "section_unavailable", "section": "snapshot_collateral",
                       "what": "Collateral metrics",
                       "reason": "this borrower's settings do not point to a "
                                 "pledged-asset map, so what secures the loan is "
                                 "not summarised here."})
        return None
    cm = _maybe_json(ml.resolve_path(path))
    if cm is None:
        absent.append({"kind": "section_unavailable", "section": "snapshot_collateral",
                       "what": "Collateral metrics",
                       "reason": "the pledged-asset map named in this borrower's "
                                 f"settings was not on file ({path}), so what "
                                 "secures the loan is not summarised here."})
        return None
    # `coverage_cuts` maps a cut's name to the component tags it counts. A key
    # beginning with an underscore is a note the analyst wrote to the next reader of
    # the file, not a cut -- counted as one, its text ends up on the page one
    # character at a time.
    raw_cuts = cm.get("coverage_cuts") or {}
    cuts = ([{"name": k, "counts": list(v)} for k, v in raw_cuts.items()
             if not str(k).startswith("_") and isinstance(v, (list, tuple))]
            if isinstance(raw_cuts, dict)
            else [c for c in raw_cuts if isinstance(c, dict)])
    comps = [c for c in (cm.get("pledged_components") or []) if isinstance(c, dict)]

    # Cover is measured against the debt the MAP says it secures, which is not
    # always the whole funded-debt line: here the syndicate slice is pledged and
    # the founders' subordinated notes are not. Using the module's total instead
    # would understate cover against the debt this lien actually stands behind.
    to_usd = (lambda v: v) if unit is None else unit.to_dollars
    fds = cm.get("funded_debt_source") or {}
    secured, secured_basis, secured_ref = None, None, None
    fd_rows = [r for r in (fds.get("model_rows") or []) if r]
    if fd_rows and model is not None:
        vals = [(r, to_usd(_num(model.get_line(r, month=period, sheet="Balance Sheet"))))
                for r in fd_rows]
        got = [(r, v) for r, v in vals if v is not None]
        if got:
            secured = sum(v for _r, v in got)
            secured_basis = ("The debt this lien secures, as the sum of "
                             + ", ".join(f'"{r}"' for r, _v in got)
                             + f" at {_period_label(period)}, per the borrower's "
                             "pledged-asset map.")
            secured_ref = [model.provenance(r, sheet="Balance Sheet", month=period)
                           for r, _v in got]
            miss = [r for r, v in vals if v is None]
            if miss:
                absent.append({
                    "kind": "coverage_incomplete", "section": "snapshot_collateral",
                    "what": "Debt secured by the lien",
                    "reason": "the model has no row named "
                              + ", ".join(f'"{r}"' for r in miss)
                              + ", so the debt cover is measured against is short by "
                                "whatever that row holds."})
    if secured is None and funded_debt:
        secured, secured_basis = funded_debt, (
            "Funded debt as this month's reporting states it. The borrower's "
            "pledged-asset map names no debt rows of its own, so cover is measured "
            "against the whole funded-debt line.")

    # Read each pledged component at the reported month, keeping the row it came
    # from so the figure answers when the reader clicks it.
    # The plan the components are compared against, read once for the whole box.
    # It gates its own column, so the box is built the same way whether or not a
    # plan is on file.
    plan_src = read_budget_plan(cfg, absent)
    plan_store = (_store_of(plan_src.given_path) if plan_src else None)
    plan_in_store = (_in_store(plan_src.given_path) if plan_src else None)

    priced, by_tier, unresolved = [], {}, []
    prior_m = _month_offset(period, -1)
    year_m = _month_offset(period, -12)
    for c in comps:
        spec, tier = c.get("model_line"), str(c.get("tier") or "")
        rows = _rows_of(spec)
        label = c.get("label") or (rows[0] if rows else None)
        v, refs, missing = read_pledged(model, spec, period, to_usd)
        # Collateral is read for movement as well as level: a reader needs to know
        # whether inventory built or ran down, not only where it stands.
        prior = read_pledged(model, spec, prior_m, to_usd)[0] if prior_m else None
        year = read_pledged(model, spec, year_m, to_usd)[0] if year_m else None
        # The plan for the SAME month, in the same dollars as the level beside it:
        # the budget states its figures on the model's own unit, so it converts the
        # same way the reported level does.
        plan, plan_refs = (None, [])
        if plan_src is not None:
            raw, cells = plan_src.value(label, period)
            plan = to_usd(raw)
            plan_refs = [dict(r, store=plan_store, path_in_store=plan_in_store)
                         for r in cells]
        if missing:
            unresolved.append(f'{label} ({", ".join(chr(34) + m + chr(34) for m in missing)})')
        priced.append({"label": label, "tier": tier,
                       "model_rows": rows, "value": v, "source_ref": refs,
                       "prior_month": prior, "prior_year": year,
                       "plan": plan, "plan_ref": plan_refs,
                       "plan_abs": (None if (v is None or plan is None)
                                    else v - plan),
                       "plan_pct": (None if (v is None or not plan)
                                    else v / plan - 1.0),
                       "plan_calc": None if (v is None or plan is None) else
                       f"The reported {label} of {v:,.0f} at "
                       f"{_period_label(period)} against the {plan:,.0f} the "
                       f"pinned budget plans for the same month.",
                       # The page shows each move as a dollar amount with its
                       # percentage beside it, so the dollar amount is a payload
                       # figure too -- a subtraction done at render time is one
                       # nothing gates and nothing can check.
                       "mom_abs": (None if (v is None or prior is None)
                                   else v - prior),
                       "yoy_abs": (None if (v is None or year is None)
                                   else v - year),
                       "mom": (None if (v is None or not prior) else v / prior - 1.0),
                       "yoy": (None if (v is None or not year) else v / year - 1.0),
                       "calc": None if v is None else
                       f"The reported {label} at {_period_label(period)}, as "
                       + (f'the model row "{rows[0]}"' if len(rows) == 1 else
                          "the sum of " + ", ".join(f'"{r}"' for r in rows))
                       + "."})
        if v is not None and tier:
            by_tier[tier] = by_tier.get(tier, 0.0) + v
    if unresolved:
        absent.append({
            "kind": "coverage_incomplete", "section": "snapshot_collateral",
            "what": "Pledged assets not counted",
            "reason": "the model has no row under the name the pledged-asset map "
                      "uses for " + "; ".join(unresolved) + ", so the cover shown "
                      "counts everything except these."})

    for cut in cuts:
        tiers = [t for t in (cut.get("counts") or [])]
        have = [t for t in tiers if t in by_tier]
        cut["value"] = sum(by_tier[t] for t in have) if have else None
        cut["tiers_missing"] = [t for t in tiers if t not in by_tier]
        cut["coverage"] = (None if (cut["value"] is None or not secured)
                           else cut["value"] / secured)
        cut["calc"] = (None if cut["coverage"] is None else
                       f"{cut.get('name')} of {cut['value']:,.0f} divided by the "
                       f"{secured:,.0f} of debt the lien secures, both at "
                       f"{_period_label(period)}.")
        if cut["tiers_missing"]:
            absent.append({
                "kind": "coverage_incomplete", "section": "snapshot_collateral",
                "what": f"{cut.get('name')} coverage",
                "reason": "the model has no rows for part of what this cut counts "
                          f"({', '.join(cut['tiers_missing'])}), so the cover shown "
                          "leaves that part out."})

    if model is None:
        absent.append({"kind": "coverage_incomplete", "section": "snapshot_collateral",
                       "what": "Collateral cover",
                       "reason": "the model workbook was not given, so what is "
                                 "pledged is listed without its reported values."})
    elif not secured:
        absent.append({"kind": "coverage_incomplete", "section": "snapshot_collateral",
                       "what": "Collateral cover",
                       "reason": "nothing on file states the debt total for the lien to be "
                                 "measured against, so cover is not shown."})
    # The table's own total row, computed here so it is a payload figure the gate
    # holds like any other -- a sum typed at render time is one nothing checks.
    with_values = [c for c in priced if c.get("value") is not None]
    total = None
    if with_values:
        def _sum(field):
            got = [c[field] for c in with_values if c.get(field) is not None]
            return sum(got) if got else None
        tv, tm, ty = _sum("value"), _sum("mom_abs"), _sum("yoy_abs")
        # The total is compared against plan only when the budget plans EVERY
        # component the total counts. A total against a plan covering three of five
        # lines reads as a business under plan when what is short is the plan.
        # `unplanned` is read off the SAME list the total is summed from, so the
        # note under the box and the total's own cell can never disagree: a
        # component with no reported value is in neither, and one the budget maps
        # but leaves empty this month is in both.
        unplanned = ([c["label"] for c in with_values if c.get("plan") is None]
                     if plan_src is not None else [])
        tp = None if unplanned else _sum("plan")
        total = {
            "label": "Total pledged components", "value": tv,
            "mom_abs": tm, "yoy_abs": ty,
            "plan": tp,
            "plan_abs": None if (tv is None or tp is None) else tv - tp,
            "plan_pct": None if (tv is None or not tp) else tv / tp - 1.0,
            "plan_calc": None if (tv is None or tp is None) else
            f"The pledged components' total of {tv:,.0f} against the {tp:,.0f} "
            f"the pinned budget plans for the same month.",
            "plan_ref": [r for c in with_values
                         for r in (c.get("plan_ref") or [])][:4],
            "mom": (None if (tv is None or tm is None or tv == tm)
                    else tm / (tv - tm)),
            "yoy": (None if (tv is None or ty is None or tv == ty)
                    else ty / (tv - ty)),
            "source_ref": [r for c in with_values
                           for r in (c.get("source_ref") or [])][:4],
            "calc": (f"The {len(with_values)} pledged components above, added up "
                     f"at each period end."),
        }
        if unplanned:
            absent.append({
                # The page reads this as "<what> is not shown", so it names the plan
                # rather than the assets: the assets themselves are all on the page.
                "kind": "column_partial", "section": "snapshot_collateral",
                "what": "The plan for some pledged assets",
                "reason": "the pinned budget plans no balance for "
                          + ", ".join(unplanned) + " this month, so those lines are "
                          "compared against earlier months only and the total is "
                          "not compared against plan."})
    return {
        "lien_type": cm.get("lien_type"),
        "security_summary": cm.get("security_summary"),
        "pledged_components": priced,
        "pledged_total": total,
        "coverage_cuts": cuts,
        "secured_debt": secured,
        "secured_basis": secured_basis,
        "secured_ref": secured_ref,
        "funded_debt_note": fds.get("note"),
        "period": period,
        "source_ref": {"store": _store_of(path), "path_in_store": _in_store(path)},
    }


def revenue_vs_last_year(sections, period):
    """Revenue year to date against the same months a year earlier.

    The Performance box's headline asks whether the business is growing, which is
    the prior-year comparison -- against-plan is a different question and has its
    own column. Both sides are already published by the run, so this states the
    change between them rather than reading anything new.
    """
    figs = (sections.get("income_review") or {}).get("figures") or []

    def find(label):
        for f in figs:
            if str(f.get("label") or "").strip().lower() == label:
                return f
        return None

    now, before = find("revenue — ytd"), find("revenue — prior ytd")
    if not now or not before:
        return None
    a, b = _num(now.get("value")), _num(before.get("value"))
    if a is None or not b:
        return None
    return {
        "id": "performance.revenue-ytd-vs-prior-year",
        "section": "snapshot_performance",
        "label": "Revenue year to date against last year",
        "value": a / b - 1.0, "unit": "%", "scaled_by": 1.0,
        "calc": f"Revenue of {a:,.0f} for the months to {_period_label(period)} "
                f"against {b:,.0f} for the same months a year earlier.",
        "source_ref": (now.get("source_ref") or []) + (before.get("source_ref") or []),
        "hops": (now.get("hops") or []) + (before.get("hops") or []),
    }


def facility_figures(terms, ref):
    """The facility's money terms as payload figures, seated on the Deal tab.

    The size, the funded balance and anything still available to draw are printed,
    so they are declared -- the deal background is their source, not this run's
    arithmetic, and the figure says so.
    """
    figs = []
    for key, label, calc in (
            ("commitment_usd", "Facility size",
             "The facility's committed size as the deal background records it, "
             "from the executed loan documents."),
            ("funded_usd", "Funded",
             "The amount advanced under the facility as the deal background "
             "records it, from the executed loan documents."),
            ("undrawn_usd", "Still available to draw", terms.get("undrawn_note"))):
        v = _num(terms.get(key))
        if v is None:
            continue
        figs.append({
            "id": f"facility.{_slug(label)}",
            "section": "deal_terms",
            "label": label, "value": v, "unit": "usd", "scaled_by": 1.0,
            "calc": calc,
            "source_ref": [ref] if ref else [],
            "hops": [],
        })
    return figs


def _plan_figures(stem, label, row):
    """A pledged line's against-plan numbers: the plan itself and the gap to it.

    Both are printed, so both are declared -- the percentage alone would leave the
    dollar gap beside it unchecked by the gate.
    """
    if row.get("plan") is None:
        return []
    refs = row.get("plan_ref") or []
    figs = [{
        "id": f"{stem}.plan", "section": "snapshot_collateral",
        "label": f"{label} planned", "value": row["plan"], "unit": "usd",
        "scaled_by": 1.0,
        "calc": f"The {row['plan']:,.0f} the pinned budget plans for {label} "
                "this month.",
        "source_ref": refs, "hops": [],
    }]
    for key, unit in (("plan_abs", "usd"), ("plan_pct", "%")):
        if row.get(key) is None:
            continue
        figs.append({
            "id": f"{stem}.{key.replace('_', '-')}",
            "section": "snapshot_collateral",
            "label": f"{label} against plan",
            "value": row[key], "unit": unit, "scaled_by": 1.0,
            "calc": row.get("plan_calc"),
            "source_ref": refs, "hops": [],
        })
    return figs


def collateral_figures(col):
    """The collateral box's numbers as payload figures.

    A figure the memo prints has to be a figure the payload declares, or the gate
    cannot check the page against the run -- and a number the gate cannot check is
    a number nobody checks.
    """
    figs = []
    for c in (col.get("pledged_components") or []):
        if c.get("value") is None:
            continue
        stem = f"collateral.{_slug(c.get('label'))}"
        figs.append({
            "id": stem,
            "section": "snapshot_collateral",
            "label": c.get("label"), "value": c.get("value"), "unit": "usd",
            "scaled_by": 1.0, "calc": c.get("calc"),
            "source_ref": c.get("source_ref") or [], "hops": [],
        })
        for key, window, base in (("mom", "the month before", c.get("prior_month")),
                                  ("yoy", "the same month a year earlier",
                                   c.get("prior_year"))):
            if c.get(key) is None:
                continue
            figs.append({
                "id": f"{stem}.{key}", "section": "snapshot_collateral",
                "label": f"{c.get('label')} against {window}",
                "value": c[key], "unit": "%", "scaled_by": 1.0,
                "calc": f"{c.get('label')} of {c['value']:,.0f} against "
                        f"{base:,.0f} {window}.",
                "source_ref": c.get("source_ref") or [], "hops": [],
            })
        # The plan is read from the budget workbook, so its figures carry the
        # budget's own cells rather than the model row beside them.
        figs.extend(_plan_figures(stem, c.get("label"), c))
    total = col.get("pledged_total")
    if total and total.get("value") is not None:
        figs.append({
            "id": "collateral.pledged-total", "section": "snapshot_collateral",
            "label": total.get("label"), "value": total["value"], "unit": "usd",
            "scaled_by": 1.0, "calc": total.get("calc"),
            "source_ref": total.get("source_ref") or [], "hops": [],
        })
        for key, window in (("mom", "the month before"),
                            ("yoy", "the same month a year earlier")):
            if total.get(key) is None:
                continue
            figs.append({
                "id": f"collateral.pledged-total.{key}",
                "section": "snapshot_collateral",
                "label": f"{total.get('label')} against {window}",
                "value": total[key], "unit": "%", "scaled_by": 1.0,
                "calc": f"The pledged components' total against {window}.",
                "source_ref": total.get("source_ref") or [], "hops": [],
            })
        figs.extend(_plan_figures("collateral.pledged-total",
                                  total.get("label"), total))
    if col.get("secured_debt") is not None:
        figs.append({
            "id": "collateral.secured-debt", "section": "snapshot_collateral",
            "label": "Debt secured", "value": col["secured_debt"], "unit": "usd",
            "scaled_by": 1.0, "calc": col.get("secured_basis"),
            "source_ref": col.get("secured_ref") or [], "hops": [],
        })
    # A cut is a ratio of figures already read, so its provenance is theirs: the
    # rows it counted plus the debt rows it was divided by.
    for c in (col.get("coverage_cuts") or []):
        if c.get("coverage") is None:
            continue
        counted = set(c.get("counts") or [])
        refs = [r for comp in (col.get("pledged_components") or [])
                if comp.get("tier") in counted and comp.get("value") is not None
                for r in (comp.get("source_ref") or [])]
        refs += [r for r in (col.get("secured_ref") or []) if r]
        figs.append({
            "id": f"collateral.cover.{_slug(c.get('name'))}",
            "section": "snapshot_collateral",
            "label": f"{c.get('name')} cover", "value": c["coverage"], "unit": "x",
            "scaled_by": 1.0, "calc": c.get("calc"),
            "source_ref": refs, "hops": [],
        })
    return figs


# ------------------------------------------------------------ the Deal summary tab
# The words a memo uses for a fact the deal file does not carry. Fixed, because a
# reader learns to skim them: an absent fact reads the same in every card and on
# every borrower.
NOT_ON_FILE = "not on file in the deal file"

# Key terms, in the order the card prints them. Each row: its id, the label the
# reader sees, and the deal-record field names that state it -- several, because
# the same fact is spelled differently across the deals. `facility` is composed
# rather than read, so it names no field.
TERM_ROWS = [
    ("borrower", "Borrower", ("borrowers", "borrower")),
    ("lender_agent", "Lender / agent",
     ("lenders_and_agent", "lender_and_agent", "lenders", "agent")),
    ("sponsor", "Sponsor", ()),                     # from static.sponsor_status
    ("close_date", "Closing date", ("close_date", "closing_date")),
    ("facility", "Facility", ()),                   # composed, see deal_terms
    ("maturity_date", "Maturity", ("maturity_date",)),
    ("amortization", "Amortization", ("amortization", "amortisation")),
    ("interest_rate", "Cash interest rate", ("interest_rate", "cash_interest_rate")),
    ("pik_rate", "PIK rate", ("pik_rate", "pik", "agent_pik")),
    ("oid", "OID", ("oid", "original_issue_discount")),
    ("fees", "Fees / kicker", ("fees",)),
    ("equity_kicker", "Equity kicker", ("equity_kicker", "warrants")),
    ("prepayment", "Prepayment",
     ("prepayment", "prepayment_protection", "call_protection")),
    ("ecf_sweep", "ECF sweep",
     ("ecf_sweep", "excess_cash_flow_sweep", "cash_flow_sweep")),
    ("security", "Collateral", ("security", "collateral")),
    ("borrowing_base", "Borrowing base", ("borrowing_base",)),
]

# The five elements the transaction paragraph is built from, in the fixed house
# order, and where the deal record states each.
TRANSACTION_ELEMENTS = [
    ("facility_structure", "the facility and how it is structured",
     ("instrument", "facility_structure", "structure")),
    ("lenders_agent", "who lent and who acts as agent",
     ("lenders_and_agent", "lender_and_agent", "lenders", "agent")),
    ("use_of_proceeds", "what the money was used for",
     ("use_of_proceeds", "proceeds")),
    ("proforma_leverage_at_close", "leverage at close and the earnings it was measured on",
     ("pro_forma_at_close", "proforma_leverage_at_close", "pro_forma_leverage_at_close")),
    ("security", "what secures the loan", ("security", "collateral")),
]

# The security summary's eight fixed columns, in order.
SECURITY_COLUMNS = [
    "Holder", "Hold at close (par)", "Cost (% of par)",
    "Amortized cost at close", "Principal o/s incl. PIK",
    "Fair value", "Enterprise value", "Loan-to-EV",
]

# The six company facts, three to a column, and where each is stated.
COMPANY_FACTS = [
    ("ownership", "left", "Ownership", ("ownership",)),
    ("business_model", "left", "Business model", ("revenue_model",)),
    ("management", "left", "Management", ("management",)),
    ("scale", "right", "Scale", ("key_operating_kpis", "headcount")),
    ("customers", "right", "Customers", ("customers",)),
    ("end_markets", "right", "End markets", ("end_markets",)),
]


# Where a deal record states how much of the facility is still available to draw.
AVAILABILITY_FIELDS = ("available_usd", "undrawn_usd", "remaining_commitment_usd",
                       "unfunded_commitment_usd")


def availability(fac):
    """How much of the facility is still there to draw, when the record says so.

    A facility size beside a funded balance invites the reader to subtract, and the
    difference is very often not available money: here every advance was
    consolidated into one term loan with no further lender commitment, so a
    $21,200,000 size against $17,721,846 funded looked like $3,478,154 of headroom
    that does not exist. So availability is only ever the figure the deal record
    states, never a subtraction, and a record that does not state it gets the
    sentence saying so.
    """
    for key in AVAILABILITY_FIELDS:
        v = _num(fac.get(key))
        if v is not None:
            return {"undrawn_usd": v,
                    "undrawn_note": f"As the deal background records it, under "
                                    f"{key.replace('_', ' ')}."}
    # A deal whose unfunded room is an accordion at the lender's discretion has no
    # available figure, and saying "not recorded" of a deal whose own draw structure
    # describes two accordions reads as a gap in the file rather than the answer.
    structure = fac.get("draw_structure")
    if structure:
        return {"undrawn_usd": None,
                "undrawn_note": "No figure is stated for what is still available to "
                                "draw, because what is left is not a committed "
                                "balance — see the draw structure below. The gap "
                                "between the facility size and the funded balance is "
                                "not availability."}
    return {"undrawn_usd": None,
            "undrawn_note": "The deal record does not say how much of the "
                            "facility is still available to draw. The gap between "
                            "the size and the funded balance is not it: on a "
                            "consolidated or fully-advanced loan there is nothing "
                            "left to draw."}


def read_deal_context(cfg, absent):
    """The deal record, once, for every card on the Deal summary tab.

    A missing deal record is a memo whose Deal tab says the terms could not be
    stated, not a memo that quietly prints funded debt alone -- so the run
    continues and every card that needed the record says what it needed.
    """
    path = (cfg.get("paths") or {}).get("deal_context_path") \
        or cfg.get("deal_context_path")
    dc = _maybe_json(ml.resolve_path(path)) if path else None
    if dc is None:
        for section, what in (("deal_terms", "Key terms"),
                              ("deal_company", "Company"),
                              ("deal_transaction", "Transaction summary")):
            absent.append({
                "kind": "deal_record_absent", "section": section, "what": what,
                "reason": "the deal record for this borrower is not on file, so the "
                          "facility, the company profile and the transaction are "
                          "not stated on this memo."})
    return dc, path


def doc_ref(path, stores, name=None, dated=None):
    """One document as the memo cites it: where it lives and the link to open it.

    The URL is joined from the store table the run's own results carry, so the
    memo's documents open through exactly the same shared folders its figures
    trace through, and nothing derives a second URL of its own.
    """
    if not path:
        return None
    store, in_store = _store_of(path), _in_store(path)
    url, _why = _url(store, in_store, stores)
    safe_name = os.path.basename(str(in_store or path).replace("\\", "/"))
    return {"name": name or safe_name, "dated": dated, "store": store,
            "path_in_store": in_store, "url": url}


def build_deal_company(dc, cfg, stores, absent):
    """What the company does, and the six facts the card states beside it.

    Six fixed rows in two columns of three, so the card reads the same on every
    borrower and a fact the deal record does not carry is visibly absent rather
    than silently dropped -- an empty column would read as a company with no end
    markets rather than a file with no answer.
    """
    static = (dc or {}).get("static") or {}
    profile = static.get("business_profile") or {}
    memo_path = (cfg.get("paths") or {}).get("ic_uw_memo_path")
    facts, missing = [], []
    for fid, column, label, fields in COMPANY_FACTS:
        text = flatten_value(first_value(profile, *fields))
        facts.append({"id": fid, "column": column, "label": label,
                      "value": text or NOT_ON_FILE, "present": bool(text)})
        if not text:
            missing.append(label.lower())
    if missing:
        absent.append({
            "kind": "deal_fact_absent", "section": "deal_company",
            "what": "Company facts",
            "reason": "the deal record on file does not state "
                      + ", ".join(sorted(missing))
                      + " for this borrower, so those rows say so."})
    return {"one_liner": static.get("business_one_liner"),
            "facts": facts,
            "source_ref": doc_ref(memo_path, stores)}


def covenant_summary_row(spec, absent):
    """The one covenant row the Key terms card carries, in plain words.

    Thresholds live on the Covenants tab, so this row says what the agreement
    TESTS and how the tests combine -- an either/or the reader has to know about
    before reading a leg's own breach word, and the standalone tests beside it.
    """
    covs = [c for c in ((spec or {}).get("covenants") or [])
            if isinstance(c, dict) and c.get("primitive") and not is_analytical(c)]
    if not covs:
        absent.append({
            "kind": "covenant_summary_absent", "section": "deal_terms",
            "what": "Financial covenants",
            "reason": "the agreement on file sets no financial covenant that is "
                      "measured against a figure, so the Key terms card has no "
                      "covenant row."})
        return None
    labels = either_or_labels(spec=spec)
    shorts = covenant_short_names([c.get("name") or c.get("id") for c in covs])
    groups = {}
    standalone = []
    for c in covs:
        name = shorts.get(str(c.get("name") or c.get("id"))) \
            or str(c.get("name") or c.get("id"))
        grp = c.get("either_or_group")
        if grp:
            groups.setdefault(str(grp), []).append(name)
        else:
            standalone.append(name)
    parts = []
    for grp in sorted(groups):
        legs = sorted(str(n) for n in groups[grp])
        parts.append(f"{labels.get(grp) or grp}: only one of "
                     + " or ".join(legs) + " has to be met")
    if standalone:
        names = sorted(str(n) for n in standalone)
        listed = (names[0] if len(names) == 1
                  else ", ".join(names[:-1]) + " and " + names[-1])
        parts.append(("plus " if parts else "") + listed
                     + (", each tested on its own" if len(names) > 1
                        else ", tested on its own"))
    return {"id": "financial_covenants", "label": "Financial covenants",
            "value": "; ".join(parts) + ".",
            "link": "#sec-covenant", "present": True}


def build_deal_terms(dc, dc_path, spec, stores, absent, borrower_name=None):
    """The deal's economics as the Key terms card prints them.

    Sixteen fixed rows in one fixed order, each stating what the executed
    documents set. The facility row is COMPOSED -- size, what is drawn and how it
    may be drawn read as one sentence rather than three cells a reader has to
    add up -- and it carries the availability sentence, because a size beside a
    funded balance invites a subtraction that is very often not available money.
    """
    static = (dc or {}).get("static") or {}
    fac = static.get("facility") or {}
    fees = fac.get("fees") if isinstance(fac.get("fees"), dict) else {}
    ref = doc_ref(dc_path, stores)
    avail = availability(fac)
    rows, missing = [], []
    for rid, label, fields in TERM_ROWS:
        note = fac.get("_" + rid + "_note")
        if rid == "sponsor":
            raw = static.get("sponsor_status")
        elif rid == "facility":
            raw = _facility_sentence(fac)
            note = avail.get("undrawn_note")
        else:
            raw = first_value(fac, *fields) or first_value(fees, *fields)
        text = flatten_value(raw)
        # The borrower is known -- it is the name at the top of this memo -- so a
        # facility record that does not repeat it says the borrower's name rather
        # than "not on file", which reads as a deal record with no borrower.
        if not text and rid == "borrower" and borrower_name:
            text = str(borrower_name)
        kind = ("date" if rid.endswith("_date")
                else "prose" if text and len(text) > 140 else "text")
        rows.append({"id": rid, "label": label,
                     "value": text or NOT_ON_FILE, "kind": kind,
                     "note": note, "source_ref": ref, "present": bool(text)})
        if not text:
            missing.append(label.lower())
    if missing:
        absent.append({
            "kind": "term_absent", "section": "deal_terms", "what": "Key terms",
            "reason": "the executed documents on file do not state "
                      + ", ".join(missing) + " for this deal, so those rows say so."})
    return {"rows": rows,
            "covenant_summary_row": covenant_summary_row(spec, absent),
            # Carried for the figures the card prints, not for the page.
            "commitment_usd": _num(fac.get("commitment_usd")),
            "funded_usd": _num(fac.get("funded_usd")),
            "undrawn_usd": avail.get("undrawn_usd"),
            "undrawn_note": avail.get("undrawn_note"),
            "source_ref": ref}


def _facility_sentence(fac):
    """Size, drawn balance and draw structure as one sentence."""
    size, funded = _num(fac.get("commitment_usd")), _num(fac.get("funded_usd"))
    instrument = flatten_value(fac.get("instrument"))
    bits = []
    if size is not None:
        bits.append(f"${size:,.0f}" + (f" {instrument}" if instrument else " facility"))
    elif instrument:
        bits.append(instrument)
    if funded is not None:
        bits.append(f"${funded:,.0f} advanced to date")
    structure = flatten_value(fac.get("draw_structure"))
    if structure:
        bits.append(structure)
    return "; ".join(bits) or None


def build_deal_transaction(dc, cfg, stores, absent):
    """The five elements the transaction paragraph rests on, and the gaps in them.

    The paragraph itself is the analyst's; what the payload owes it is which of
    the five elements the file actually carries, so an element no document states
    is named to the reader instead of being written around.
    """
    static = (dc or {}).get("static") or {}
    fac = static.get("facility") or {}
    paths = cfg.get("paths") or {}
    elements, gaps = {}, []
    ref = doc_ref(paths.get("credit_agreement_path"), stores)
    for eid, words, fields in TRANSACTION_ELEMENTS:
        raw = first_value(fac, *fields)
        if eid == "facility_structure" and raw is not None:
            raw = _facility_sentence(fac) or raw
        text = flatten_value(raw)
        elements[eid] = {"present": bool(text), "value": text or None,
                         "source_ref": ref if text else None}
        if not text:
            gaps.append({"id": eid, "words": words})
    if gaps:
        absent.append({
            "kind": "transaction_element_absent", "section": "deal_transaction",
            "what": "Transaction summary",
            "reason": "no document on file states "
                      + ", ".join(g["words"] for g in gaps)
                      + ", so the summary names those as not on file."})
    sources = []
    for path, dated in ([(paths.get("credit_agreement_path"), None)]
                        + [(p, None) for p in
                           (paths.get("credit_agreement_amendment_paths") or [])]
                        + [(paths.get("ic_uw_memo_path"), None)]
                        + [(p, None) for p in
                           (paths.get("ic_uw_memo_supplemental_paths") or [])]):
        d = doc_ref(path, stores, dated=dated)
        if d:
            sources.append(d)
    return {"elements": elements, "gaps": gaps, "sources": sources}


def build_deal_security(dc, cfg, period, stores, absent):
    """Who holds the loan and on what basis, as of period end.

    Two rows under eight fixed columns. The holder row restates the executed
    documents and nothing else: no document on file allocates the loan across
    the fund vehicles, so that split is LEFT UNDETERMINED rather than
    inferred from fund names. Cost at close is par net of the discount the deal
    record states; the loan-level balance outstanding is an open request whenever
    the model's funded-debt line mixes our loan with the borrower's others; marks
    and enterprise value live with fund accounting, so they are not on file here.
    """
    static = (dc or {}).get("static") or {}
    fac = static.get("facility") or {}
    agreement_ref = doc_ref((cfg.get("paths") or {}).get("credit_agreement_path"),
                            stores)
    as_of_iso = ml.month_end_iso(_period_label(period)) if period else None
    par = _num(fac.get("funded_usd"))
    oid = leading_pct(fac.get("oid"))

    def cell(value=None, unit=None, words=None, ref=None):
        return {"value": value, "unit": unit, "absent_words": words,
                "source_ref": ref}

    holder = [
        cell(words=("Lender of record per the executed loan and security agreement"),
             ref=agreement_ref),
        cell(words="split undetermined"),
        cell(words="—"), cell(words="—"), cell(words="—"),
        cell(words="—"), cell(words="—"), cell(words="—"),
    ]
    cost_pct = None if oid is None else 1.0 - oid
    total = [
        cell(words="Total"),
        (cell(par, "usd", ref=agreement_ref) if par is not None
         else cell(words=NOT_ON_FILE)),
        (cell(cost_pct, "%", ref=agreement_ref) if cost_pct is not None
         else cell(words="not stated in the deal file")),
        (cell(par * cost_pct, "usd", ref=agreement_ref)
         if (par is not None and cost_pct is not None) else cell(words=NOT_ON_FILE)),
        cell(words="not separable from the model's funded-debt line"),
        cell(words="not on file"), cell(words="not on file"), cell(words="n/a"),
    ]
    if par is None:
        absent.append({
            "kind": "security_basis_absent", "section": "deal_security",
            "what": "Hold at close",
            "reason": "the deal record on file does not state how much was advanced "
                      "at close, so the holding is not stated at par."})
    if cost_pct is None:
        absent.append({
            "kind": "security_basis_absent", "section": "deal_security",
            "what": "Cost at close",
            "reason": "the deal record does not state a discount to par as a "
                      "percentage, so cost at close is not shown against par."})
    absent.append({
        "kind": "security_basis_absent", "section": "deal_security",
        "what": "Principal outstanding, fair value and enterprise value",
        "reason": "the monthly reporting states one funded-debt figure covering "
                  "every borrowing, so the balance of our own loan including "
                  "accrued PIK cannot be separated out; the mark and the "
                  "enterprise value sit with fund accounting and the valuation "
                  "process rather than in the monthly package."})
    return {"as_of": _us_date(as_of_iso),
            "columns": [(c + f" · {_us_date(as_of_iso)}"
                         if c.startswith("Principal") and as_of_iso else c)
                        for c in SECURITY_COLUMNS],
            "rows": [{"id": "holder_of_record", "kind": "holder", "cells": holder},
                     {"id": "total", "kind": "total", "cells": total}]}


def build_deal_capstructure(dc, cfg, model, period, group, stores, absent,
                            unit=None):
    """The capital structure, deduced, with what it could not attribute stated.

    Rows run in one fixed order: our own instrument from the executed documents,
    then any other lender the file names with its ranking against our lien, then
    the residual between those and the funded-debt line the monthly reporting
    carries -- stated, never absorbed, because a balance the package does not
    explain is a finding. The total is the figure the leverage test is measured
    on; memo rows below it are debt-like items that are not funded debt; equity
    is book value and total capitalization is the two added.
    """
    static = (dc or {}).get("static") or {}
    fac = static.get("facility") or {}
    dc_path = (cfg.get("paths") or {}).get("deal_context_path")
    dc_ref = doc_ref(dc_path, stores)
    rows = []
    ours = _num(fac.get("funded_usd"))
    rows.append({
        "id": "our_instrument",
        "label": flatten_value(fac.get("instrument")) or "Our loan",
        "ranking": flatten_value(fac.get("security")),
        "value": ours, "unit": "usd" if ours is not None else None,
        "source_ref": dc_ref, "sub": False})

    # Another lender only ever appears because a document names it. Nothing here
    # reads a lender out of a balance: an unexplained balance is the residual row
    # below, which says it is unexplained.
    named = first_value(static, "other_debt", "prior_liens", "other_lenders") or []
    if isinstance(named, dict):
        named = [named]
    others_total = 0.0
    for i, entry in enumerate(named if isinstance(named, list) else []):
        if not isinstance(entry, dict):
            continue
        v = _num(first_value(entry, "amount_usd", "balance_usd", "value"))
        if v is not None:
            others_total += v
        rows.append({
            "id": f"other_lender.{_slug(entry.get('lender') or entry.get('name') or i)}",
            "label": flatten_value(first_value(entry, "lender", "name")) or "Other lender",
            "ranking": flatten_value(first_value(entry, "ranking", "lien", "priority")),
            "value": v, "unit": "usd" if v is not None else None,
            "source_ref": dc_ref, "sub": False})
    if not rows[1:]:
        absent.append({
            "kind": "debt_schedule_absent", "section": "deal_capstructure",
            "what": "Other lenders",
            "reason": "the borrower delivers no schedule of its borrowings, so no "
                      "other lender is named and the balance below our own loan is "
                      "shown as unattributed."})

    total = _num((group or {}).get("total_ex_leases"))
    if total is None:
        absent.append({
            "kind": "funded_debt_absent", "section": "deal_capstructure",
            "what": "Total funded debt",
            "reason": "this month's review did not produce a funded-debt total for "
                      "the borrower, so the capital structure shows our own loan and "
                      "the equity without a total across them, and nothing states "
                      "what else the borrower owes."})
    fd_ref = None
    if group:
        # The total is the module's own traced group, so the row it seats points at
        # the same rows the group footed.
        fd_ref = {"store": _store_of((cfg.get("paths") or {}).get("model_output_path")),
                  "path_in_store": _in_store((cfg.get("paths") or {}).get(
                      "model_output_path")), "sheet": "Balance Sheet",
                  "period": period}
        fd_ref["url"] = _url(fd_ref["store"], fd_ref["path_in_store"], stores)[0]
    residual = (None if (total is None or ours is None)
                else total - ours - others_total)
    if residual is not None:
        rows.append({
            "id": "unattributed_residual",
            "label": "Other borrowings and leases, unattributed",
            "ranking": "Composition not stated in the monthly reporting",
            "value": residual, "unit": "usd", "source_ref": fd_ref, "sub": False})
    rows.append({
        "id": "total_funded_debt", "label": "Total funded debt",
        "ranking": (group or {}).get("basis"),
        "value": total, "unit": "usd" if total is not None else None,
        "source_ref": fd_ref, "sub": False})
    lease_total = _num((group or {}).get("lease_total"))
    if lease_total:
        rows.append({
            "id": "memo.capital_leases", "label": "Memo: capital leases",
            "ranking": "Included in funded debt only on the lease-inclusive basis",
            "value": lease_total, "unit": "usd", "source_ref": fd_ref, "sub": True})

    equity, equity_ref = None, None
    equity_line = (cfg.get("analysis_lines") or {}).get("total_equity")
    to_usd = (lambda v: v) if unit is None else unit.to_dollars
    if model is not None and equity_line:
        equity = to_usd(_num(model.get_line(equity_line, month=period,
                                            sheet="Balance Sheet")))
        equity_ref = model.provenance(equity_line, sheet="Balance Sheet",
                                      month=period)
    if equity is None:
        absent.append({
            "kind": "equity_absent", "section": "deal_capstructure",
            "what": "Equity at book",
            "reason": "the monthly balance sheet has no equity total under the name "
                      "this borrower's settings use for it, so equity and total "
                      "capitalization are not shown."})
    rows.append({
        "id": "total_equity", "label": "Total equity, at book",
        "ranking": "Book value, not a market figure",
        "value": equity, "unit": "usd" if equity is not None else None,
        "source_ref": equity_ref, "sub": False})
    cap = (None if (total is None or equity is None) else total + equity)
    rows.append({
        "id": "total_capitalization", "label": "Total capitalization, at book",
        "ranking": "Funded debt plus equity at book",
        "value": cap, "unit": "usd" if cap is not None else None,
        "source_ref": fd_ref, "sub": False})
    return {"rows": rows, "funded_debt": group}


def capstructure_figures(cap):
    """Every capital-structure amount the memo prints, declared as a figure.

    The table's own rows are already in dollars, so each is published with the
    document or model row it was read from -- a number the gate cannot check is a
    number nobody checks.
    """
    figs = []
    for row in (cap.get("rows") or []):
        if row.get("value") is None or row.get("unit") != "usd":
            continue
        ref = row.get("source_ref")
        figs.append({
            "id": f"capstructure.{_slug(row.get('id'))}",
            "section": "deal_capstructure",
            "label": row.get("label"), "value": row["value"], "unit": "usd",
            "scaled_by": 1.0,
            "calc": (f"{row.get('label')} at {row.get('ranking') or 'book'}, as the "
                     "capital structure is deduced from the executed documents and "
                     "the month's balance sheet."),
            "source_ref": [ref] if isinstance(ref, dict) else (ref or []),
            "hops": [],
        })
    return figs


def security_figures(sec):
    """The security summary's two stated amounts and the cost basis behind them."""
    figs = []
    columns = sec.get("columns") or []
    for row in (sec.get("rows") or []):
        if row.get("kind") != "total":
            continue
        for i, c in enumerate(row.get("cells") or []):
            if c.get("value") is None or i == 0:
                continue
            label = columns[i] if i < len(columns) else f"column {i}"
            figs.append({
                "id": f"security.{_slug(label)}",
                "section": "deal_security",
                "label": label, "value": c["value"], "unit": c.get("unit"),
                "scaled_by": 1.0,
                "calc": (f"{label} for the whole holding, as the executed loan "
                         "documents state it — par is the amount advanced and cost "
                         "is par net of the discount at which it was bought."),
                "source_ref": ([c["source_ref"]] if isinstance(c.get("source_ref"), dict)
                               else (c.get("source_ref") or [])),
                "hops": [],
            })
    return figs


def _store_of(path):
    logical = recorded_path(path)
    return logical.split(":", 1)[0] if logical else None


def _in_store(path):
    logical = recorded_path(path)
    return logical.split(":", 1)[1] if logical else None


# ------------------------------------------------------------ the Covenants tab
# The three words the Result column is allowed to carry. Fixed, because the word
# is the reader's summary of the row and two rows in the same state must read
# identically.
RESULT_WORDS = {"pass": "PASS", "at_risk": "PASS (at-risk)",
                "breach": "NOT PASSED"}

# How bad each state is, for ordering the grid: what is broken leads, what is
# close to broken follows, what is fine sits at the bottom.
STATE_RANK = {"breach": 0, "at_risk": 1, "pass": 2}


def _threshold_words(primitive, value, unit):
    """A covenant's test as the grid prints it: "≥ $750,000", "≤ 3.50x"."""
    if value is None:
        return None
    op = "≤" if primitive == "max_ratio" else "≥"
    if unit == "usd":
        return f"{op} ${value:,.0f}"
    return f"{op} {value:,.2f}x"


def _metric_words(cov):
    """What a covenant measures, in the agreement's own nouns.

    The spec names each side of a ratio by the definition it shares with every
    module, and the definition names are written for a reader ("Total Senior
    Debt", "Adjusted EBITDA"), so the underscores come out and nothing else does.
    """
    def side(d):
        ref = (d or {}).get("ref") if isinstance(d, dict) else d
        return str(ref or "").replace("_", " ").strip() or None

    num, den = side(cov.get("numerator")), side(cov.get("denominator"))
    if num and den:
        return f"{num} ÷ {den}"
    return side(cov.get("metric")) or num or den


def covenant_preamble(env, spec, cfg, certified, stores, absent):
    """What governs the covenants: the agreement, its amendments, the rule that
    combines the tests, and who signed this month's certificate.

    Every one of these is something the reader can open, so each carries the
    document rather than a description of it. A certificate whose signer and date
    were not captured is said out loud: an unsigned or unattributed certification
    is a finding on the credit, not a blank in a memo.
    """
    paths = cfg.get("paths") or {}
    agreement_ref = doc_ref(paths.get("credit_agreement_path"), stores)
    amendments = []
    for i, path in enumerate(paths.get("credit_agreement_amendment_paths") or []):
        d = doc_ref(path, stores)
        if d:
            amendments.append({"id": f"amendment-{i + 1}", **d})
    stated = [flatten_value(a) for a in ((spec or {}).get("amendments_reflected") or [])]
    for i, text in enumerate([t for t in stated if t]):
        if i < len(amendments):
            amendments[i]["text"] = text
        else:
            amendments.append({"id": f"amendment-{i + 1}", "name": text, "text": text,
                              "dated": None, "store": None, "path_in_store": None,
                               "url": None})
    covs = [c for c in ((spec or {}).get("covenants") or [])
            if isinstance(c, dict) and c.get("primitive")]
    labels = either_or_labels(env, spec)
    grouped = sorted({str(c.get("either_or_group")) for c in covs
                      if c.get("either_or_group")})
    if grouped:
        rule = ("The agreement tests these as alternatives: under "
                + " and ".join(labels.get(g) or g for g in grouped)
                + " only one leg has to be met, so the governing verdict for the "
                  "month is the group's, not a single leg's.")
    else:
        rule = ("Every financial covenant below has to be met in the same month; "
                "the agreement offers no alternative between them.")
    cert = {"signer": flatten_value(first_value(certified or {}, "signer", "signed_by",
                                               "certificate_signer")),
            "signed_date": flatten_value(first_value(certified or {}, "signed_date",
                                                    "certificate_date", "dated")),
            "source_ref": None}
    cert_name = (certified or {}).get("source") or env.get("certified_source")
    if cert_name:
        cert["source_ref"] = {"name": str(cert_name), "store": None,
                              "path_in_store": None, "url": None}
    if not cert_name:
        absent.append({
            "kind": "certificate_absent", "section": "covenant_detail",
            "what": "Compliance certificate",
            "reason": "no compliance certificate for this month has been read, so the "
                      "borrower's own certified figures are not shown beside ours."})
    elif not cert["signer"] or not cert["signed_date"]:
        absent.append({
            "kind": "certificate_signature_absent", "section": "covenant_detail",
            "what": "Certificate signer and date",
            "reason": f"who signed {cert_name} and on what date was not captured when "
                      "its figures were read, so the memo names the certificate "
                      "without naming its signatory."})
    return {"agreement": {"text": flatten_value((spec or {}).get("agreement")),
                          "dated": (spec or {}).get("as_of_agreement_date"),
                          "source_ref": agreement_ref},
            "amendments": amendments,
            "rule_plain": rule,
            "certificate": cert}


def covenant_grid(env, spec, cfg, period, stores, absent):
    """One row per FINANCIAL covenant at this month's test date.

    A covenant that is measured against a figure earns a row; a reporting or
    affirmative obligation does not -- it appears in the memo only when broken,
    and then as a follow-up. Both sides are stated: what the borrower certified
    and what our own model reads, each with its own verdict and headroom, so a
    disagreement between them is visible rather than resolved silently. The Result
    word takes the WORSE of the two -- a breach on either side is not a clean pass.
    """
    by_id = {str(c.get("id")): c for c in ((spec or {}).get("covenants") or [])
             if isinstance(c, dict)}
    labels = either_or_labels(env, spec)
    scale = float(((env.get("reporting_unit") or {}).get("scale")) or 1.0)
    current_end = ml.month_end_iso(_period_label(period)) if period else None
    agreement_ref = doc_ref((cfg.get("paths") or {}).get("credit_agreement_path"),
                            stores)
    grid = []
    for r in (env.get("results") or []):
        if not isinstance(r, dict) or str(r.get("test_date") or "") != current_end:
            continue
        cov = by_id.get(str(r.get("id"))) or {}
        primitive = r.get("primitive") or cov.get("primitive")
        if not primitive or is_analytical(r):
            continue
        unit = "ratio" if primitive in ("max_ratio", "min_coverage") else "usd"
        mul = scale if unit == "usd" else 1.0

        def money(v, _mul=mul):
            n = _num(v)
            return None if n is None else n * _mul

        m_status = str(r.get("status") or "").lower()
        c_status = str(r.get("cert_status") or "").lower()
        worst = min([s for s in (m_status, c_status) if s in STATE_RANK],
                    key=lambda s: STATE_RANK[s], default=None)
        threshold_ref = None
        if agreement_ref:
            threshold_ref = dict(agreement_ref)
            for k in ("section", "page", "quote"):
                v = (cov.get("agreement_ref") or {}).get(k) \
                    if isinstance(cov.get("agreement_ref"), dict) else None
                if v is not None:
                    threshold_ref[k] = v
        stem = f"covenant.{_slug(r.get('id'))}.{_slug(current_end)}"
        grid.append({
            "id": str(r.get("id")),
            "name": r.get("name") or cov.get("name") or str(r.get("id")),
            "sub_label": (labels.get(str(r.get("either_or_group")))
                          or _metric_words(cov)),
            "test": _threshold_words(primitive, money(r.get("required")), unit),
            "threshold_ref": threshold_ref,
            "certified": {"value": money(r.get("cert_value")), "unit": unit,
                          "status": r.get("cert_status"),
                          "headroom": {"abs": money(r.get("cert_headroom_abs")),
                                       "pct": _num(r.get("cert_headroom_pct"))}},
            "model": {"value": money(r.get("actual")), "unit": unit,
                      "status": r.get("status"),
                      "headroom": {"abs": money(r.get("headroom_abs")),
                                   "pct": _num(r.get("headroom_pct"))},
                      "figure_id": f"{stem}.actual"},
            "result_word": RESULT_WORDS.get(worst),
        })
        if r.get("cert_value") is None:
            absent.append({
                "kind": "certified_figure_absent", "section": "covenant_detail",
                "what": f"{r.get('name') or r.get('id')} — the borrower's certified figure",
                "reason": "this month's compliance certificate carries no figure for "
                          "this covenant, so only our own reading of it is shown."})
    grid.sort(key=lambda g: (STATE_RANK.get(
        str((g.get("model") or {}).get("status") or "").lower(), 3),
        STATE_RANK.get(str((g.get("certified") or {}).get("status") or "").lower(), 3),
        str(g.get("name") or "")))
    return grid


def covenant_far_out(env, spec, period, cfg, stores, absent):
    """The covenants whose first test is more than a year away.

    Page-one space belongs to what can move inside the year, so a covenant first
    measured next year gets no chip and no grid row -- one line here instead,
    naming when it first bites and at what level. A covenant already being tested
    is never far out, whatever its schedule says.
    """
    current_end = ml.month_end_iso(_period_label(period)) if period else None
    if not current_end:
        return []
    horizon = _iso_plus_months(current_end, 12)
    tested = {str(r.get("id")) for r in (env.get("results") or [])
              if isinstance(r, dict) and str(r.get("test_date") or "") <= current_end
              and r.get("actual") is not None}
    scale = float(((env.get("reporting_unit") or {}).get("scale")) or 1.0)
    agreement_ref = doc_ref((cfg.get("paths") or {}).get("credit_agreement_path"),
                            stores)
    out = []
    for cov in ((spec or {}).get("covenants") or []):
        if not isinstance(cov, dict) or not cov.get("primitive"):
            continue
        cid = str(cov.get("id"))
        if cid in tested or is_analytical(cov):
            continue
        first = _leading_iso_date(cov.get("first_test_date"))
        # A future test date the run itself already knows about beats the spec's
        # own words -- the schedule is what the engine measured against. It has to be
        # THIS covenant's date: read across all of them, the soonest next test of any
        # covenant became every candidate's first test, pulled each one inside the
        # year, and dropped the whole block -- so a covenant that cannot be tested
        # for years went unmentioned.
        future = sorted(str(r.get("test_date")) for r in (env.get("results") or [])
                        if isinstance(r, dict) and str(r.get("id")) == cid
                        and str(r.get("test_date") or "") > current_end)
        if future:
            first = min([d for d in (first, future[0]) if d] or [None])
        if not first or first <= horizon:
            continue
        primitive = cov.get("primitive")
        unit = "ratio" if primitive in ("max_ratio", "min_coverage") else "usd"
        level = _num(cov.get("threshold"))
        if level is None:
            sched = [s for s in (cov.get("threshold_schedule") or [])
                     if isinstance(s, dict)]
            first_step = min(sched, key=lambda s: str(s.get("effective") or "")) \
                if sched else None
            level = _num((first_step or {}).get("value"))
        ref = None
        if agreement_ref:
            ref = dict(agreement_ref)
            for k in ("section", "page", "quote"):
                v = (cov.get("agreement_ref") or {}).get(k) \
                    if isinstance(cov.get("agreement_ref"), dict) else None
                if v is not None:
                    ref[k] = v
        out.append({
            "id": cid, "name": cov.get("name") or cid,
            "first_test_date": first,
            "threshold": _threshold_words(
                primitive, (None if level is None
                            else level * (scale if unit == "usd" else 1.0)), unit),
            "threshold_ref": ref})
        absent.append({
            "kind": "covenant_not_yet_tested", "section": "covenant_detail",
            "what": cov.get("name") or cid,
            "reason": f"this covenant is first measured on {first}, more than a year "
                      "after the month being reported, so it carries no verdict yet."})
    return sorted(out, key=lambda x: (x.get("first_test_date") or "",
                                      x.get("name") or ""))


def covenant_stress_inputs(env, spec, period, stores, cfg, absent):
    """The traced facts a forward read rests on.

    Three of them, and no scenario: where the agreement's own schedule goes next,
    how much room the binding covenant has today, and what size of move lands each
    test exactly on its threshold. The forward narrative is the analyst's; these
    are the arithmetic it has to agree with.
    """
    current_end = ml.month_end_iso(_period_label(period)) if period else None
    scale = float(((env.get("reporting_unit") or {}).get("scale")) or 1.0)
    by_id = {str(c.get("id")): c for c in ((spec or {}).get("covenants") or [])
             if isinstance(c, dict)}
    agreement_ref = doc_ref((cfg.get("paths") or {}).get("credit_agreement_path"),
                            stores)
    steps = []
    for cid in sorted(by_id):
        cov = by_id[cid]
        sched = [s for s in (cov.get("threshold_schedule") or [])
                 if isinstance(s, dict) and s.get("effective")]
        if not sched:
            continue
        sched.sort(key=lambda s: str(s.get("effective")))
        # The step in force now, and every step after it: a reader needs the level
        # being tested today to make sense of the ones that follow.
        live = [s for s in sched if str(s["effective"]) <= (current_end or "")]
        keep = ([live[-1]] if live else []) + \
               [s for s in sched if str(s["effective"]) > (current_end or "")]
        primitive = cov.get("primitive")
        unit = "ratio" if primitive in ("max_ratio", "min_coverage") else "usd"
        for s in keep:
            v = _num(s.get("value"))
            steps.append({
                # Which covenant steps, as well as when: two covenants on
                # schedules would otherwise produce one unreadable list.
                "covenant": cid,
                "test_date": str(s["effective"]),
                "required": (None if v is None
                             else v * (scale if unit == "usd" else 1.0)),
                "unit": unit,
                "source_ref": agreement_ref})
    if not steps:
        absent.append({
            "kind": "schedule_absent", "section": "covenant_detail",
            "what": "Covenant step-downs",
            "reason": "no covenant in this agreement tightens on a schedule, so there "
                      "are no future step levels to read forward against."})

    current = [r for r in (env.get("results") or [])
               if isinstance(r, dict) and str(r.get("test_date") or "") == current_end
               and r.get("actual") is not None and r.get("primitive")
               and not is_analytical(r)]
    headroom_now = None
    binding = [r for r in current if _num(r.get("headroom_pct")) is not None]
    if binding:
        tight = min(binding, key=lambda r: (_num(r.get("headroom_pct")),
                                            str(r.get("id"))))
        # Which covenant this is comes from the figure it points at, so it is
        # named once rather than twice.
        headroom_now = {
            "value": _num(tight.get("headroom_pct")),
            "unit": "%",
            "figure_id": (f"covenant.{_slug(tight.get('id'))}."
                          f"{_slug(current_end)}.headroom_pct")}

    breakevens = []
    for r in sorted(current, key=lambda r: str(r.get("id"))):
        cid = str(r.get("id"))
        primitive, name = r.get("primitive"), r.get("name") or cid
        actual, required = _num(r.get("actual")), _num(r.get("required"))
        if actual is None or required is None:
            continue
        unit = "ratio" if primitive in ("max_ratio", "min_coverage") else "usd"
        ref = r.get("source_ref") or []
        if primitive == "min_level":
            breakevens.append({
                "id": f"{cid}.floor", "label": f"{name} — the level it breaches below",
                "value": required * scale, "unit": "usd",
                "calc": (f"The agreement's floor. It stands at "
                         f"${actual * scale:,.0f} today, "
                         f"${(actual - required) * scale:,.0f} above the floor."),
                "source_ref": ref})
        elif primitive == "max_ratio" and actual:
            move = 1.0 - actual / required if required else None
            breakevens.append({
                "id": f"{cid}.earnings_cushion",
                "label": f"{name} — fall in earnings that reaches the limit",
                "value": move, "unit": "%",
                "calc": (f"The test reads {actual:,.2f}x against a limit of "
                         f"{required:,.2f}x, so the earnings it is measured on can "
                         f"fall {abs(move):.1%} before the ratio reaches the limit."
                         if move is not None and move >= 0 else
                         f"The test already reads {actual:,.2f}x against a limit of "
                         f"{required:,.2f}x; earnings have to rise "
                         f"{abs(move):.1%} to bring it back to the limit."),
                "source_ref": ref})
        elif primitive == "min_coverage" and actual:
            move = required / actual - 1.0
            breakevens.append({
                "id": f"{cid}.coverage_gap",
                "label": f"{name} — move needed to reach the test",
                "value": move, "unit": "%",
                "calc": (f"The test reads {actual:,.2f}x against a minimum of "
                         f"{required:,.2f}x, so the earnings it is measured on have "
                         f"to rise {move:.1%} to reach it." if move >= 0 else
                         f"The test reads {actual:,.2f}x against a minimum of "
                         f"{required:,.2f}x, so it can fall {abs(move):.1%} before "
                         f"reaching the minimum."),
                "source_ref": ref})
    return {"schedule_steps": steps, "headroom_now": headroom_now,
            "breakevens": breakevens}


# --------------------------------------------------- the two module-gated blocks
# The ARR bridge's legs, in the order a bridge is read: where the base started,
# what was added, what came off, where it ended.
ARR_LEGS = [
    ("opening", "Opening recurring revenue"),
    ("new", "New customers"),
    ("expansion", "Expansion on existing customers"),
    ("contraction", "Contraction on existing customers"),
    ("churn", "Churn"),
    ("acquired", "Acquired"),
    ("adjustments", "Adjustments"),
    ("closing", "Closing recurring revenue"),
]

# Ageing buckets in the order every export states them. A bucket a borrower's own
# export names differently follows, sorted, so nothing is dropped.
BUCKET_ORDER = ["current", "1-30", "31-60", "61-90", ">90"]


def revenue_quality_block(env):
    """The recurring-revenue bridge and retention, when the module ran.

    Both come from the module's own archived record: the bridge legs read straight
    off the borrower's retention workbook, and gross and net retention as the
    module computed them from those legs. Nothing is recomputed here -- a second
    retention figure is a third number nobody can check.
    """
    sections = {str(r.get("section")): r for r in (env.get("results") or [])
                if isinstance(r, dict)}
    bridge = sections.get("arr_bridge") or {}
    ret = sections.get("retention") or {}
    if not bridge.get("available") and not ret.get("available"):
        return None
    scale = float(((env.get("reporting_unit") or {}).get("scale")) or 1.0)
    waterfall = []
    for key, label in ARR_LEGS:
        v = _num(bridge.get(key))
        if v is None:
            continue
        waterfall.append({"id": key, "label": label, "value": v * scale,
                          "unit": "usd",
                          "source_ref": bridge.get("source_ref") or []})
    return {
        "waterfall": waterfall,
        "retention": {"gross": _num(ret.get("grr_pct")),
                      "net": _num(ret.get("nrr_pct")),
                      "basis": bridge.get("basis") or ret.get("basis")},
    }


def aging_block(env):
    """Receivables and payables by age, when the module ran.

    The buckets and the totals are the borrower's own export as the module read
    it; the share each bucket is of its total comes with them. Money is stated in
    dollars.
    """
    scale = float(((env.get("reporting_unit") or {}).get("scale")) or 1.0)
    out, dates, refs = {}, {}, []
    for r in (env.get("results") or []):
        if not isinstance(r, dict):
            continue
        side = str(r.get("side") or "").lower()
        latest = r.get("latest") if isinstance(r.get("latest"), dict) else None
        if side not in ("ar", "ap") or not latest:
            continue
        dist = latest.get("bucket_distribution") or {}
        names = [b for b in BUCKET_ORDER if b in dist] + \
                sorted(b for b in dist if b not in BUCKET_ORDER)
        buckets = []
        for b in names:
            entry = dist.get(b) or {}
            amt = _num(entry.get("amount"))
            buckets.append({"label": b,
                            "value": None if amt is None else amt * scale,
                            "pct": _num(entry.get("pct_of_total"))})
        total = _num(latest.get("total"))
        side_refs = latest.get("source_ref") or []
        out[side] = {"buckets": buckets,
                     "total": None if total is None else total * scale,
                     "source_refs": side_refs,
                     "calc": latest.get("calc")}
        if latest.get("as_of"):
            dates[side] = str(latest["as_of"]).strip()
        refs.extend(side_refs)
    if not out:
        return None
    out.setdefault("ar", {"buckets": [], "total": None})
    out.setdefault("ap", {"buckets": [], "total": None})
    # The two sides do not always stand at the same date: a month the borrower
    # sent a payables export but not a receivables one leaves the receivables a
    # month behind. One date over both would put the wrong date on one of them,
    # so where they differ the block says both.
    seen = sorted(set(dates.values()))
    if len(seen) > 1:
        out["as_of"] = "; ".join(f"{side.upper()} {dates[side]}"
                                 for side in ("ar", "ap") if side in dates)
    else:
        out["as_of"] = seen[0] if seen else None
    out["source_refs"] = refs
    return out


# ---------------------------------------------------------- the Files & audit tab
# The registry's groups, in the order the tab prints them.
FILE_GROUPS = ["loan_documentation", "underwriting", "monthly_reporting", "budget",
               "working_analysis"]


def _first_sentence(text):
    """The first sentence of a note, for a table cell that has one line to give.

    A pinned-budget note carries which tab, which columns, who pinned it and what
    it was checked against -- all of it worth keeping, none of it worth a column
    forty lines tall. The cell says what the document covers; the row's own words
    carry the rest.
    """
    s = str(text or "").strip()
    if not s:
        return None
    cut = re.search(r"\.\s", s)
    return (s[:cut.start() + 1] if cut else s)


def _after_first_sentence(text):
    """Whatever the note says after its first sentence, or nothing."""
    s = str(text or "").strip()
    first = _first_sentence(s)
    rest = s[len(first):].strip() if first else ""
    return rest or None


def build_files_registry(cfg, archive, archive_dir, period, stores, absent, spec=None):
    """Every original document this memo rests on, grouped and linked.

    The registry answers one question -- where did this come from -- so it lists
    sources and repeats no figure. A document the memo relies on but cannot open
    is still listed, with where it sits said in words, because a row missing from
    the registry reads as a document the memo never used.
    """
    paths = cfg.get("paths") or {}
    docs = []

    def add(group, path, name=None, dated=None, covers=None, doc_id=None,
            words=None):
        d = doc_ref(path, stores, name=name, dated=dated)
        if d is None:
            return
        docs.append({"id": doc_id or f"{group}.{_slug(d['name'])}", "group": group,
                     "name": d["name"], "dated": dated or covers, "words": words,
                     "store": d["store"], "path_in_store": d["path_in_store"],
                     "url": d["url"], "linked": bool(d["url"])})

    # The agreement's own date is already settled elsewhere in this payload -- the
    # covenant spec records it, because every threshold cites it. A registry that
    # says "--" beside the document governing every covenant figure sends the
    # reader to open the file to learn something the run already knows.
    add("loan_documentation", paths.get("credit_agreement_path"),
        dated=(spec or {}).get("as_of_agreement_date"),
        doc_id="loan_documentation.credit_agreement")
    for i, p in enumerate(paths.get("credit_agreement_amendment_paths") or []):
        add("loan_documentation", p, doc_id=f"loan_documentation.amendment-{i + 1}")
    add("underwriting", paths.get("ic_uw_memo_path"),
        doc_id="underwriting.ic_uw_memo")
    for i, p in enumerate(paths.get("ic_uw_memo_supplemental_paths") or []):
        add("underwriting", p, doc_id=f"underwriting.supplemental-{i + 1}")

    # This period's reporting package, read off the map the model build wrote when
    # it filled each column -- through the SAME reader every figure trace uses, so
    # the registry lists exactly the files the figures hop to and nothing else.
    # One row per file, however many statements it fed.
    smap = None
    if paths.get("model_output_path"):
        try:
            smap = ml.model_source_map(ml.resolve_path(paths["model_output_path"]))
        except (OSError, ValueError):
            smap = None
    pkg = set()
    for sheet, _section, _flow in STATEMENT_SHEETS:
        for entry in ml.upstream_sources(smap, sheet, period):
            if isinstance(entry, dict) and entry.get("file"):
                pkg.add(str(entry["file"]))
    for p in sorted(pkg):
        add("monthly_reporting", p, dated=period,
            doc_id=f"monthly_reporting.{_slug(os.path.basename(_in_store(p) or str(p)))}")
    if not pkg:
        absent.append({
            "kind": "package_not_listed", "section": "files_registry",
            "what": "This month's reporting package",
            "reason": "the record of which file filled each month of the model does "
                      "not name one for this month, so the package is not listed."})
    cert_name = _certificate_name(archive)
    if cert_name:
        docs.append({
            "id": "monthly_reporting.compliance_certificate",
            "group": "monthly_reporting", "name": cert_name, "dated": period,
            "store": None, "path_in_store": None, "url": None, "linked": False})

    add("budget", paths.get("budget_path"),
        covers=_first_sentence(paths.get("budget_version_note")),
        words=_after_first_sentence(paths.get("budget_version_note")),
        doc_id="budget.plan")
    if not paths.get("budget_path"):
        absent.append({
            "kind": "budget_absent", "section": "files_registry",
            "what": "The year's plan",
            "reason": "no budget is on file for this borrower, so the registry lists "
                      "none and the memo compares the month against earlier months "
                      "only."})

    add("working_analysis", paths.get("model_output_path"),
        doc_id="working_analysis.model")
    for mod in sorted(archive):
        name = archive[mod]["name"]
        stored = ml.unresolve_path(os.path.join(archive_dir, name)) or name
        d = doc_ref(stored, stores, name=name)
        docs.append({"id": f"working_analysis.{_slug(mod)}", "group": "working_analysis",
                     "name": name, "dated": period,
                     "store": (d or {}).get("store"),
                     "path_in_store": (d or {}).get("path_in_store"),
                     "url": (d or {}).get("url"),
                     "linked": bool((d or {}).get("url"))})
    # `linked` says one thing only: whether this row can be opened. It is settled
    # here, from the URL, so no row can claim a link it does not have and the
    # registry cannot emit a dead one.
    for d in docs:
        d["linked"] = bool(d.get("url"))
    unlinked = [d["name"] for d in docs if not d["linked"]]
    if unlinked:
        absent.append({
            "kind": "document_not_linked", "section": "files_registry",
            "what": "Documents the registry cannot open",
            "reason": "the registry lists " + ", ".join(sorted(unlinked))
                      + " without a link, because the memo knows the document by "
                        "name but not where it sits in a shared folder."})
    order = {g: i for i, g in enumerate(FILE_GROUPS)}
    return sorted(docs, key=lambda d: (order.get(d["group"], len(FILE_GROUPS)),
                                       str(d["name"])))


def _sibling(model_path, name):
    """A file the model build leaves beside the workbook, in store form."""
    store, in_store = _store_of(model_path), _in_store(model_path)
    if store and in_store:
        base = in_store.replace("\\", "/").rsplit("/", 1)[0]
        return f"{store}:{base}/{name}"
    return os.path.join(os.path.dirname(str(model_path)), name)


def _certificate_name(archive):
    """The compliance certificate this month's covenant work read, by name."""
    entry = archive.get("covenant-compliance")
    env = (entry or {}).get("data") or {}
    return env.get("certified_source") or None


def build_data_quality(cfg, period, degrade, absent, stores, budget_present):
    """The notes the RUN itself knows, written for whoever reads the credit.

    Everything here is mechanical: which links could not be built, whether a plan
    was on file, which months the borrower restated, and which figures the model
    could not find. The judgment notes -- capitalization policy, an unreconciled
    delta, what a format cost us -- are the analyst's and arrive separately. Every
    note names what it affects and what to trust, in words a reader can act on;
    the engine's own phrasing is put through `plain_reason` first, because a
    sentence written for whoever wrote the engine teaches a fund analyst nothing
    and worries them.
    """
    notes, seen = [], set()

    def note(nid, lead, text):
        text = " ".join(str(text or "").split())
        # One note per thing to say. The same gap is reported by every figure it
        # touched, and a reader does not need the same sentence eleven times.
        if text and nid not in seen:
            seen.add(nid)
            notes.append({"id": nid, "lead": lead, "text": text})

    if not stores:
        note("links.no-shared-folders", "Files are named but do not open",
             "Every figure on this memo names the file it came from, and none of "
             "them opens: this borrower's documents sit in a folder the monitor has "
             "not been given. The figures are unaffected.")
    if "no_store" in degrade:
        note("links.file-not-in-a-known-folder", "Some files are named but do not open",
             "A few figures name a file that sits outside the folders this memo can "
             "reach, so those figures say where they came from without opening it. "
             "The figures themselves are unaffected.")
    if "no_url_base" in degrade:
        note("links.folder-has-no-address", "Some links could not be built",
             "For some files the monitor knows the folder but not its web address, so "
             "those figures name the file without a link rather than guess one.")
    if "no_upstream" in degrade:
        note("links.stops-at-the-model", "Some figures trace to the model, not past it",
             "Some figures trace to the cell of the monitoring workbook they were read "
             "from but not on to the borrower's own monthly file for that month. The "
             "workbook cell is the furthest the trail goes for those figures.")
    if not budget_present:
        note("plan.none-on-file", "No plan to compare against",
             "No budget for this year is on file, so every comparison on this memo is "
             "against earlier months and the same months last year, never against plan.")

    for a in absent:
        kind = a.get("kind")
        if kind == "line_not_found":
            note(f"model.line-missing.{_slug(a.get('what'))}",
                 "A statement line was not found",
                 f"The monitoring workbook holds no line called \"{a.get('what')}\", so "
                 f"any figure that needs it is left off rather than shown as zero.")
        elif kind == "unit_unstated":
            # The figure's own label ends with the field it was read from
            # ("EBITDA (as-reported) — actual"); the reader wants the metric.
            metric = str(a.get("what") or "").split(" — ")[0]
            note(f"figure.withheld.{_slug(metric)}",
                 "A figure was withheld",
                 f"{metric} is not shown: the record behind it does not say whether "
                 f"it is a money amount, and this borrower's statements are not kept "
                 f"in whole dollars, so stating it either way could be wrong by a "
                 f"factor of a thousand.")
        elif kind == "ambiguous_row":
            note("model.repeated-row-names", "Two lines share one name",
                 plain_reason(a.get("reason")))
        elif kind == "series_absent":
            note(f"chart.line-missing.{_slug(a.get('what'))}",
                 "A line is missing from the trend chart",
                 plain_reason(a.get("reason")))
        elif kind == "quarter_absent":
            note(f"chart.quarter-missing.{_slug(a.get('what'))}",
                 "A quarter total is missing from the trend chart",
                 plain_reason(a.get("reason")))
        elif kind == "document_not_linked":
            note("links.documents-cannot-be-opened",
                 "Some source documents are listed but do not open",
                 plain_reason(a.get("reason")) + " The documents themselves are "
                 "unchanged; only the link from this memo is missing.")

    for r in restatement_notes(cfg, stores):
        if r["id"] not in seen:
            seen.add(r["id"])
            notes.append(r)
    # Grouped the way the tab reads: what the memo could not link, what it had no
    # plan to compare against, what the model could not find, what was withheld,
    # then the months the borrower restated -- and alphabetical inside a group, so
    # two runs over the same archive list them in the same order.
    rank = {"links": 0, "plan": 1, "model": 2, "chart": 3, "figure": 4,
            "restated": 5}
    return sorted(notes, key=lambda n: (rank.get(n["id"].split(".")[0], 9),
                                        n["id"]))


def restatement_notes(cfg, stores):
    """One note per figure the borrower restated after first reporting it.

    A restated month is the reason two memos can quote different numbers for the
    same month, so each one is stated: which month, which line, what it was and
    what it became, and the file that carries the new figure.
    """
    model_path = (cfg.get("paths") or {}).get("model_output_path")
    if not model_path:
        return []
    entries = _maybe_json(_sibling(model_path, "restatements.json"))
    if not isinstance(entries, list):
        return []
    unit = ml.reporting_unit(cfg)
    out, seen = [], set()
    for e in entries:
        if not isinstance(e, dict):
            continue
        old, new = _num(e.get("old")), _num(e.get("new"))
        line, month = e.get("line"), e.get("period")
        nid = f"restated.{_slug(month)}.{_slug(line)}.{_slug(old)}.{_slug(new)}"
        if nid in seen:
            continue
        seen.add(nid)
        src = doc_ref(e.get("source"), stores)
        moved = ("" if (old is None or new is None) else
                 f" from ${unit.to_dollars(old):,.0f} to ${unit.to_dollars(new):,.0f}")
        out.append({
            "id": nid,
            "lead": f"{month} was restated",
            "text": (f"The borrower's later reporting restates \"{line}\" for "
                     f"{month}{moved}"
                     + (f", per {src['name']}." if src else ".")
                     + (" The change is large enough to affect a comparison across "
                        "that month." if e.get("material") else ""))})
    return out


# --------------------------------------------------------------------- follow-ups
# The three tiers the memo prints follow-ups in, and how a tier is recognised.
# Tier 1 is anything touching a covenant, tier 2 a required report that did not
# arrive, tier 3 everything else.
TIER_OF_CATEGORY = {"covenant": 1, "missing-deliverable": 2, "missing_deliverable": 2}
TIER_1_WORDS = ("covenant", "compliance certificate", "breach", "leverage",
                "headroom", "cure")
TIER_2_WORDS = ("not delivered", "not been delivered", "was not filed",
                "not filed", "not received", "never been delivered",
                "has not arrived", "resume monthly")


def _tier_of(category, text):
    """Which tier a follow-up belongs in.

    The category the deal record carries GOVERNS -- including when it says the item
    is not one of the first two tiers. Reading the words instead would promote an
    item the analyst deliberately filed as ordinary just because it mentions a
    covenant test in passing. Only an item with no category at all is classified
    from its own words, so an item written before the field existed still sorts.
    """
    stated = str(category or "").strip().lower()
    if stated:
        return TIER_OF_CATEGORY.get(stated, 3)
    low = str(text or "").lower()
    if any(w in low for w in TIER_1_WORDS):
        return 1
    if any(w in low for w in TIER_2_WORDS):
        return 2
    return 3


def build_flags(dc, absent):
    """The open items the memo numbers, sorted the way it prints them.

    Three tiers -- what touches a covenant, what the borrower owes and has not
    sent, then everything else -- and inside a tier the dollar-valued items lead,
    largest first, because size is the only ranking a reader can check. An item
    with no dollar figure follows the valued ones rather than being dropped.
    """
    cf = (dc or {}).get("carry_forward") or {}
    items = []
    for key, category_default in (("open_requests_to_borrower", None),
                                  ("unresolved_flags", None),
                                  ("analyst_focus", "other")):
        for e in (cf.get(key) or []):
            if not isinstance(e, dict):
                text = flatten_value(e)
                if not text:
                    continue
                e = {"item": text}
            text = flatten_value(first_value(e, "item", "flag", "request",
                                             "question"))
            if not text:
                continue
            category = (e.get("category") or category_default or "other")
            items.append({
                "id": f"{_slug(key)}.{_slug(text[:60])}",
                "tier": _tier_of(e.get("category"), text),
                "category": str(category),
                "amount_usd": _num(e.get("amount_usd")),
                "text": text})
    if not items:
        absent.append({
            "kind": "no_open_items", "section": "flags",
            "what": "Flags and follow-ups",
            "reason": "the deal record carries no open request, unresolved question or "
                      "focus item for this borrower, so there is nothing to follow up."})
    return sorted(items, key=lambda i: (i["tier"],
                                        0 if i["amount_usd"] is not None else 1,
                                        -(i["amount_usd"] or 0.0),
                                        i["text"]))


# ------------------------------------------------------------------ the plan
# The Performance box compares the month against plan in its own `vs Budget`
# column, and drops that column when no budget is on file. The plan figures
# already exist -- budget-vs-actual published them -- so they are read from its
# result rather than left out, which is how the column went missing in silence.
BUDGET_METRICS = (
    ("revenue", "revenue"),
    ("gross_profit", "gross profit"),
    ("ebitda", "ebitda"),
    ("opex", "total opex"),
)


def latest_month_budget(archive):
    """The reported month's plan figures, in the shape `performance_views` wants.

    Stated on the MODEL'S OWN UNIT, which is how budget-vs-actual archived them.
    `performance_views.build` converts the plan beside the actual it is compared
    against, so converting here too would state the plan a thousandfold high for a
    borrower kept in thousands.

    Metric names carry the basis they were measured on ("EBITDA (as-reported)",
    "Total Opex (normalized)"), so each key matches on the metric's opening words
    rather than the whole string. Returns None when budget-vs-actual did not run,
    which is the honest "no budget on file" case the column is allowed to drop for.
    """
    entry = archive.get("budget-vs-actual")
    if not entry or entry.get("skipped"):
        return None
    out = {}
    for r in ((entry.get("data") or {}).get("results") or []):
        if not isinstance(r, dict) or str(r.get("scope") or "").lower() != "latest":
            continue
        name = str(r.get("metric") or "").strip().lower()
        for key, prefix in BUDGET_METRICS:
            if key not in out and name.startswith(prefix):
                v = _num(r.get("budget"))
                if v is not None:
                    out[key] = v
    if not out:
        return None
    rev = out.get("revenue")
    if rev:
        if "gross_profit" in out:
            out.setdefault("gross_margin", out["gross_profit"] / rev)
        if "ebitda" in out:
            out.setdefault("ebitda_margin", out["ebitda"] / rev)
        if "opex" in out:
            out.setdefault("opex_pct", out["opex"] / rev)
    return out


# ------------------------------------------------------------------ Performance box
def build_performance(cfg_path, period, absent, budget=None):
    """The Performance box, from the engine that already owns it.

    `performance_views.build` writes its payload to a file, so it gets a temp file
    in the system temp dir — never beside the payload, where a cloud-synced output
    folder can refuse the delete and strand the scratch file next to the memo. Its
    `generated` date is dropped: the memo has to re-render byte-identically from
    the same run, and a date stamp would make every re-render differ.

    `budget` carries the month's plan so the box keeps its `vs Budget` column.
    """
    fd, tmp = tempfile.mkstemp(prefix="memo-perf-", suffix=".json")
    os.close(fd)
    try:
        perf = performance_views.build(cfg_path, tmp, period=period, budget=budget)
    except Exception as e:                       # noqa: BLE001 - degrade, never fail
        absent.append({"kind": "section_unavailable",
                       "section": "snapshot_performance",
                       "what": "Performance",
                       "reason": f"the comparison views could not be built: {e}"})
        return None
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass
    perf.pop("generated", None)
    return perf


# ------------------------------------------------------------------ build
def build(cfg_path, period, out_path, archive_dir=None, model_path=None,
          trend_months=18):
    cfg = ml.load_config(cfg_path)
    borrower = (cfg.get("borrower") or {}).get("name") or cfg.get("borrower_name") \
        or (cfg.get("borrower") if isinstance(cfg.get("borrower"), str) else None)
    paths = cfg.get("paths") or {}
    out_folder = (paths.get("monitor_output_folder")
                  or cfg.get("monitor_output_folder"))
    archive_dir = archive_dir or (
        os.path.join(ml.resolve_path(out_folder), "_archive") if out_folder else None)
    if not archive_dir:
        raise PayloadError(
            "The borrower's settings do not say where its monthly results are "
            "saved, so there is nothing to build the memo from.")
    archive = read_archive(archive_dir, period)

    recorded_model_path = paths.get("model_output_path") or cfg.get("model_output_path")
    model_path = model_path or recorded_model_path
    recorded_model_path = recorded_path(recorded_model_path, model_path)
    model = ml.open_model(ml.resolve_path(model_path)) if model_path else None

    # Every figure read straight from the model is stated in dollars, which needs
    # the unit the model is kept in. A borrower whose config does not declare it
    # stops here rather than publishing a memo that may be out by a thousand.
    unit = ml.reporting_unit(cfg)

    sections = {s: {"figures": [], "chips": [], "groups": [],
                    **{k: (list(v) if isinstance(v, list)
                           else dict(v) if isinstance(v, dict) else v)
                       for k, v in (SECTION_SEEDS.get(s) or {}).items()}}
                for s in SECTIONS}
    absent, degrade = [], []

    # The store table every link on this memo is joined from. It comes from the
    # run's own results rather than the local machine, so the same archive gives
    # the same links wherever the memo is rendered.
    stores = next((archive[m]["data"].get("stores") for m in sorted(archive)
                   if not archive[m]["skipped"]
                   and archive[m]["data"].get("stores")), {}) or {}

    for mod in sorted(archive):
        entry = archive[mod]
        env = entry["data"]
        # A result file written before the envelope carried the model's unit says
        # nothing about what one of its figures stands for, and every money figure
        # in it would then be published unscaled -- Example Borrower's cash reading $10,700
        # against a $3,000 covenant floor rather than $10.7MM against $3.0MM. The
        # borrower's own settings DECLARE the unit, so that declaration stands in;
        # a file that states its own unit is never overridden, because a module may
        # deliberately publish in dollars whatever the model is kept in.
        if not (env.get("reporting_unit") or {}).get("scale"):
            env["reporting_unit"] = unit.as_dict()
            # Said out loud only where it changes a figure: a file in the folder
            # that carries no results carries no money either.
            if any(env.get(k) for k in ("results", "liquidity", "funded_debt")):
                absent.append({
                    "kind": "unit_assumed", "section": "data_quality",
                    "what": module_name(mod),
                    "reason": f"this month's {module_name(mod).lower()} result does "
                              f"not record what one figure in it stands for, so the "
                              f"unit this borrower's settings declare ({unit.plain}) "
                              f"was used to state its money figures in dollars.",
                    "module": mod})
        if entry["skipped"]:
            raw = env.get("reason") or ""
            absent.append({
                "kind": "module_skipped",
                "section": MODULE_SECTIONS.get(mod) or "data_quality",
                "what": module_name(mod),
                "reason": plain_reason(raw)
                          or "its inputs were not on file for this month",
                "detail": raw,
                "module": mod,
            })
            continue
        if mod == "covenant-compliance":
            figs, abs_, dg, chips, hist, grps = _covenant_figures(mod, env, period)
            sections["covenant_detail"]["figures"].extend(figs)
            sections["covenant_detail"]["history"] = hist
            sections["covenant_strip"]["chips"].extend(chips)
            sections["covenant_strip"]["either_or"] = grps
            # A breach at an earlier test date belongs on the front of the memo, not
            # only in the history box further down.
            sections["covenant_strip"]["prior_breaches"] = \
                unresolved_earlier_breaches(hist, period)
            absent.extend(abs_)
            degrade.extend(dg)
        else:
            fixed = MODULE_SECTIONS.get(mod, "data_quality")

            def section_of(r, _fixed=fixed):
                if _fixed is None:
                    return TSA_SECTIONS.get(str(r.get("section") or ""), "income_review")
                return _fixed

            figs, dg = _result_figures(mod, env, section_of)
            absent.extend(absent_reads(mod, env, section_of))
            lfigs, ldg, lgaps, lblock = liquidity_figures(env)
            figs += lfigs
            dg = sorted(set(dg) | set(ldg))
            absent.extend(lgaps)
            if lblock:
                sections["balance_review"]["liquidity"] = lblock
            for f in figs:
                # A figure the payload could not state safely becomes a line in the
                # memo's "could not cover" section rather than a number on the page.
                if f.get("unsafe"):
                    absent.append({"kind": "unit_unstated", "section": f["section"],
                                   "what": f.get("label") or f["id"],
                                   "reason": f["unsafe"], "module": mod})
                    continue
                seat(sections, f["section"])["figures"].append(f)
            degrade.extend(dg)
            # The two blocks that exist only when their module ran: the key is left
            # OFF the section when it did not, and the skip record covers it.
            if mod == "revenue-quality":
                rq = revenue_quality_block(env)
                if rq is not None:
                    sections["income_review"]["revenue_quality"] = rq
            if mod == "ap-ar-aging":
                ag = aging_block(env)
                if ag is not None:
                    sections["balance_review"]["aging"] = ag
        grp = debt_group(env)
        if grp:
            seat(sections, grp["section"], "traced group")["groups"].append(grp)
        for e in (env.get("exceptions") or []):
            if isinstance(e, dict) and str(e.get("kind") or "").endswith("coverage_gap"):
                absent.append({
                    "kind": "partial_months",
                    "section": "data_quality",
                    "what": e.get("label") or e.get("input") or mod,
                    "reason": e.get("message") or e.get("detail") or str(e),
                    "module": mod,
                })
        for miss in (env.get("lines_not_found") or []):
            absent.append({
                "kind": "line_not_found",
                "section": "data_quality",
                "what": str(miss),
                "reason": "the model has no row under this name, so figures that "
                          "need it are not shown",
                "module": mod,
            })

    # The Performance box and the windows come from the engines that already own
    # them, rather than being recomputed here.
    windows = period_windows.windows(period)
    # The against-plan column may be dropped only when there is genuinely no plan,
    # and the memo has to say which case it is -- a box that quietly comes out with
    # three columns instead of four looks finished.
    budget = latest_month_budget(archive)
    if budget is None:
        bva = archive.get("budget-vs-actual")
        absent.append({
            "kind": "column_absent", "section": "snapshot_performance",
            "what": "The against-plan column",
            "reason": ("no plan for this month was on file, so Performance compares "
                       "the month against earlier months only"
                       if bva is None or bva.get("skipped") else
                       "the plan on file carries no figure for any of the lines "
                       "Performance shows, so it compares the month against earlier "
                       "months only")})
    perf = build_performance(cfg_path, period, absent, budget=budget)
    if perf is not None:
        sections["snapshot_performance"]["performance_views"] = perf

    # What secures the loan, from the control file that was on disk all along.
    fdg = next((g for g in sections["deal_capstructure"]["groups"]
                if g.get("id") == "funded_debt"), None)
    collateral = build_collateral(
        cfg, model, period,
        (fdg or {}).get("total_ex_leases"), absent, unit)
    if collateral is not None:
        sections["snapshot_collateral"]["collateral"] = collateral
        # Publish each collateral figure the memo prints, so the gate holds the
        # page to the same agreement it holds every other number to.
        sections["snapshot_collateral"]["figures"].extend(
            collateral_figures(collateral))

    # The Deal summary tab: the deal as underwritten and documented, read once
    # from the deal record and the executed documents the settings name.
    dc, dc_path = read_deal_context(cfg, absent)
    spec = _maybe_json(paths.get("covenant_spec_path")) or {}
    if not spec:
        absent.append({
            "kind": "covenant_spec_absent", "section": "covenant_detail",
            "what": "The agreement's covenant terms",
            "reason": "the distilled terms of this borrower's agreement are not on "
                      "file, so thresholds, the rule that combines the tests and the "
                      "step-down schedule are not stated."})
    static = (dc or {}).get("static") or {}
    sections["header"]["borrower_legal_name"] = borrower
    sections["header"]["one_liner"] = static.get("business_one_liner")
    # The deal team is a plain list of the analyst's own colleagues, and a deal
    # record that carries none simply omits the line -- an empty label under the
    # title is worse than no line. A name the analyst did not confirm is dropped
    # and the run says so, so the banner never states a team nobody confirmed.
    team, team_gap = settled_deal_team(static)
    sections["header"]["deal_team"] = team
    if team_gap:
        absent.append({"kind": "deal_team_not_confirmed", "section": "header",
                       "what": "The deal team on the header", "reason": team_gap})
    view, disclaimed = settled_view((dc or {}).get("our_view"))
    if disclaimed:
        absent.append({
            "kind": "view_not_signed_off", "section": "header",
            "what": "Our standing view of the credit",
            "reason": "the deal record's own note says its watchlist and "
                      "direction were written during onboarding and still need the "
                      "covering analyst's sign-off, so they are left off this memo."})
    sections["header"]["our_view"] = view

    company = build_deal_company(dc, cfg, stores, absent)
    sections["deal_company"].update(company)
    terms = build_deal_terms(dc, dc_path, spec, stores, absent,
                             borrower_name=borrower)
    sections["deal_terms"]["rows"] = terms["rows"]
    sections["deal_terms"]["covenant_summary_row"] = terms["covenant_summary_row"]
    sections["deal_terms"]["figures"].extend(
        facility_figures(terms, terms.get("source_ref")))
    sections["deal_transaction"].update(
        build_deal_transaction(dc, cfg, stores, absent))
    security = build_deal_security(dc, cfg, period, stores, absent)
    sections["deal_security"].update(security)
    sections["deal_security"]["figures"].extend(security_figures(security))
    capstructure = build_deal_capstructure(dc, cfg, model, period, fdg, stores,
                                           absent, unit)
    sections["deal_capstructure"].update(capstructure)
    sections["deal_capstructure"]["figures"].extend(
        capstructure_figures(capstructure))

    # The Covenants tab's own blocks, over the covenant work already archived.
    cov_entry = archive.get("covenant-compliance")
    cov_env = (cov_entry or {}).get("data") or {}
    if cov_entry and not cov_entry.get("skipped"):
        # The borrower's own certified figures for the month, transcribed beside
        # the run's results -- read for the certificate's name, signer and date.
        certified = _maybe_json(
            f"{str(out_folder).rstrip('/')}/certified-values-{period}.json") \
            if out_folder else None
        sections["covenant_detail"]["preamble"] = covenant_preamble(
            cov_env, spec, cfg, certified, stores, absent)
        sections["covenant_detail"]["grid"] = covenant_grid(
            cov_env, spec, cfg, period, stores, absent)
        sections["covenant_detail"]["far_out"] = covenant_far_out(
            cov_env, spec, period, cfg, stores, absent)
        sections["covenant_detail"]["stress_inputs"] = covenant_stress_inputs(
            cov_env, spec, period, stores, cfg, absent)
        # Every group the strip has to name gets an entry, whether or not the run
        # emitted a group row. An archived run from before the engine emitted them
        # left the strip with an id and no label, and the id is what reached the
        # page -- so the label is settled here, where the spec is in hand.
        seated = {str(g.get("id")) for g in sections["covenant_strip"]["either_or"]}
        labels = either_or_labels(cov_env, spec)
        for gid in dict.fromkeys(
                str(c.get("either_or_group"))
                for c in sections["covenant_strip"]["chips"]
                if c.get("either_or_group")):
            if gid not in seated:
                sections["covenant_strip"]["either_or"].append(
                    {"id": gid, "label": labels.get(gid) or "Either/or test",
                     "name": None, "status": None})

    sections["files_registry"]["documents"] = build_files_registry(
        cfg, archive, archive_dir, period, stores, absent, spec)
    sections["flags"]["items"] = build_flags(dc, absent)

    growth = revenue_vs_last_year(sections, period)
    if growth is not None:
        sections["snapshot_performance"]["figures"].append(growth)

    if model is not None:
        for section, table in statement_tables(model, period,
                                               unit=unit).items():
            sections[section]["table"] = table
        for _sheet, section, _flow in STATEMENT_SHEETS:
            if not sections[section].get("table"):
                absent.append({
                    "kind": "table_absent", "section": section,
                    "what": section.replace("_", " "),
                    "reason": "the model has no rows this borrower's account list names "
                              "for this statement, so the table is not shown"})
        chart = trend_series(model, cfg, period, trend_months, unit=unit)
        sections["trend_chart"]["series"] = chart["series"]
        sections["trend_chart"]["months"] = chart["months"]
        idx, break_label, break_reason = basis_break_seat(cfg, chart["months"])
        sections["trend_chart"]["basis_break_index"] = idx
        quarters = quarter_callouts(chart["months"], chart["series"], idx)
        sections["trend_chart"]["quarters"] = quarters
        sections["trend_chart"]["notes"] = trend_notes(unit, break_label, quarters)
        if break_label:
            absent.append({
                "kind": "basis_break", "section": "trend_chart",
                "what": f"months either side of {break_label}",
                "reason": plain_reason(break_reason)
                          or "the borrower changed how it reports these figures at "
                             "this month, so months either side are not comparable"})
        for a in chart["absent"]:
            absent.append({"kind": "series_absent", "section": "trend_chart",
                           "what": a.get("series"), "reason": a.get("reason")})
        # A gap in the run of quarter boxes is only readable if it is stated. The
        # window's own first quarter is clipped by the window and is expected; a
        # quarter the chart draws in full and still cannot total is a data gap.
        drawn = {q["label"] for q in quarters}
        for label, _short_label, mos, _idx in whole_quarters_drawn(chart["months"]):
            if label not in drawn:
                absent.append({
                    "kind": "quarter_absent", "section": "trend_chart",
                    "what": f"the {label} total on the trend chart",
                    "reason": f"{label} is drawn in full but at least one of "
                              f"{', '.join(mos)} carries no figure, so the quarter "
                              f"cannot be added up and no box is shown for it"})
        dup = model.duplicate_label_warning()
        if dup:
            absent.append({"kind": "ambiguous_row", "section": "data_quality",
                           "what": "repeated row names", "reason": dup})
    else:
        absent.append({"kind": "series_absent", "section": "trend_chart",
                       "what": "all", "reason": "the model workbook was not given"})
        sections["trend_chart"]["notes"] = trend_notes(unit, None)

    # Written last, because the notes are about everything above them.
    sections["data_quality"]["notes"] = build_data_quality(
        cfg, period, set(degrade), absent, stores, budget is not None)

    # A record naming a section the memo does not carry would silently disappear
    # from the page, so it stops the run the same way a stray figure does.
    for a in absent:
        seat(sections, a.get("section") or "data_quality", "not-covered record")

    envelopes = {m: archive[m]["name"] for m in sorted(archive)}
    first_env = next((archive[m]["data"] for m in sorted(archive)
                      if not archive[m]["skipped"]), {})
    payload = {
        "schema": SCHEMA,
        "identity": {
            "borrower": borrower,
            "period": period,
            "fund": cfg.get("fund") or None,
            "model_path": recorded_model_path,
            "archive_dir": recorded_path(archive_dir),
            "archive_files": [envelopes[m] for m in sorted(envelopes)],
            "modules": sorted(envelopes),
            "reporting_unit": first_env.get("reporting_unit") or {},
            "months": list(model.months) if model is not None else [],
        },
        "windows": windows,
        "sections": {s: sections[s] for s in SECTIONS},
        "absent": sorted(absent, key=lambda a: (a.get("section") or "",
                                                a.get("kind") or "",
                                                str(a.get("what") or ""))),
        "degrade": sorted(set(degrade)),
        "stores": stores,
    }
    text = json.dumps(payload, indent=2, sort_keys=True, ensure_ascii=False)
    if out_path:
        with open(out_path, "w", encoding="utf-8", newline="\n") as fh:
            fh.write(text + "\n")
    return payload


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--borrower-config", required=True)
    ap.add_argument("--period", required=True, help='the month being reported, "YYYY-MM"')
    ap.add_argument("--out", required=True, help="where to write memo-payload.json")
    ap.add_argument("--archive", help="the run's _archive folder (defaults to the "
                                      "monitor output folder's _archive)")
    ap.add_argument("--model", help="the model workbook (defaults to the configured path)")
    ap.add_argument("--trend-months", type=int, default=18)
    a = ap.parse_args()
    if not re.fullmatch(r"\d{4}-\d{2}", a.period):
        print(f'--period must be "YYYY-MM" (got {a.period!r})', file=sys.stderr)
        return 2
    try:
        p = build(a.borrower_config, a.period, a.out, archive_dir=a.archive,
                  model_path=a.model, trend_months=a.trend_months)
    except (PayloadError, ml.ConfigError, ml.StoreAccessError) as e:
        print(str(e), file=sys.stderr)
        return 1
    n = sum(len(v.get("figures") or []) for v in p["sections"].values())
    print(f"OK — {n} figure(s), {len(p['absent'])} not covered, "
          f"{len(p['identity']['modules'])} result file(s) → {a.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
