#!/usr/bin/env python3
"""Workbook and memo content comparison — repackaging vs real edits.

Raw bytes change when Excel or SharePoint re-wraps a file; the cells may not.
Keep raw sha256 for exact-copy verification (manifests), but decide whether
recomputation is needed by comparing what a person would read: every cell's
value/formula for workbooks, visible text for memos.

Stdlib + openpyxl only. If openpyxl is missing, fall back to raw bytes.
"""

from __future__ import annotations

import html
import io
import os
import re
from html.parser import HTMLParser
from zipfile import BadZipFile, ZipFile

try:
    import openpyxl  # type: ignore
except ImportError:  # pragma: no cover
    openpyxl = None  # type: ignore





def _num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def _same_cell(a, b):
    if _blank(a) and _blank(b):
        return True
    if _num(a) and _num(b):
        return abs(a - b) <= max(abs(a), abs(b), 1.0) * 1e-12
    return a == b


def _grid_size(*sheets):
    return (max(ws.max_row for ws in sheets), max(ws.max_column for ws in sheets))


def _defined_names(book):
    """Workbook- and sheet-scoped named ranges, name -> formula/target text.

    A name redefinition (e.g. `TotalRevenue` now points at a different cell)
    changes what any formula referencing it computes, with no visible cell
    difference of its own -- it must not disappear from this comparison.
    """
    names = {name: dn.value for name, dn in book.defined_names.items()}
    for ws in book.worksheets:
        for name, dn in dict(getattr(ws, "defined_names", None) or {}).items():
            names[f"{ws.title}!{name}"] = dn.value
    return names


def _has_external_links(path):
    """Whether the raw zip package references an external workbook link.

    Checked directly against the package structure -- not via openpyxl's
    private, version-dependent `_external_links` attribute, which may be
    empty or incompletely populated for a link openpyxl cannot fully parse.
    An external link's cached values or target are not modelled by this
    comparison at all, present or not, so this must not silently trust an
    internal object's opinion of whether one exists.
    """
    try:
        with ZipFile(path) as z:
            return any(name.lower().startswith("xl/externallinks/") for name in z.namelist())
    except (OSError, BadZipFile):
        return True


def same_workbook_content(path_a, path_b):
    """True when two workbook files carry the same cells (values/formulas).

    Matches CAG's comparison standard: cell values/formulas, defined names
    (calculation-relevant, not a cell value), and external-link presence.
    Container-only re-wraps (SharePoint/Excel) that preserve cells are
    treated as same content; unfamiliar package parts are not separately
    checked -- CAG does not distinguish them from cell-level changes.
    """
    if openpyxl is None:
        with open(path_a, "rb") as fa, open(path_b, "rb") as fb:
            return fa.read() == fb.read()
    try:
        book_a = openpyxl.load_workbook(path_a, data_only=False)
        book_b = openpyxl.load_workbook(path_b, data_only=False)
    except Exception:
        return False
    if book_a.sheetnames != book_b.sheetnames:
        return False
    if _has_external_links(path_a) or _has_external_links(path_b):
        return False
    if _defined_names(book_a) != _defined_names(book_b):
        return False
    for name in book_a.sheetnames:
        ws_a, ws_b = book_a[name], book_b[name]
        rows, cols = _grid_size(ws_a, ws_b)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                if not _same_cell(ws_a.cell(r, c).value, ws_b.cell(r, c).value):
                    return False
    return True


def _col_empty(ws, c, rows):
    return all(_blank(ws.cell(r, c).value) for r in range(1, rows + 1))


def _held_workbook_reason(path_a, path_b):
    """Why a changed workbook cannot be taken as a single reported figure.

    Mirrors CAG's conservative classification: structural changes, formula
    cells, cleared figures and text-in-figure cells are held for explicit
    user clarification rather than silently pinned as a source figure.
    Returns None only when every difference is a plain reported-figure
    value change (number to number, or a number filled into a previously
    blank cell of an existing row and column).

    Fail-closed: an unreadable file, a missing classifier, a defined-name
    change or an external-link presence is held, never treated as eligible.
    """
    if openpyxl is None:
        return "the workbook could not be classified (openpyxl unavailable)"
    try:
        book_a = openpyxl.load_workbook(path_a, data_only=False)
        book_b = openpyxl.load_workbook(path_b, data_only=False)
    except Exception:
        return "the workbook could not be opened as a workbook"
    if book_a.sheetnames != book_b.sheetnames:
        return "a sheet was added or removed"
    if _defined_names(book_a) != _defined_names(book_b):
        return "a defined name was added, removed or repointed"
    if _has_external_links(path_a) or _has_external_links(path_b):
        return "the workbook references an external workbook link"
    for name in book_a.sheetnames:
        ws_a, ws_b = book_a[name], book_b[name]
        rows, cols = _grid_size(ws_a, ws_b)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                a, b = ws_a.cell(r, c).value, ws_b.cell(r, c).value
                if _same_cell(a, b):
                    continue
                coord = ws_a.cell(r, c).coordinate
                if _is_formula(a) or _is_formula(b):
                    return f"a formula cell changed ({name} {coord})"
                if _blank(b) and not _blank(a):
                    return f"a figure was cleared ({name} {coord})"
                if not ((_num(a) or _blank(a)) and (_num(b) or _blank(b))):
                    return f"text was typed in a figure cell ({name} {coord})"
                if _blank(a) and not _blank(b):
                    row_empty = all(_blank(ws_a.cell(r, cc).value) for cc in range(1, cols + 1))
                    if row_empty:
                        return f"a new row was added ({name} row {r})"
                    if _col_empty(ws_a, c, rows):
                        return f"a new column was added ({name} column {c})"
                    continue
                if _num(a) and _num(b):
                    continue
                return f"text was typed in a figure cell ({name} {coord})"
    return None


class _Text(HTMLParser):
    BLOCKS = {"p", "div", "br", "h1", "h2", "h3", "h4", "h5", "h6", "li", "tr", "section", "article"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "template"):
            self.skip += 1
        if tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "template") and self.skip:
            self.skip -= 1
        if tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def visible_lines(raw):
    """Memo as a reader sees it: tags gone, entities decoded, whitespace normalised."""
    p = _Text()
    p.feed(raw)
    p.close()
    text = html.unescape("".join(p.parts))
    lines = [re.sub(r"\s+", " ", ln).strip() for ln in text.split("\n")]
    return [ln for ln in lines if ln]


def same_memo_content(path_a, path_b):
    """True when two HTML memos carry the same visible text."""
    try:
        with open(path_a, encoding="utf-8", errors="replace") as fa:
            a = visible_lines(fa.read())
        with open(path_b, encoding="utf-8", errors="replace") as fb:
            b = visible_lines(fb.read())
    except OSError:
        return False
    return a == b


def same_content(path_a, path_b):
    """Whether two copies of one file carry the same content.

    - .xlsx/.xlsm: same cells (values/formulas), ignoring container repackaging
    - .html/.htm: same visible text, ignoring tags/whitespace
    - else: raw bytes equal
    """
    if not os.path.exists(path_a) or not os.path.exists(path_b):
        return False
    # Fast path: bytes equal -> content equal
    try:
        with open(path_a, "rb") as fa, open(path_b, "rb") as fb:
            if fa.read() == fb.read():
                return True
    except OSError:
        return False
    ext = os.path.splitext(str(path_a))[1].lower()
    # Use the extension of either path; they should match but be lenient
    if ext not in (".xlsx", ".xlsm", ".html", ".htm"):
        ext = os.path.splitext(str(path_b))[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return same_workbook_content(path_a, path_b)
    if ext in (".html", ".htm"):
        return same_memo_content(path_a, path_b)
    # Unknown type: bytes already differed -> not same
    return False


def content_changed(frozen_path, edited_path):
    """Classify an edited file against its frozen copy.

    Returns "same" (bytes equal), "same_content" (repackaging only),
    "changed" (real edit), or "missing" (edited file not there).
    """
    if not os.path.exists(edited_path):
        return "missing"
    if not os.path.exists(frozen_path):
        return "changed"
    try:
        with open(frozen_path, "rb") as fa, open(edited_path, "rb") as fb:
            if fa.read() == fb.read():
                return "same"
    except OSError:
        return "changed"
    if same_content(frozen_path, edited_path):
        return "same_content"
    return "changed"
